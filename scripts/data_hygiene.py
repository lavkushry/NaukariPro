import openpyxl
import datetime
import shutil
import sys
import os

def run_data_hygiene(tracker_path="Naukri_Live_Application_Tracker.xlsx"):
    if not os.path.exists(tracker_path):
        print(f"[!] {tracker_path} does not exist.")
        return False

    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    bak = f"{tracker_path}.bak.{ts}"
    shutil.copy2(tracker_path, bak)
    print(f"[+] Backed up tracker to {bak}")

    wb = openpyxl.load_workbook(tracker_path)

    # 1. Archive LinkedIn SUBMITTED_DRY_RUN rows
    if "LinkedIn_Applications_Log" in wb.sheetnames:
        li_ws = wb["LinkedIn_Applications_Log"]
        li_rows = list(li_ws.iter_rows(values_only=True))
        if li_rows:
            header = li_rows[0]
            data_rows = li_rows[1:]

            if "DryRun_Archive" not in wb.sheetnames:
                archive_ws = wb.create_sheet(title="DryRun_Archive")
                archive_ws.append(list(header))
            else:
                archive_ws = wb["DryRun_Archive"]

            dry_count = 0
            live_rows = []
            for r in data_rows:
                if len(r) > 6 and str(r[6]).upper() == "SUBMITTED_DRY_RUN":
                    archive_ws.append(list(r))
                    dry_count += 1
                else:
                    live_rows.append(r)

            # Clear and rebuild LinkedIn_Applications_Log with only live rows
            wb.remove(li_ws)
            new_li_ws = wb.create_sheet(title="LinkedIn_Applications_Log")
            new_li_ws.append(list(header))
            for lr in live_rows:
                new_li_ws.append(list(lr))
            print(f"[+] Moved {dry_count} SUBMITTED_DRY_RUN rows to DryRun_Archive. Live rows remaining in LinkedIn: {len(live_rows)}")

    # 2. Add columns to Applications_Log
    if "Applications_Log" in wb.sheetnames:
        app_ws = wb["Applications_Log"]
        current_headers = [cell.value for cell in app_ws[1] if cell.value is not None]
        new_cols = [
            "Responded (Y/N)",
            "Response date",
            "Call received (Y/N)",
            "Call source (Resdex / application / LinkedIn / referral)",
            "Outcome"
        ]
        cols_to_add = [c for c in new_cols if c not in current_headers]
        if cols_to_add:
            start_col = len(current_headers) + 1
            for offset, c in enumerate(cols_to_add):
                col_idx = start_col + offset
                app_ws.cell(row=1, column=col_idx, value=c)

            # Backfill defaults
            for row_idx in range(2, app_ws.max_row + 1):
                app_ws.cell(row=row_idx, column=12, value="N")
                app_ws.cell(row=row_idx, column=13, value="")
                app_ws.cell(row=row_idx, column=14, value="N")
                app_ws.cell(row=row_idx, column=15, value="application")
                app_ws.cell(row=row_idx, column=16, value="Under Review")
            print(f"[+] Added columns {cols_to_add} to Applications_Log and backfilled defaults across {app_ws.max_row - 1} rows.")

    wb.save(tracker_path)
    print(f"[✓] Data hygiene complete on {tracker_path}!")
    return True

if __name__ == "__main__":
    path = sys.argv[1] if len(sys.argv) > 1 else "Naukri_Live_Application_Tracker.xlsx"
    run_data_hygiene(path)
