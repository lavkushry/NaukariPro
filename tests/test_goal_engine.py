"""
Unit & Integration Tests for 5-Calls/Day Recruiter Outreach & Goal Pacing Engine
Validates:
- 5 message templates generation
- Strict <= 300 character limit on Connection Request Notes
- Recruiter Outreach CRM (JSON & Excel Recruiter_Outreach_Log sync)
- Funnel pacing calculations, conversion rates & diagnostic matrix
- 3-Block execution orchestrator
- CLI command dispatching & dry-run isolation
100% offline, deterministic, zero network calls.
"""

import os
import json
import pytest
import openpyxl
import asyncio
from datetime import datetime
from unittest.mock import MagicMock, AsyncMock, patch

from naukari_pro.modules.recruiter_outreach import (
    RecruiterTemplateGenerator,
    RecruiterOutreachCRM,
    VALID_STATUSES,
    TOP_GCC_TARGETS,
    CANDIDATE_FACTS
)
from naukari_pro.modules.linkedin_outreach import (
    LinkedInOutreachSender,
    DIRECT_CONNECT_SELECTORS,
    OVERFLOW_CONNECT_SELECTORS,
    MORE_SELECTORS,
    ADD_NOTE_SELECTORS,
    TEXTAREA_SELECTORS,
    DISMISS_SELECTORS,
    SEND_SELECTORS,
    PENDING_SELECTORS,
    CONNECTED_SELECTORS
)
from naukari_pro.modules.goal_5calls import (
    Goal5CallsEngine,
    DAILY_TARGETS,
    DIAGNOSTIC_RULES
)
from naukari_pro.cli import handle_goal_5calls, main


# ==============================================================================
# SECTION 1: TEMPLATE GENERATION & CHARACTER CONSTRAINTS
# ==============================================================================

def test_template_generator_initialization():
    gen = RecruiterTemplateGenerator()
    assert gen.name == "Lavkush Kumar"
    assert "2.5 YOE" in gen.yoe
    assert gen.phone == "+91-9122036484"
    assert gen.email == "lavkushry@gmail.com"


@pytest.mark.parametrize("recruiter_name,company,stack", [
    ("Priya Sharma", "Walmart Global Tech", "Azure Databricks, PySpark & Delta Lake"),
    ("Rajesh Kumar", "Microsoft India Development Center", "Databricks & Delta Lake"),
    ("Sneha", "Amazon Web Services (AWS)", "Distributed PySpark & Kafka"),
    ("Christopher Nolan-Smith-Montgomery", "A Very Long Enterprise Corporation Private Limited", "Databricks, PySpark, Delta Lake, ADF, Kafka, SQL, ADLS Gen2, Terraform"),
    ("", "", "Databricks")
])
def test_connection_note_strict_300_character_limit(recruiter_name, company, stack):
    """
    CRITICAL CONSTRAINT: LinkedIn connection notes must be strictly <= 300 characters.
    Must never exceed 300 chars even with long names or long company names.
    """
    gen = RecruiterTemplateGenerator()
    note = gen.generate_connection_note(recruiter_name, company, stack)
    assert len(note) <= 300, f"Connection note exceeded 300 chars: {len(note)} chars ('{note}')"
    assert "Lavkush" in note
    assert "2.5 YOE" in note
    assert "15-day notice" in note.lower() or "15 days" in note.lower()


def test_direct_inmail_template_content():
    """Verify Direct InMail includes 2.5 YOE, 2M+ events, 15-day notice, and contact details."""
    gen = RecruiterTemplateGenerator()
    inmail = gen.generate_direct_inmail("Anjali Roy", "JPMorgan Chase")
    assert "subject" in inmail and "body" in inmail
    subj = inmail["subject"]
    body = inmail["body"]

    assert "Azure Data Engineer" in subj
    assert "15-Day Notice" in subj
    assert "Lavkush Kumar" in subj

    assert "Hi Anjali," in body
    assert "JPMorgan Chase" in body
    assert "2M+ events/day" in body
    assert "35%" in body
    assert "+91-9122036484" in body
    assert "lavkushry@gmail.com" in body
    assert "linkedin.com/in/lavkushry" in body
    assert "₹20–25 LPA" in body


def test_tech_lead_pitch_content():
    """Verify Tech Lead Outreach highlights Lakehouse scaling, Delta ACID, and DataOps."""
    gen = RecruiterTemplateGenerator()
    pitch = gen.generate_tech_lead_pitch("Amit Verma", "PNC Bank GCC")
    assert "subject" in pitch and "body" in pitch
    subj = pitch["subject"]
    body = pitch["body"]

    assert "Architecture & Lakehouse Scaling" in subj
    assert "Hi Amit," in body
    assert "PNC Bank GCC" in body
    assert "Lakehouse Scaling" in body
    assert "Delta Lake ACID transactions" in body
    assert "Query & Pipeline Optimization" in body
    assert "35%" in body
    assert "Terraform" in body


def test_48h_followup_urgency():
    """Verify 48-Hour Follow-Up is concise, high urgency, and reminds of 15-day notice."""
    gen = RecruiterTemplateGenerator()
    followup = gen.generate_48h_followup("Rohan Gupta", "Barclays")
    assert "subject" in followup and "body" in followup
    assert "Re:" in followup["subject"]
    body = followup["body"]
    assert "Hi Rohan," in body
    assert "Barclays" in body
    assert "active interview rounds" in body.lower()
    assert "15-day notice period" in body.lower()
    assert "+91-9122036484" in body


def test_application_followup():
    """Verify external portal application follow-up references portal and stack."""
    gen = RecruiterTemplateGenerator()
    app_f = gen.generate_application_followup("Vikram", "Adobe", req_id_or_portal="Greenhouse #9812")
    assert "subject" in app_f and "body" in app_f
    assert "Greenhouse #9812" in app_f["subject"]
    body = app_f["body"]
    assert "Greenhouse #9812" in body
    assert "Adobe" in body
    assert "2M+ events/day" in body


def test_template_dispatcher():
    """Verify generate_template dispatches correctly to all 5 template types."""
    gen = RecruiterTemplateGenerator()
    t1 = gen.generate_template("connect", "Sarah", "Google")
    assert len(t1["body"]) <= 300

    t2 = gen.generate_template("inmail", "Sarah", "Google")
    assert "2M+ events" in t2["body"]

    t3 = gen.generate_template("tech-lead", "Dave", "Google")
    assert "Architecture" in t3["subject"]

    t4 = gen.generate_template("48h", "Sarah", "Google")
    assert "active interview rounds" in t4["body"]

    t5 = gen.generate_template("app-followup", "Sarah", "Google", portal_or_req_id="Workday R-102")
    assert "Workday R-102" in t5["body"]


# ==============================================================================
# SECTION 2: RECRUITER OUTREACH CRM & EXCEL LOGGING
# ==============================================================================

def test_crm_json_and_excel_synchronization(tmp_path):
    """
    Verify CRM persists leads to JSON and synchronizes with
    Recruiter_Outreach_Log sheet in Excel tracker.
    """
    json_file = str(tmp_path / "outreach_tracker.json")
    excel_file = str(tmp_path / "Tracker.xlsx")

    # Create dummy workbook
    wb = openpyxl.Workbook()
    wb.active.title = "Applications_Log"
    wb.save(excel_file)

    crm = RecruiterOutreachCRM(json_path=json_file, excel_path=excel_file)

    # 1. Log an outreach
    lead = crm.log_outreach(
        recruiter_name="Sunil Mehta",
        company="Microsoft",
        target_role="Azure Data Engineer",
        platform="LinkedIn",
        message_type="Direct InMail",
        status="CONTACTED",
        notes="First outreach note",
        message_content="Sample pitch"
    )

    assert lead["lead_id"].startswith("LEAD-")
    assert lead["status"] == "CONTACTED"
    assert os.path.exists(json_file)

    # Verify JSON content
    with open(json_file, "r", encoding="utf-8") as f:
        data = json.load(f)
    assert len(data) == 1
    assert data[0]["recruiter_name"] == "Sunil Mehta"

    # Verify Excel Sheet content
    wb_read = openpyxl.load_workbook(excel_file, read_only=True)
    assert "Recruiter_Outreach_Log" in wb_read.sheetnames
    ws = wb_read["Recruiter_Outreach_Log"]
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) == 2  # header + 1 data row
    assert rows[0] == (
        "Date", "Recruiter Name", "Company", "Target Role",
        "Platform", "Message Type", "Status", "Notes"
    )
    assert rows[1][1] == "Sunil Mehta"
    assert rows[1][2] == "Microsoft"
    assert rows[1][6] == "CONTACTED"


def test_crm_update_status(tmp_path):
    """Verify lead status update updates both JSON and Excel."""
    json_file = str(tmp_path / "outreach_tracker.json")
    excel_file = str(tmp_path / "Tracker.xlsx")

    crm = RecruiterOutreachCRM(json_path=json_file, excel_path=excel_file)
    lead = crm.log_outreach("Ananya", "Walmart Global Tech", status="CONTACTED")

    # Update to CALL_SCHEDULED
    updated = crm.update_status("Walmart", "CALL_SCHEDULED", notes="Round 1 booked for Thursday")
    assert updated is not None
    assert updated["status"] == "CALL_SCHEDULED"
    assert "Round 1 booked for Thursday" in updated["notes"]

    # Verify invalid status throws ValueError
    with pytest.raises(ValueError):
        crm.update_status("Walmart", "INVALID_STATUS")

    # Verify Excel reflected update
    wb = openpyxl.load_workbook(excel_file, read_only=True)
    ws = wb["Recruiter_Outreach_Log"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[1][6] == "CALL_SCHEDULED"
    assert "Round 1 booked for Thursday" in rows[1][7]


def test_crm_batch_generate_gcc_outreach(tmp_path):
    """Verify batch generating GCC outreach logs 10 target companies."""
    json_file = str(tmp_path / "outreach_tracker.json")
    excel_file = str(tmp_path / "Tracker.xlsx")

    crm = RecruiterOutreachCRM(json_path=json_file, excel_path=excel_file)
    pitches = crm.batch_generate_gcc_outreach()

    assert len(pitches) == 10
    companies = [p["lead"]["company"] for p in pitches]
    assert "Microsoft" in companies
    assert "Walmart Global Tech" in companies
    assert "JPMorgan Chase" in companies
    assert "Amazon" in companies

    # Check that all generated connection notes strictly adhere to <= 300 chars
    for p in pitches:
        assert len(p["connection_note"]) <= 300

    summary = crm.get_summary()
    assert summary["total_leads"] == 10
    assert summary["today_outreaches"] == 10


# ==============================================================================
# SECTION 3: 5-CALLS/DAY PACING & FUNNEL DIAGNOSTICS
# ==============================================================================

def test_goal_pacing_and_diagnostics_calculation(tmp_path):
    """
    Verify funnel metrics computation, conversion ratios, and diagnostic matrix rules.
    """
    json_file = str(tmp_path / "outreach_tracker.json")
    excel_file = str(tmp_path / "Tracker.xlsx")

    # Set up workbook with mock application and performance rows
    wb = openpyxl.Workbook()
    ws_app = wb.active
    ws_app.title = "Applications_Log"
    ws_app.append(["Application Date", "Company Name", "Job Title", "Match Score (%)"])
    today_str = datetime.now().strftime("%Y-%m-%d %H:%M")
    # Add 30 applications today
    for i in range(30):
        ws_app.append([today_str, f"Company_{i}", "Azure Data Engineer", "85%"])

    # Performance log: low contact views
    ws_perf = wb.create_sheet(title="Recruiter_Performance_Log")
    ws_perf.append(["Timestamp", "Search Appearances", "Profile Views", "Downloads", "Contact Views"])
    ws_perf.append([today_str, 250, 15, 2, 2])

    wb.save(excel_file)

    engine = Goal5CallsEngine(excel_path=excel_file, json_path=json_file)

    # 1. Add mock outreach leads
    engine.crm.log_outreach("Recruiter A", "Microsoft", status="CONNECTED")
    engine.crm.log_outreach("Recruiter B", "Walmart", status="CALL_SCHEDULED")
    engine.crm.log_outreach("Recruiter C", "Amazon", status="PITCH_SENT")

    metrics = engine.calculate_funnel_metrics()

    assert metrics["applications_today"] == 30
    assert metrics["applications_target"] == 35
    assert metrics["outreaches_today"] == 3
    assert metrics["calls_today"] == 1
    assert metrics["calls_target"] == 5
    assert metrics["search_appearances"] == 250
    assert metrics["contact_views"] == 2

    # Diagnostics verification
    diagnostics = engine.run_diagnostics(metrics)
    diag_ids = [d["id"] for d in diagnostics]

    # Should flag LOW_OUTREACH_VOLUME (3 < 10)
    assert "LOW_OUTREACH_VOLUME" in diag_ids
    # Should flag LOW_CONTACT_VIEWS (contact views 2 < 5 with searches 250)
    assert "LOW_CONTACT_VIEWS" in diag_ids
    # Should flag LOW_CALL_CONVERSION (1 < 5)
    assert "LOW_CALL_CONVERSION" in diag_ids


def test_goal_dashboard_string_format(tmp_path):
    """Verify format_dashboard outputs human-readable terminal dashboard with metrics."""
    json_file = str(tmp_path / "outreach_tracker.json")
    excel_file = str(tmp_path / "Tracker.xlsx")

    engine = Goal5CallsEngine(excel_path=excel_file, json_path=json_file)
    engine.crm.batch_generate_gcc_outreach()

    dashboard = engine.format_dashboard()
    assert "5-CALLS/DAY GOAL PACING & RECRUITER OUTREACH DASHBOARD" in dashboard
    assert "DAILY FUNNEL VELOCITY" in dashboard
    assert "Applications Sent:" in dashboard
    assert "Recruiter Outreaches:" in dashboard
    assert "Recruiter Calls Scheduled:" in dashboard
    assert "DAILY 3-BLOCK SCHEDULE:" in dashboard


# ==============================================================================
# SECTION 4: 3-BLOCK EXECUTION ORCHESTRATOR
# ==============================================================================

def test_daily_3_block_orchestrator(tmp_path):
    """Verify execution of Block 1, Block 2, and Block 3."""
    json_file = str(tmp_path / "outreach_tracker.json")
    excel_file = str(tmp_path / "Tracker.xlsx")

    engine = Goal5CallsEngine(excel_path=excel_file, json_path=json_file)

    # Block 1 (Morning)
    b1 = engine.run_block(block_number=1, dry_run=True)
    assert b1["status"] == "COMPLETED"
    assert len(b1["actions_executed"]) == 2
    assert "Resdex" in b1["actions_executed"][0]["action"]
    assert b1["actions_executed"][1]["target_count"] == 25

    # Block 2 (Afternoon)
    b2 = engine.run_block(block_number=2, dry_run=False)
    assert b2["status"] == "COMPLETED"
    assert b2["actions_executed"][0]["count"] == 10
    assert len(engine.crm.leads) == 10

    # Block 3 (Evening)
    b3 = engine.run_block(block_number=3, dry_run=False)
    assert b3["status"] == "COMPLETED"
    assert len(b3["actions_executed"]) == 3
    # Follow-ups executed
    assert b3["actions_executed"][1]["count"] == 5

    # Invalid block throws ValueError
    with pytest.raises(ValueError):
        engine.run_block(block_number=99)


# ==============================================================================
# SECTION 5: CLI DISPATCHING & INTEGRATION
# ==============================================================================

def test_cli_handle_goal_5calls_pacing(capsys):
    """Verify CLI --pacing option prints dashboard cleanly."""
    mock_args = MagicMock()
    mock_args.pacing = True
    mock_args.outreach = False
    mock_args.status = False
    mock_args.block = None
    mock_args.update_status = None
    mock_args.dry_run = True

    ret = handle_goal_5calls(args=mock_args)
    assert ret == 0
    captured = capsys.readouterr()
    assert "5-CALLS/DAY GOAL PACING" in captured.out


def test_cli_handle_goal_5calls_single_outreach(capsys):
    """Verify CLI single recruiter outreach generation in dry-run mode."""
    mock_args = MagicMock()
    mock_args.pacing = False
    mock_args.outreach = True
    mock_args.status = False
    mock_args.block = None
    mock_args.update_status = None
    mock_args.dry_run = True
    mock_args.company = "Microsoft"
    mock_args.recruiter = "Kavita Singh"
    mock_args.role = "Azure Data Engineer"
    mock_args.msg_type = "tech-lead"
    mock_args.notes = "UnitTest"

    ret = handle_goal_5calls(args=mock_args)
    assert ret == 0
    captured = capsys.readouterr()
    assert "RECRUITER OUTREACH TEMPLATES: MICROSOFT" in captured.out
    assert "Kavita Singh" in captured.out
    assert "LINKEDIN CONNECTION NOTE" in captured.out
    assert "Dry-run mode" in captured.out


def test_cli_handle_goal_5calls_block_execution(capsys):
    """Verify CLI --block 1 executes block 1 cleanly."""
    mock_args = MagicMock()
    mock_args.pacing = False
    mock_args.outreach = False
    mock_args.status = False
    mock_args.block = 1
    mock_args.update_status = None
    mock_args.dry_run = True

    ret = handle_goal_5calls(args=mock_args)
    assert ret == 0
    captured = capsys.readouterr()
    assert "DAILY BLOCK 1 EXECUTION SUMMARY" in captured.out
    assert "Resdex Profile Timestamp Refresh" in captured.out


def test_cli_handle_goal_5calls_status_view(capsys):
    """Verify CLI --status displays active CRM summary."""
    mock_args = MagicMock()
    mock_args.pacing = False
    mock_args.outreach = False
    mock_args.status = True
    mock_args.block = None
    mock_args.update_status = None
    mock_args.dry_run = True

    ret = handle_goal_5calls(args=mock_args)
    assert ret == 0
    captured = capsys.readouterr()
    assert "RECRUITER OUTREACH CRM & ACTIVE PIPELINE STATUS" in captured.out


# ==============================================================================
# SECTION 6: ADVERSARIAL STRESS & CORNER-CASE TESTS
# ==============================================================================

def test_salutation_intelligence_anti_bot():
    """
    CRITICAL: Verify salutation generator NEVER emits bot-like greetings
    like 'Hi Technical,', 'Hi Senior,', 'Hi Corporate,', 'Hi AWS,'.
    """
    gen = RecruiterTemplateGenerator()

    # 1. Genuine personal names (with or without titles)
    assert gen._format_salutation("Priya Sharma", "Walmart") == "Hi Priya,"
    assert gen._format_salutation("Mr. Rajesh Kumar", "Microsoft") == "Hi Rajesh,"
    assert gen._format_salutation("Dr. Amit Verma", "PNC Bank") == "Hi Amit,"
    assert gen._format_salutation("Prof. Sarah Jenkins", "MIT") == "Hi Sarah,"
    assert gen._format_salutation("Sneha Kulkarni (Senior Recruiter)", "HSBC") == "Hi Sneha,"

    # 2. Departmental / Role titles (must fall back to team or generic salutation)
    sal1 = gen._format_salutation("Technical Talent Acquisition - Data & AI", "Microsoft")
    assert sal1 == "Hi Microsoft Recruiting Team,"
    assert "Technical" not in sal1

    sal2 = gen._format_salutation("Senior Tech Recruiter - Data Engineering", "Walmart Global Tech")
    assert sal2 == "Hi Walmart Global Tech Team,"
    assert "Senior" not in sal2

    sal3 = gen._format_salutation("Corporate Technology Talent Acquisition", "JPMorgan Chase")
    assert sal3 == "Hi JPMorgan Chase Recruiting Team,"
    assert "Corporate" not in sal3

    sal4 = gen._format_salutation("AWS / Retail Data Tech Recruiting", "Amazon")
    assert sal4 == "Hi Amazon Recruiting Team,"
    assert "AWS" not in sal4

    sal5 = gen._format_salutation("GCC Campus & Lateral Talent Acquisition", "PNC Bank GCC")
    assert sal5 == "Hi PNC Bank GCC Recruiting Team,"
    assert not sal5.startswith("Hi GCC,")

    # 3. Empty or None recruiter
    assert gen._format_salutation("", "Barclays") == "Hi Barclays Recruiting Team,"
    assert gen._format_salutation("", "") == "Hi there,"


def test_connection_note_intact_signoff_extreme_company_length():
    """
    Verify that extreme company lengths (150+ chars) never slice off the
    candidate sign-off '- Lavkush' and strictly remain <= 300 characters.
    """
    gen = RecruiterTemplateGenerator()
    giant_company = "Global Enterprise Mega Financial Services Cloud And Big Data Technologies Platform Solutions International Private Limited"
    note = gen.generate_connection_note("Priya Sharma", giant_company)

    assert len(note) <= 300, f"Exceeded 300 chars: {len(note)}"
    assert note.endswith("- Lavkush"), f"Sign-off was truncated: '{note}'"
    assert "Hi Priya," in note


def test_template_dispatcher_flexible_portal_kwargs():
    """Verify generate_template accepts 'portal' directly as keyword argument."""
    gen = RecruiterTemplateGenerator()
    res = gen.generate_template("app-followup", "Neha", "Salesforce", portal="Workday R-90210")
    assert "Workday R-90210" in res["subject"]
    assert "Workday R-90210" in res["body"]


def test_crm_preserves_lead_id_and_tracks_call_scheduled_at(tmp_path):
    """
    Verify:
    1. Re-logging an existing lead preserves the original lead_id.
    2. Setting CALL_SCHEDULED records call_scheduled_at timestamp.
    3. Updating notes later does not overwrite call_scheduled_at.
    """
    json_file = str(tmp_path / "outreach.json")
    excel_file = str(tmp_path / "tracker.xlsx")
    crm = RecruiterOutreachCRM(json_path=json_file, excel_path=excel_file)

    entry1 = crm.log_outreach("Kavita Rao", "Microsoft", status="CONTACTED")
    original_id = entry1["lead_id"]
    assert original_id.startswith("LEAD-")

    # Update with new message type (e.g. 48h followup)
    entry2 = crm.log_outreach("Kavita Rao", "Microsoft", message_type="48-Hour Follow-Up", status="CONNECTED")
    assert entry2["lead_id"] == original_id, "lead_id was wiped or regenerated on update!"

    # Update to CALL_SCHEDULED
    entry3 = crm.update_status("Microsoft", "CALL_SCHEDULED", notes="First round booked")
    assert entry3["status"] == "CALL_SCHEDULED"
    call_ts = entry3.get("call_scheduled_at")
    assert call_ts is not None

    # Later update notes only
    entry4 = crm.update_status("Microsoft", "CALL_SCHEDULED", notes="Added interviewer prep notes")
    assert entry4["call_scheduled_at"] == call_ts, "call_scheduled_at was mutated on subsequent note update!"


def test_excel_styling_and_column_widths(tmp_path):
    """Verify that Recruiter_Outreach_Log sheet is formatted with executive styling."""
    json_file = str(tmp_path / "outreach.json")
    excel_file = str(tmp_path / "tracker.xlsx")
    crm = RecruiterOutreachCRM(json_path=json_file, excel_path=excel_file)

    crm.log_outreach("Priya Sharma", "Walmart Global Tech", status="CALL_SCHEDULED", notes="Screening scheduled")
    crm.log_outreach("Rohan Gupta", "JPMorgan", status="PITCH_SENT")

    wb = openpyxl.load_workbook(excel_file)
    assert "Recruiter_Outreach_Log" in wb.sheetnames
    ws = wb["Recruiter_Outreach_Log"]

    # Header row height and fill
    assert ws.row_dimensions[1].height == 28
    assert ws.cell(row=1, column=1).fill.start_color.rgb in ["001E3A8A", "1E3A8A"]
    assert ws.cell(row=1, column=1).font.color.rgb in ["00FFFFFF", "FFFFFF"]

    # Column widths configured
    assert ws.column_dimensions["A"].width == 18
    assert ws.column_dimensions["C"].width == 25
    assert ws.column_dimensions["H"].width == 45

    # Status cell styling
    # Row 2 is Walmart (CALL_SCHEDULED) -> soft emerald fill
    assert ws.cell(row=2, column=7).fill.start_color.rgb in ["00D1FAE5", "D1FAE5"]


def test_cli_interactive_handler_clean_exit(monkeypatch):
    """Verify handle_goal_5calls_interactive terminates cleanly on choice '0'."""
    from naukari_pro.cli import handle_goal_5calls_interactive
    monkeypatch.setattr("builtins.input", lambda prompt="": "0")
    ret = handle_goal_5calls_interactive()
    assert ret == 0


# ==============================================================================
# SECTION 5: AUTONOMOUS LINKEDIN RECRUITER OUTREACH BOT (PLAYWRIGHT STEALTH)
# ==============================================================================

def make_mock_outreach_locator(count: int = 1, visible: bool = True, text: str = "", attr_map: dict = None):
    """Factory creating an asynchronous Playwright locator mock."""
    loc = MagicMock()
    loc.count = AsyncMock(return_value=count)
    loc.is_visible = AsyncMock(return_value=visible)
    loc.inner_text = AsyncMock(return_value=text)
    attr_map = attr_map or {}
    loc.get_attribute = AsyncMock(side_effect=lambda a: attr_map.get(a, None))
    loc.click = AsyncMock(return_value=None)
    loc.focus = AsyncMock(return_value=None)
    loc.fill = AsyncMock(return_value=None)
    loc.type = AsyncMock(return_value=None)
    loc.press_sequentially = AsyncMock(return_value=None)
    loc.first = loc
    loc.last = loc
    loc.nth = MagicMock(return_value=loc)
    return loc


def make_mock_outreach_page(url="https://www.linkedin.com/in/kavita-rao/", locator_map=None, body_content=""):
    """Factory creating an asynchronous Playwright page mock."""
    page = MagicMock()
    page.url = url
    page.goto = AsyncMock(return_value=None)
    page.content = AsyncMock(return_value=f"<html><body>{body_content}</body></html>")
    page.wait_for_timeout = AsyncMock(return_value=None)

    locator_map = locator_map or {}
    fallback_loc = make_mock_outreach_locator(count=0, visible=False)

    def get_locator(selector: str):
        if selector in locator_map:
            return locator_map[selector]
        return fallback_loc

    page.locator = MagicMock(side_effect=get_locator)
    return page


def test_outreach_direct_connect_add_note_send(tmp_path):
    """
    Verify full live outreach execution:
    1. Direct Connect button click
    2. Add a note modal trigger
    3. Humanized typing into message textarea (<= 300 chars)
    4. Send button click
    5. CRM logging with status CONTACTED
    """
    json_path = str(tmp_path / "outreach.json")
    excel_path = str(tmp_path / "tracker.xlsx")
    crm = RecruiterOutreachCRM(json_path=json_path, excel_path=excel_path)
    sender = LinkedInOutreachSender(crm=crm)

    connect_loc = make_mock_outreach_locator(count=1, visible=True)
    add_note_loc = make_mock_outreach_locator(count=1, visible=True)
    textarea_loc = make_mock_outreach_locator(count=1, visible=True)
    send_loc = make_mock_outreach_locator(count=1, visible=True)

    locator_map = {
        DIRECT_CONNECT_SELECTORS[0]: connect_loc,
        ADD_NOTE_SELECTORS[0]: add_note_loc,
        TEXTAREA_SELECTORS[0]: textarea_loc,
        SEND_SELECTORS[0]: send_loc,
    }
    mock_page = make_mock_outreach_page(locator_map=locator_map)

    res = asyncio.run(
        sender.send_outreach(
            recruiter_name="Priya Sharma",
            company="Walmart Global Tech",
            profile_url="https://www.linkedin.com/in/priya-sharma-walmart/",
            dry_run=False,
            page=mock_page
        )
    )

    assert res["status"] == "SENT"
    assert res["sent"] is True
    assert res["dry_run"] is False
    assert res["note_length"] <= 300
    assert "Walmart" in res["note"]

    # Verify DOM actions
    connect_loc.click.assert_awaited_once()
    add_note_loc.click.assert_awaited_once()
    send_loc.click.assert_awaited_once()
    assert textarea_loc.type.await_count > 0 or textarea_loc.press_sequentially.await_count > 0 or textarea_loc.fill.await_count > 0

    # Verify CRM updated to CONTACTED
    leads = crm.list_active_leads()
    assert len(leads) == 1
    assert leads[0]["recruiter_name"] == "Priya Sharma"
    assert leads[0]["company"] == "Walmart Global Tech"
    assert leads[0]["status"] == "CONTACTED"
    assert leads[0]["message_type"] == "Connection Request Note"


def test_outreach_overflow_more_menu_connect(tmp_path):
    """
    Verify fallback from missing direct Connect button to More overflow menu:
    1. Direct Connect buttons are absent
    2. More actions dropdown button clicked
    3. Overflow Connect menu item clicked
    4. Note added and sent
    """
    json_path = str(tmp_path / "outreach.json")
    excel_path = str(tmp_path / "tracker.xlsx")
    crm = RecruiterOutreachCRM(json_path=json_path, excel_path=excel_path)
    sender = LinkedInOutreachSender(crm=crm)

    more_loc = make_mock_outreach_locator(count=1, visible=True)
    overflow_connect_loc = make_mock_outreach_locator(count=1, visible=True)
    add_note_loc = make_mock_outreach_locator(count=1, visible=True)
    textarea_loc = make_mock_outreach_locator(count=1, visible=True)
    send_loc = make_mock_outreach_locator(count=1, visible=True)

    locator_map = {
        MORE_SELECTORS[0]: more_loc,
        OVERFLOW_CONNECT_SELECTORS[0]: overflow_connect_loc,
        ADD_NOTE_SELECTORS[0]: add_note_loc,
        TEXTAREA_SELECTORS[0]: textarea_loc,
        SEND_SELECTORS[0]: send_loc,
    }
    mock_page = make_mock_outreach_page(locator_map=locator_map)

    res = asyncio.run(
        sender.send_outreach(
            recruiter_name="Rohan Gupta",
            company="JPMorgan Chase",
            dry_run=False,
            page=mock_page
        )
    )

    assert res["status"] == "SENT"
    assert res["sent"] is True
    more_loc.click.assert_awaited_once()
    overflow_connect_loc.click.assert_awaited_once()
    send_loc.click.assert_awaited_once()


def test_outreach_already_connected_detection(tmp_path):
    """Verify that an existing 1st-degree connection is detected without clicking Connect."""
    json_path = str(tmp_path / "outreach.json")
    crm = RecruiterOutreachCRM(json_path=json_path)
    sender = LinkedInOutreachSender(crm=crm)

    connected_loc = make_mock_outreach_locator(count=1, visible=True, text="1st")
    connect_loc = make_mock_outreach_locator(count=1, visible=True)

    locator_map = {
        CONNECTED_SELECTORS[0]: connected_loc,
        DIRECT_CONNECT_SELECTORS[0]: connect_loc,
    }
    mock_page = make_mock_outreach_page(locator_map=locator_map)

    res = asyncio.run(
        sender.send_outreach(
            recruiter_name="Amit Verma",
            company="PNC Bank GCC",
            page=mock_page
        )
    )

    assert res["status"] == "ALREADY_CONNECTED"
    assert res["sent"] is False
    connect_loc.click.assert_not_called()

    leads = crm.list_active_leads()
    assert len(leads) == 1
    assert leads[0]["status"] == "CONNECTED"
    assert "Already 1st-degree connection" in leads[0]["notes"]


def test_outreach_pending_invitation_detection(tmp_path):
    """Verify that a pending invitation is detected without re-sending."""
    json_path = str(tmp_path / "outreach.json")
    crm = RecruiterOutreachCRM(json_path=json_path)
    sender = LinkedInOutreachSender(crm=crm)

    pending_loc = make_mock_outreach_locator(count=1, visible=True, text="Pending")
    connect_loc = make_mock_outreach_locator(count=1, visible=True)

    locator_map = {
        PENDING_SELECTORS[0]: pending_loc,
        DIRECT_CONNECT_SELECTORS[0]: connect_loc,
    }
    mock_page = make_mock_outreach_page(locator_map=locator_map)

    res = asyncio.run(
        sender.send_outreach(
            recruiter_name="Kavita Rao",
            company="Microsoft",
            page=mock_page
        )
    )

    assert res["status"] == "PENDING"
    assert res["sent"] is False
    connect_loc.click.assert_not_called()

    leads = crm.list_active_leads()
    assert len(leads) == 1
    assert "pending" in leads[0]["notes"].lower()


def test_outreach_strict_character_limit_compliance(tmp_path):
    """
    CRITICAL CONSTRAINT: Connection note must never exceed 300 characters.
    Even with an adversarial 500-character custom note, it is clamped to <= 300.
    """
    crm = RecruiterOutreachCRM(json_path=str(tmp_path / "outreach.json"))
    sender = LinkedInOutreachSender(crm=crm)

    connect_loc = make_mock_outreach_locator(count=1, visible=True)
    add_note_loc = make_mock_outreach_locator(count=1, visible=True)
    textarea_loc = make_mock_outreach_locator(count=1, visible=True)
    dismiss_loc = make_mock_outreach_locator(count=1, visible=True)

    locator_map = {
        DIRECT_CONNECT_SELECTORS[0]: connect_loc,
        ADD_NOTE_SELECTORS[0]: add_note_loc,
        TEXTAREA_SELECTORS[0]: textarea_loc,
        DISMISS_SELECTORS[0]: dismiss_loc,
    }
    mock_page = make_mock_outreach_page(locator_map=locator_map)

    excessive_note = "A" * 500
    res = asyncio.run(
        sender.send_outreach(
            recruiter_name="Ananya Sen",
            company="Amazon",
            note=excessive_note,
            dry_run=True,
            page=mock_page
        )
    )

    assert res["status"] == "DRY_RUN_VERIFIED"
    assert res["note_length"] == 300
    assert len(res["note"]) == 300


def test_outreach_dry_run_simulation_mode(tmp_path):
    """
    Verify dry-run simulation:
    1. Connect and Add a note are clicked
    2. Note is typed into textarea
    3. Dismiss / Cancel is clicked
    4. Send button is NEVER clicked
    5. Lead is logged with [DRY RUN] note
    """
    crm = RecruiterOutreachCRM(json_path=str(tmp_path / "outreach.json"))
    sender = LinkedInOutreachSender(crm=crm)

    connect_loc = make_mock_outreach_locator(count=1, visible=True)
    add_note_loc = make_mock_outreach_locator(count=1, visible=True)
    textarea_loc = make_mock_outreach_locator(count=1, visible=True)
    dismiss_loc = make_mock_outreach_locator(count=1, visible=True)
    send_loc = make_mock_outreach_locator(count=1, visible=True)

    locator_map = {
        DIRECT_CONNECT_SELECTORS[0]: connect_loc,
        ADD_NOTE_SELECTORS[0]: add_note_loc,
        TEXTAREA_SELECTORS[0]: textarea_loc,
        DISMISS_SELECTORS[0]: dismiss_loc,
        SEND_SELECTORS[0]: send_loc,
    }
    mock_page = make_mock_outreach_page(locator_map=locator_map)

    res = asyncio.run(
        sender.send_outreach(
            recruiter_name="Vikram Malhotra",
            company="Barclays",
            dry_run=True,
            page=mock_page
        )
    )

    assert res["status"] == "DRY_RUN_VERIFIED"
    assert res["sent"] is False
    assert res["dry_run"] is True
    connect_loc.click.assert_awaited_once()
    add_note_loc.click.assert_awaited_once()
    dismiss_loc.click.assert_awaited_once()
    send_loc.click.assert_not_called()

    leads = crm.list_active_leads()
    assert len(leads) == 1
    assert "[DRY RUN]" in leads[0]["notes"]


def test_outreach_rate_limit_http_429_graceful_handling(tmp_path):
    """Verify HTTP 429 rate limit is gracefully captured without uncaught exception."""
    crm = RecruiterOutreachCRM(json_path=str(tmp_path / "outreach.json"))
    sender = LinkedInOutreachSender(crm=crm)

    mock_page = make_mock_outreach_page(body_content="HTTP 429 Too Many Requests. IP rate limited.")

    res = asyncio.run(
        sender.send_outreach(
            recruiter_name="Sneha Kulkarni",
            company="HSBC Technology",
            page=mock_page
        )
    )

    assert res["status"] == "RATE_LIMITED"
    assert res["sent"] is False
    assert "429" in res["detail"]


def test_outreach_checkpoint_and_authwall_detection(tmp_path):
    """Verify security checkpoint and authwall redirects are safely detected."""
    crm = RecruiterOutreachCRM(json_path=str(tmp_path / "outreach.json"))
    sender = LinkedInOutreachSender(crm=crm)

    # Checkpoint
    checkpoint_page = make_mock_outreach_page(url="https://www.linkedin.com/checkpoint/challenge/12345")
    res1 = asyncio.run(sender.send_outreach("Arjun Nair", "Lloyds", page=checkpoint_page))
    assert res1["status"] == "CHECKPOINT"
    assert res1["sent"] is False

    # Authwall
    authwall_page = make_mock_outreach_page(url="https://www.linkedin.com/authwall?trk=bf")
    res2 = asyncio.run(sender.send_outreach("Arjun Nair", "Lloyds", page=authwall_page))
    assert res2["status"] == "AUTH_WALL"
    assert res2["sent"] is False


def test_outreach_daily_safety_throttling(tmp_path):
    """
    Verify safety throttle halts outreach when daily connection limit (e.g. 8) is reached.
    """
    json_path = str(tmp_path / "outreach.json")
    crm = RecruiterOutreachCRM(json_path=json_path)

    # Pre-populate CRM with 8 real outreaches today
    today_prefix = datetime.now().strftime("%Y-%m-%d %H:%M")
    for i in range(8):
        crm.log_outreach(
            recruiter_name=f"Recruiter {i}",
            company=f"Company {i}",
            platform="LinkedIn",
            status="CONTACTED",
            notes="Real invitation sent",
            date_str=today_prefix
        )

    sender = LinkedInOutreachSender(crm=crm, max_daily_limit=8)
    mock_page = make_mock_outreach_page()

    res = asyncio.run(
        sender.send_outreach(
            recruiter_name="Neha Kapoor",
            company="Salesforce",
            dry_run=False,
            page=mock_page
        )
    )

    assert res["status"] == "DAILY_LIMIT_REACHED"
    assert res["sent"] is False
    assert "Safety throttle" in res["detail"]
    mock_page.goto.assert_not_called()


def test_cli_goal_5calls_send_outreach_flag(monkeypatch):
    """Verify CLI --send-outreach dispatches cleanly and returns 0."""
    with patch("naukari_pro.modules.linkedin_outreach.LinkedInOutreachSender.run_outreach") as mock_run:
        mock_run.return_value = {
            "status": "DRY_RUN_VERIFIED",
            "recruiter_name": "Kavita Rao",
            "company": "Microsoft",
            "sent": False,
            "note": "Hi Kavita, ...",
            "note_length": 250,
            "detail": "Dry-run verified"
        }
        ret = handle_goal_5calls(
            send_outreach=True,
            company="Microsoft",
            recruiter="Kavita Rao",
            dry_run=True
        )
        assert ret == 0
        mock_run.assert_called_once()


def test_cli_goal_5calls_batch_send_outreach(monkeypatch):
    """Verify CLI batch outreach dispatches cleanly."""
    with patch("naukari_pro.modules.linkedin_outreach.LinkedInOutreachSender.run_batch") as mock_batch:
        mock_batch.return_value = {
            "status": "COMPLETED",
            "total_processed": 3,
            "sent_count": 0,
            "dry_run_count": 3,
            "results": [
                {"company": "Microsoft", "recruiter_name": "Kavita", "status": "DRY_RUN_VERIFIED", "note": "Note 1", "note_length": 150},
                {"company": "Walmart", "recruiter_name": "Priya", "status": "DRY_RUN_VERIFIED", "note": "Note 2", "note_length": 160},
                {"company": "Amazon", "recruiter_name": "Ananya", "status": "DRY_RUN_VERIFIED", "note": "Note 3", "note_length": 170},
            ]
        }
        ret = handle_goal_5calls(
            send_outreach=True,
            limit=3,
            dry_run=True
        )
        assert ret == 0
        mock_batch.assert_called_once()


def test_cli_interactive_handler_bot_option(monkeypatch):
    """Verify interactive menu handles choice 'B' then exits cleanly on '0'."""
    from naukari_pro.cli import handle_goal_5calls_interactive

    inputs = iter([
        "B",        # Choose Bot
        "1",        # Batch mode
        "3",        # Limit 3
        "Y",        # Dry-run Y
        "n",        # Headed N
        "",         # Press enter to continue
        "0"         # Exit
    ])
    monkeypatch.setattr("builtins.input", lambda prompt="": next(inputs))

    with patch("naukari_pro.modules.linkedin_outreach.LinkedInOutreachSender.run_batch") as mock_batch:
        mock_batch.return_value = {
            "status": "COMPLETED",
            "total_processed": 3,
            "sent_count": 0,
            "dry_run_count": 3,
            "results": []
        }
        ret = handle_goal_5calls_interactive()
        assert ret == 0
        mock_batch.assert_called_once()


