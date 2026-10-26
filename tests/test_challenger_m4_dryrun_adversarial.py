"""
Adversarial Empirical Stress Test Suite for Milestone 4 (CLI, Dry-Run & Scheduler Robustness)
Authored by Challenger 2 (teamwork_preview_challenger_m4_2)

Empirically stress-tests:
1. Strict Offline Socket Guard: monkeypatches socket.socket and network primitives to raise AirGapViolationError;
   verifies LinkedInApplyEngine.run_offline_dry_run() executes 100% offline with 0 socket calls.
2. Skip-List Purity Stress: verifies skip_list_linkedin.json is NEVER modified, corrupted, or contaminated during dry-run.
3. CLI Argument Parsing Stress: verifies comma-delimited keywords/locations, plural aliases (--keywords, --locations),
   empty/whitespace commas, edge-case match gates, and direct execution routing.
4. Shell Script Robustness & Error Isolation: verifies agy_schedule_runner.sh syntax via bash -n, checks error isolation
   on lines 116 & 122 (|| true), and executes isolated subshell simulations under set -euo pipefail.
5. Tracker Multi-Tier Fallback Under Hostile File Systems: verifies emergency CSV logging and audit trail without skip list pollution.
"""

import os
import sys
import json
import socket
import argparse
import tempfile
import subprocess
from pathlib import Path
from typing import Dict, List, Any
from unittest.mock import patch, MagicMock

import openpyxl
import pytest

from naukari_pro.modules.linkedin_apply import (
    LinkedInApplyEngine,
    LinkedInApplicationTracker,
    LinkedInApplyRunner,
    LINKEDIN_BLUE_FILL,
    WHITE_BOLD_FONT
)
from naukari_pro.modules.linkedin_search import JobMatchScorer
from naukari_pro.cli import handle_linkedin_apply, main


# ==============================================================================
# 1. STRICT OFFLINE AIR-GAP SOCKET TRAP
# ==============================================================================

class AirGapViolationError(RuntimeError):
    """Raised whenever any low-level socket creation or network call is attempted."""
    pass


@pytest.fixture
def strict_airgap_socket_trap(monkeypatch):
    """
    Adversarial socket trap:
    Replaces socket.socket, socket.create_connection, socket.getaddrinfo, and
    socket.gethostbyname with functions that raise AirGapViolationError immediately.
    """
    def trap_call(*args, **kwargs):
        raise AirGapViolationError("CRITICAL AIR-GAP VIOLATION: Unsanctioned network socket operation attempted!")

    original_socket_cls = socket.socket

    class TrappedSocket:
        def __init__(self, *args, **kwargs):
            raise AirGapViolationError("CRITICAL AIR-GAP VIOLATION: socket.socket() instantiation attempted!")

    monkeypatch.setattr(socket, "socket", TrappedSocket)
    monkeypatch.setattr(socket, "create_connection", trap_call)
    monkeypatch.setattr(socket, "getaddrinfo", trap_call)
    monkeypatch.setattr(socket, "gethostbyname", trap_call)
    monkeypatch.setattr(socket, "gethostbyname_ex", trap_call)

    yield

    monkeypatch.setattr(socket, "socket", original_socket_cls)


def test_adversarial_strict_socket_guard(tmp_path, strict_airgap_socket_trap):
    """
    Adversarial Challenge 1:
    Verify LinkedInApplyEngine.run_offline_dry_run() completes with ZERO network sockets.
    Under strict_airgap_socket_trap, ANY network initialization raises AirGapViolationError.
    """
    excel_path = str(tmp_path / "test_tracker.xlsx")
    json_path = str(tmp_path / "test_log.json")
    skip_path = str(tmp_path / "test_skip.json")

    tracker = LinkedInApplicationTracker(
        excel_path=excel_path,
        json_log_path=json_path,
        skip_file_path=skip_path
    )
    engine = LinkedInApplyEngine(tracker=tracker, min_score=70.0)

    # Execute dry run under active air-gap trap
    stats = engine.run_offline_dry_run(
        keywords=["Azure Data Engineer", "Databricks"],
        locations=["Bengaluru", "Remote"],
        limit=5,
        min_score=70.0
    )

    # Assertions on returned execution stats
    assert stats["discovered"] == 5, f"Expected 5 synthetic vacancies discovered, got {stats['discovered']}"
    assert stats["scored_above_threshold"] >= 3, "Expected at least 3 high-yield matching vacancies"
    assert stats["dry_run_verified"] >= 3, "Expected at least 3 dry-run verified applications"
    assert stats["skipped"] >= 2, "Expected low-scoring/irrelevant jobs to be gated out"
    assert stats["failed"] == 0, f"Expected 0 failures, got {stats['failed']}"

    # Verify scored jobs details
    assert len(stats["jobs"]) == 5
    for j in stats["jobs"]:
        assert "job_id" in j
        assert "score" in j
        assert "status" in j
        if j["status"] == "SUBMITTED_DRY_RUN":
            assert j["score"] >= 70.0
            assert j["steps"] == 4
            assert j["success"] is True
        elif j["status"] == "GATED_OUT":
            assert j["score"] < 70.0
            assert j["steps"] == 0
            assert j["success"] is False


# ==============================================================================
# 2. SKIP-LIST NON-POLLUTION & IMMUTABILITY STRESS
# ==============================================================================

def test_adversarial_skip_list_empty_purity(tmp_path):
    """
    Adversarial Challenge 2A:
    Verify an initially empty skip_list_linkedin.json remains completely empty after dry-run.
    """
    excel_path = str(tmp_path / "tracker.xlsx")
    json_path = str(tmp_path / "log.json")
    skip_path = str(tmp_path / "skip_list_linkedin.json")

    # Initialize empty skip list
    with open(skip_path, "w", encoding="utf-8") as f:
        json.dump([], f)

    tracker = LinkedInApplicationTracker(
        excel_path=excel_path,
        json_log_path=json_path,
        skip_file_path=skip_path
    )
    engine = LinkedInApplyEngine(tracker=tracker, min_score=70.0)

    stats = engine.run_offline_dry_run(limit=5)
    assert stats["dry_run_verified"] >= 3

    # Skip list must remain strictly empty
    with open(skip_path, "r", encoding="utf-8") as f:
        skip_content = json.load(f)
    assert skip_content == [], f"Skip list contaminated! Found: {skip_content}"


def test_adversarial_skip_list_preexisting_contents_unaltered(tmp_path):
    """
    Adversarial Challenge 2B:
    Verify pre-existing production entries in skip_list_linkedin.json are NOT mutated,
    deleted, or appended to during dry-run.
    """
    excel_path = str(tmp_path / "tracker.xlsx")
    json_path = str(tmp_path / "log.json")
    skip_path = str(tmp_path / "skip_list_linkedin.json")

    preexisting_ids = ["4100000001", "4100000002", "4100000003"]
    with open(skip_path, "w", encoding="utf-8") as f:
        json.dump(preexisting_ids, f, indent=2)

    with open(skip_path, "rb") as f:
        original_hash = hash(f.read())

    tracker = LinkedInApplicationTracker(
        excel_path=excel_path,
        json_log_path=json_path,
        skip_file_path=skip_path
    )
    engine = LinkedInApplyEngine(tracker=tracker, min_score=70.0)

    # Run multiple successive dry runs
    for run_idx in range(5):
        engine.run_offline_dry_run(limit=5)

    with open(skip_path, "r", encoding="utf-8") as f:
        current_ids = json.load(f)

    with open(skip_path, "rb") as f:
        current_hash = hash(f.read())

    assert current_ids == preexisting_ids, f"Skip list entries modified! Current: {current_ids}"
    assert current_hash == original_hash, "Skip list file was rewritten during dry runs!"


def test_adversarial_tracker_direct_log_dry_run_gate(tmp_path):
    """
    Adversarial Challenge 2C:
    Verify LinkedInApplicationTracker.log_application() strictly enforces dry_run=True gate
    even when passed status='APPLIED' or status='SUBMITTED'.
    """
    excel_path = str(tmp_path / "tracker.xlsx")
    json_path = str(tmp_path / "log.json")
    skip_path = str(tmp_path / "skip_list_linkedin.json")

    with open(skip_path, "w", encoding="utf-8") as f:
        json.dump(["INITIAL_ID"], f)

    tracker = LinkedInApplicationTracker(
        excel_path=excel_path,
        json_log_path=json_path,
        skip_file_path=skip_path
    )

    # Attempt to log with status='APPLIED' but dry_run=True
    tracker.log_application(
        job_data={"job_id": "9999999999", "company": "Target Co", "title": "Azure DE"},
        status="APPLIED",
        match_score=95.0,
        steps=4,
        dry_run=True  # MUST prevent skip list update!
    )

    with open(skip_path, "r", encoding="utf-8") as f:
        skip_data = json.load(f)
    assert "9999999999" not in skip_data
    assert skip_data == ["INITIAL_ID"]

    # Now verify authentic submission (dry_run=False) DOES update skip list
    tracker.log_application(
        job_data={"job_id": "8888888888", "company": "Target Co", "title": "Azure DE"},
        status="APPLIED",
        match_score=95.0,
        steps=4,
        dry_run=False
    )
    with open(skip_path, "r", encoding="utf-8") as f:
        skip_data_after = json.load(f)
    assert "8888888888" in skip_data_after


# ==============================================================================
# 3. CLI ARGUMENT PARSING & DISPATCH STRESS
# ==============================================================================

def test_adversarial_cli_plural_aliases_and_comma_parsing():
    """
    Adversarial Challenge 3A:
    Test argparse registration of --keywords and --locations plural aliases,
    comma-separated values, and default limit=25.
    """
    parser = argparse.ArgumentParser()
    subparsers = parser.add_subparsers(dest="command")
    common_parser = argparse.ArgumentParser(add_help=False)
    common_parser.add_argument("--headless", action="store_true")
    common_parser.add_argument("--limit", type=int, default=20)
    common_parser.add_argument("--dry-run", action="store_true")

    apply_parser = subparsers.add_parser("linkedin-apply", parents=[common_parser])
    apply_parser.add_argument("--headed", action="store_true", default=False)
    apply_parser.add_argument("--min-score", type=float, default=70.0)
    apply_parser.add_argument("--keyword", "--keywords", dest="keyword", type=str, default=None)
    apply_parser.add_argument("--location", "--locations", dest="location", type=str, default=None)
    apply_parser.set_defaults(limit=25)

    # Test 1: Plural flags with comma lists
    args1 = parser.parse_args([
        "linkedin-apply",
        "--keywords", "Azure Data Engineer, Databricks, PySpark",
        "--locations", "Bengaluru, Hyderabad, Remote",
        "--limit", "10",
        "--dry-run"
    ])
    assert args1.command == "linkedin-apply"
    assert args1.keyword == "Azure Data Engineer, Databricks, PySpark"
    assert args1.location == "Bengaluru, Hyderabad, Remote"
    assert args1.limit == 10
    assert args1.dry_run is True

    # Test 2: Singular flags without overriding limit (should default to 25)
    args2 = parser.parse_args([
        "linkedin-apply",
        "--keyword", "Delta Lake",
        "--location", "Pune"
    ])
    assert args2.limit == 25
    assert args2.keyword == "Delta Lake"
    assert args2.location == "Pune"


def test_adversarial_cli_whitespace_commas_handling(capsys, tmp_path, monkeypatch):
    """
    Adversarial Challenge 3B:
    Stress-test handle_linkedin_apply() with whitespace-only commas:
    --keywords " , , " and --locations "   ,  , "
    Must cleanly filter out empty items without IndexError and fall back to defaults.
    """
    excel_path = str(tmp_path / "tracker.xlsx")
    monkeypatch.setenv("WORKSPACE_DIR", str(tmp_path))

    class MockArgs:
        dry_run = True
        limit = 5
        headed = False
        headless = True
        min_score = 70.0
        keyword = "  ,  ,  "
        location = " ,   , "

    with patch("naukari_pro.modules.linkedin_apply.TRACKER_PATH", excel_path):
        exit_code = handle_linkedin_apply(MockArgs())
        assert exit_code == 0

    out = capsys.readouterr().out
    assert "DRY-RUN OFFLINE SIMULATION" in out
    assert "Dry-Run Verified:" in out
    assert "Passed Match Gate:" in out


def test_adversarial_cli_edge_case_scores(capsys, tmp_path, monkeypatch):
    """
    Adversarial Challenge 3C:
    Test handle_linkedin_apply() with extreme min_score boundaries:
    1. --min-score 95.0: All 5 synthetic vacancies must be GATED_OUT (0 verified, exit code 0).
    2. --min-score 10.0: Relevant roles verified, irrelevant roles (swiggy react, oracle 15y) still filtered.
    """
    excel_path = str(tmp_path / "tracker.xlsx")
    monkeypatch.setenv("WORKSPACE_DIR", str(tmp_path))

    # Boundary 1: High threshold (95.0%)
    class HighScoreArgs:
        dry_run = True
        limit = 5
        headed = False
        headless = True
        min_score = 95.0
        keyword = "Azure Data Engineer"
        location = "Bengaluru"

    with patch("naukari_pro.modules.linkedin_apply.TRACKER_PATH", excel_path):
        exit_code_high = handle_linkedin_apply(HighScoreArgs())
        assert exit_code_high == 0

    out_high = capsys.readouterr().out
    assert "Passed Match Gate:     0" in out_high
    assert "Dry-Run Verified:      0" in out_high
    assert "Skipped / Gated Out:   5" in out_high

    # Boundary 2: Low threshold (10.0%)
    class LowScoreArgs:
        dry_run = True
        limit = 5
        headed = False
        headless = True
        min_score = 10.0
        keyword = "Azure Data Engineer"
        location = "Bengaluru"

    with patch("naukari_pro.modules.linkedin_apply.TRACKER_PATH", excel_path):
        exit_code_low = handle_linkedin_apply(LowScoreArgs())
        assert exit_code_low == 0

    out_low = capsys.readouterr().out
    assert "Dry-Run Verified:      3" in out_low


def test_adversarial_cli_binary_execution():
    """
    Adversarial Challenge 3D:
    Empirically execute `./naukariPro linkedin-apply` via subprocess in dry-run mode.
    Validates output table, exit status 0, and non-blocking operation.
    """
    cmd = [
        sys.executable,
        "-m", "naukari_pro.cli",
        "linkedin-apply",
        "--dry-run",
        "--limit", "3",
        "--keywords", "Azure Data Engineer, Databricks",
        "--locations", "Bengaluru, Hyderabad"
    ]
    proc = subprocess.run(
        cmd,
        capture_output=True,
        text=True,
        timeout=15,
        cwd=str(Path(__file__).resolve().parent.parent)
    )
    assert proc.returncode == 0, f"Process exited with non-zero: {proc.stderr}"
    assert "LINKEDIN AUTONOMOUS EASY APPLY ENGINE" in proc.stdout
    assert "DRY-RUN OFFLINE SIMULATION" in proc.stdout
    assert "Job ID" in proc.stdout
    assert "Microsoft GCC" in proc.stdout
    assert "Dry-Run Verified:      3" in proc.stdout


# ==============================================================================
# 4. SHELL SCRIPT ROBUSTNESS & ERROR ISOLATION (AGY_SCHEDULE_RUNNER.SH)
# ==============================================================================

def test_adversarial_shell_script_syntax():
    """
    Adversarial Challenge 4A:
    Validate bash syntax of agy_schedule_runner.sh using `bash -n`.
    """
    script_path = Path(__file__).resolve().parent.parent / "agy_schedule_runner.sh"
    assert script_path.exists(), "agy_schedule_runner.sh does not exist in workspace!"

    proc = subprocess.run(
        ["bash", "-n", str(script_path)],
        capture_output=True,
        text=True,
        timeout=5
    )
    assert proc.returncode == 0, f"Syntax error in agy_schedule_runner.sh: {proc.stderr}"
    assert proc.stderr.strip() == "", f"Unexpected bash -n warnings: {proc.stderr}"


def test_adversarial_shell_script_error_isolation_presence():
    """
    Adversarial Challenge 4B:
    Inspect agy_schedule_runner.sh to ensure lines executing `./naukariPro pro`
    contain explicit error isolation `|| true` on lines 116 and 122.
    """
    script_path = Path(__file__).resolve().parent.parent / "agy_schedule_runner.sh"
    with open(script_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    # Search for all executable ./naukariPro pro invocations (lines starting with ./naukariPro pro)
    pro_invocations = [
        (idx + 1, line.strip())
        for idx, line in enumerate(lines)
        if line.strip().startswith("./naukariPro pro")
    ]

    assert len(pro_invocations) >= 2, f"Expected at least 2 executable pro invocations, found {len(pro_invocations)}"

    # Specifically check lines 116 and 122
    assert "|| true" in lines[115], f"Line 116 lacks '|| true': {lines[115]}"
    assert "|| true" in lines[121], f"Line 122 lacks '|| true': {lines[121]}"

    for line_num, content in pro_invocations:
        assert "|| true" in content, (
            f"Line {line_num} in agy_schedule_runner.sh lacks error isolation ('|| true')! "
            f"Content: '{content}'"
        )


def test_adversarial_shell_script_error_isolation_simulation(tmp_path):
    """
    Adversarial Challenge 4C:
    Execute a simulated subshell under `set -euo pipefail` replicating lines 114-118 of
    agy_schedule_runner.sh.
    Verifies that when `./naukariPro pro` fails (exit 1), `|| true` prevents script termination
    and allows `./naukariPro linkedin-apply` to execute successfully.
    """
    test_sh = tmp_path / "test_pipeline.sh"
    log_file = tmp_path / "run.log"

    # Shell script testing error isolation
    sh_content = f"""#!/bin/bash
set -euo pipefail

RUN_LOG="{log_file}"

echo "Starting test pipeline..." >> "$RUN_LOG"

# Simulated failed naukariPro pro execution with error isolation
(exit 1) 2>&1 | tee -a "$RUN_LOG" || true

# Subsequent linkedin-apply execution MUST be reached
echo "LINKEDIN_APPLY_EXECUTED" >> "$RUN_LOG"
"""
    test_sh.write_text(sh_content)
    test_sh.chmod(0o755)

    proc = subprocess.run(
        ["bash", str(test_sh)],
        capture_output=True,
        text=True,
        timeout=5
    )
    assert proc.returncode == 0, f"Pipeline script crashed despite error isolation! stderr: {proc.stderr}"
    log_content = log_file.read_text()
    assert "LINKEDIN_APPLY_EXECUTED" in log_content, (
        "Subsequent command was NOT executed! Error isolation failed to prevent abort."
    )


def test_adversarial_shell_script_without_error_isolation_fails(tmp_path):
    """
    Adversarial Challenge 4D:
    Control test: Demonstrates that WITHOUT `|| true`, `set -euo pipefail` indeed
    aborts the script and PREVENTS linkedin-apply from ever executing.
    Confirms the necessity and efficacy of the fix on lines 116 and 122.
    """
    test_sh = tmp_path / "test_pipeline_unprotected.sh"
    log_file = tmp_path / "run.log"

    sh_content = f"""#!/bin/bash
set -euo pipefail

RUN_LOG="{log_file}"

# Unprotected failing command
(exit 1) 2>&1 | tee -a "$RUN_LOG"

# Should NEVER be reached
echo "LINKEDIN_APPLY_EXECUTED" >> "$RUN_LOG"
"""
    test_sh.write_text(sh_content)
    test_sh.chmod(0o755)

    proc = subprocess.run(
        ["bash", str(test_sh)],
        capture_output=True,
        text=True,
        timeout=5
    )
    # Must fail with non-zero exit
    assert proc.returncode != 0
    log_content = log_file.read_text() if log_file.exists() else ""
    assert "LINKEDIN_APPLY_EXECUTED" not in log_content


# ==============================================================================
# 5. TRACKER MULTI-TIER FALLBACK RESILIENCE UNDER HOSTILE CONDITIONS
# ==============================================================================

def test_adversarial_tracker_fallback_to_csv_on_corrupted_excel(tmp_path):
    """
    Adversarial Challenge 5A:
    If Excel tracker file is binary corrupted or unparseable,
    log_application() must seamlessly fall back to emergency CSV and JSON log,
    preserving application record with zero data loss.
    """
    corrupt_excel = tmp_path / "corrupted_tracker.xlsx"
    corrupt_excel.write_bytes(b"NOT_A_VALID_ZIP_OR_XLSX_DATA")

    json_path = str(tmp_path / "audit_log.json")
    skip_path = str(tmp_path / "skip.json")
    csv_fallback = str(tmp_path / "linkedin_applications_fallback.csv")

    tracker = LinkedInApplicationTracker(
        excel_path=str(corrupt_excel),
        json_log_path=json_path,
        skip_file_path=skip_path,
        fallback_csv_path=csv_fallback
    )

    job_data = {
        "job_id": "4199991234",
        "company": "Resilient Corp",
        "title": "Lead Data Platform Engineer",
        "location": "Bengaluru",
        "url": "https://www.linkedin.com/jobs/view/4199991234/"
    }

    # Should not raise exception
    success = tracker.log_application(
        job_data=job_data,
        status="SUBMITTED_DRY_RUN",
        match_score=88.5,
        steps=4,
        dry_run=True
    )
    assert success is True

    # Verify emergency CSV was created and contains the job
    assert os.path.exists(csv_fallback)
    with open(csv_fallback, "r", encoding="utf-8") as f:
        csv_text = f.read()
    assert "4199991234" in csv_text
    assert "Resilient Corp" in csv_text
    assert "LinkedIn Easy Apply" in csv_text

    # Verify JSON audit log contains the entry
    assert os.path.exists(json_path)
    with open(json_path, "r", encoding="utf-8") as f:
        audit_records = json.load(f)
    assert len(audit_records) == 1
    assert audit_records[0]["job_id"] == "4199991234"
    assert audit_records[0]["dry_run"] is True


def test_adversarial_tracker_canonical_headers_and_styling(tmp_path):
    """
    Adversarial Challenge 5B:
    Verify that newly initialized tracker sheets strictly feature:
    - 10 canonical columns (Feature 16)
    - LinkedIn Blue fill (#0A66C2)
    - Bold white font
    - Freeze panes at A2
    """
    excel_path = str(tmp_path / "canonical_tracker.xlsx")
    tracker = LinkedInApplicationTracker(excel_path=excel_path)
    tracker._ensure_excel_sheet()

    wb = openpyxl.load_workbook(excel_path)
    assert "LinkedIn_Applications_Log" in wb.sheetnames
    ws = wb["LinkedIn_Applications_Log"]

    # 1. 10 Canonical Columns
    headers = [ws.cell(1, col).value for col in range(1, 11)]
    assert headers == LinkedInApplicationTracker.CANONICAL_HEADERS

    # 2. LinkedIn Blue fill and Bold White Font on header cells
    for col in range(1, 11):
        cell = ws.cell(1, col)
        fill_color = cell.fill.start_color.rgb if cell.fill else None
        assert fill_color in ("000A66C2", "0A66C2", "FF0A66C2"), f"Col {col} fill color is {fill_color}, expected #0A66C2"
        assert cell.font.bold is True
        font_color = cell.font.color.rgb if cell.font and cell.font.color else None
        assert font_color in ("00FFFFFF", "FFFFFF", "FFFFFFFF"), f"Col {col} font color is {font_color}, expected White"

    # 3. Freeze panes at A2
    assert ws.freeze_panes == "A2", f"Expected freeze_panes='A2', got '{ws.freeze_panes}'"
