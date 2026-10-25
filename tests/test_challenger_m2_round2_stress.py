"""
Adversarial Stress Test Suite for Milestone 2 Round 2
Challenger: teamwork_preview_challenger_m2_round2_2
Target: Multi-page search harvesting, NativeChromeRunner loop execution, URL generation invariants, and Excel schema preservation.
"""

import os
import json
import shutil
import urllib.parse
from unittest.mock import MagicMock, patch
import pytest
import openpyxl

from naukari_pro.modules.auto_apply import (
    DEFAULT_KEYWORDS,
    CANONICAL_LOCATIONS,
    EXPANDED_M2_KEYWORDS,
    EXPERIENCE_PARAM,
    build_search_url,
    log_to_excel,
    make_job_id,
    NativeChromeRunner,
    PlaywrightRunner,
)
from naukari_pro.config import TRACKER_PATH


# ==============================================================================
# ADVERSARIAL SUITE 1: Search URL Invariants & Combinatorics
# ==============================================================================

class TestAdversarialUrlConstruction:
    """Stress tests and boundary condition validation for search query generation."""

    def test_exhaustive_permutation_url_invariants(self):
        """Verify all 231 permutations (11 keywords * 7 locations * 3 pages) satisfy strict URL invariants."""
        assert len(EXPANDED_M2_KEYWORDS) == 11
        assert len(CANONICAL_LOCATIONS) == 7

        all_urls = []
        for loc in CANONICAL_LOCATIONS:
            for kw in EXPANDED_M2_KEYWORDS:
                for page in range(1, 4):
                    url = build_search_url(kw, loc, page)
                    all_urls.append(url)

                    # Invariant 1: Scheme and Host
                    parsed = urllib.parse.urlparse(url)
                    assert parsed.scheme == "https"
                    assert parsed.netloc == "www.naukri.com"

                    # Invariant 2: Path structure
                    expected_slug = "-".join(kw.lower().split())
                    expected_suffix = f"-{page}" if page > 1 else ""
                    expected_path = f"/{expected_slug}-jobs-in-{loc}{expected_suffix}"
                    assert parsed.path == expected_path

                    # Invariant 3: Query parameters
                    params = urllib.parse.parse_qs(parsed.query, keep_blank_values=True)
                    assert params.get("experience") == [str(EXPERIENCE_PARAM)]
                    assert params.get("jobAge") == ["1"]
                    assert params.get("freshness") == ["1"]
                    assert params.get("sortBy") == ["date"]
                    assert params.get("sortType") == ["1"]

                    if page == 1:
                        assert "pageNo" not in params
                    else:
                        assert params.get("pageNo") == [str(page)]

        # Verify uniqueness: 11 * 7 * 3 = 231 distinct URLs
        assert len(all_urls) == 231
        assert len(set(all_urls)) == 231

    @pytest.mark.parametrize("kw,loc,page,expected_slug", [
        ("  Azure   Data    Engineer  ", "hyderabad", 1, "azure-data-engineer"),
        ("ETL/ELT Pipeline", "bengaluru", 2, "etl/elt-pipeline"),
        ("PySpark & Databricks", "pune", 3, "pyspark-&-databricks"),
        ("Kafka", "  MUMBAI  ", 1, "kafka"),
    ])
    def test_url_whitespace_and_casing_resilience(self, kw, loc, page, expected_slug):
        """Verify build_search_url normalizes arbitrary spacing, casing, and edge inputs."""
        url = build_search_url(kw, loc, page)
        parsed = urllib.parse.urlparse(url)
        clean_loc = loc.lower().strip()
        suffix = f"-{page}" if page > 1 else ""
        assert parsed.path == f"/{expected_slug}-jobs-in-{clean_loc}{suffix}"


# ==============================================================================
# ADVERSARIAL SUITE 2: NativeChromeRunner Loop Execution & Indentation Proof
# ==============================================================================

class TestAdversarialNativeChromeRunnerLoop:
    """Empirically stress-tests NativeChromeRunner.run multi-page harvesting and job processing."""

    def test_multipage_harvesting_processes_all_pages_without_dropping_jobs(self):
        """
        Adversarial test:
        Provide 2 jobs on Page 1, 2 jobs on Page 2, and 1 job on Page 3.
        Verify that ALL 5 jobs are evaluated and processed sequentially, proving
        that Page 2 and Page 3 fetches do NOT clobber or drop Page 1 jobs.
        """
        runner = NativeChromeRunner(skip_list=set(), dry_run=True)
        runner.run_multi_apply_recommended = MagicMock(return_value=0)

        # Track which pages were called
        visited_urls = []

        def mock_exec_js(url, script, wait_sec=4):
            visited_urls.append(url)
            parsed = urllib.parse.urlparse(url)
            qs = urllib.parse.parse_qs(parsed.query)
            page_no = qs.get("pageNo", ["1"])[0]

            if page_no == "1":
                return json.dumps([
                    {
                        "title": "Azure Data Engineer P1-A",
                        "company": "Enterprise 1A",
                        "sal": "16 LPA",
                        "href": "https://www.naukri.com/job-1a",
                        "desc": "Azure Data Engineer PySpark Databricks Delta Lake Kafka ADF",
                        "tags": "PySpark, Databricks"
                    },
                    {
                        "title": "Azure Data Engineer P1-B",
                        "company": "Enterprise 1B",
                        "sal": "18 LPA",
                        "href": "https://www.naukri.com/job-1b",
                        "desc": "Azure Data Engineer PySpark Databricks Delta Lake Kafka ADF",
                        "tags": "PySpark, ADF"
                    }
                ])
            elif page_no == "2":
                return json.dumps([
                    {
                        "title": "Azure Data Engineer P2-A",
                        "company": "Enterprise 2A",
                        "sal": "20 LPA",
                        "href": "https://www.naukri.com/job-2a",
                        "desc": "Azure Data Engineer PySpark Databricks Delta Lake Kafka ADF",
                        "tags": "Delta Lake, Kafka"
                    },
                    {
                        "title": "Azure Data Engineer P2-B",
                        "company": "Enterprise 2B",
                        "sal": "22 LPA",
                        "href": "https://www.naukri.com/job-2b",
                        "desc": "Azure Data Engineer PySpark Databricks Delta Lake Kafka ADF",
                        "tags": "Azure, Databricks"
                    }
                ])
            elif page_no == "3":
                return json.dumps([
                    {
                        "title": "Azure Data Engineer P3-A",
                        "company": "Enterprise 3A",
                        "sal": "24 LPA",
                        "href": "https://www.naukri.com/job-3a",
                        "desc": "Azure Data Engineer PySpark Databricks Delta Lake Kafka ADF",
                        "tags": "PySpark, Kafka"
                    }
                ])
            return json.dumps([])

        runner.exec_js_in_tab = mock_exec_js

        # Restrict to 1 keyword and 1 location to isolate multi-page behavior
        with patch("naukari_pro.modules.auto_apply.DEFAULT_KEYWORDS", ["Azure Data Engineer"]):
            with patch("naukari_pro.modules.auto_apply.LOCATIONS", ["hyderabad"]):
                total = runner.run(limit=10)

        # All 5 jobs across pages 1, 2, and 3 should be successfully processed
        assert total == 5, f"Expected 5 jobs processed across 3 pages, but got {total}!"
        assert len(visited_urls) == 3, f"Expected 3 page fetches, got {len(visited_urls)}"
        assert not any("pageNo" in visited_urls[0] for _ in [0])
        assert "pageNo=2" in visited_urls[1]
        assert "pageNo=3" in visited_urls[2]

    def test_runner_halts_immediately_when_limit_reached_mid_page(self):
        """Verify runner terminates immediately when application limit is reached on page 1."""
        runner = NativeChromeRunner(skip_list=set(), dry_run=True)
        runner.run_multi_apply_recommended = MagicMock(return_value=0)

        def mock_exec_js(url, script, wait_sec=4):
            return json.dumps([
                {
                    "title": f"Azure Data Engineer Job {i}",
                    "company": f"Enterprise {i}",
                    "sal": "18 LPA",
                    "href": f"https://www.naukri.com/job-{i}",
                    "desc": "Azure Data Engineer PySpark Databricks Delta Lake Kafka ADF",
                    "tags": "PySpark, Databricks"
                }
                for i in range(10)
            ])

        runner.exec_js_in_tab = mock_exec_js

        with patch("naukari_pro.modules.auto_apply.DEFAULT_KEYWORDS", ["Azure Data Engineer", "PySpark"]):
            with patch("naukari_pro.modules.auto_apply.LOCATIONS", ["hyderabad"]):
                total = runner.run(limit=3)

        assert total == 3

    def test_runner_recovers_gracefully_from_corrupt_json_payloads(self):
        """Verify that malformed or non-JSON returns from exec_js_in_tab do not crash the runner."""
        runner = NativeChromeRunner(skip_list=set(), dry_run=True)
        runner.run_multi_apply_recommended = MagicMock(return_value=0)

        # Return malformed JSON on page 1
        runner.exec_js_in_tab = MagicMock(return_value="INVALID_JSON_CONTENT{{{")

        with patch("naukari_pro.modules.auto_apply.DEFAULT_KEYWORDS", ["Azure Data Engineer"]):
            with patch("naukari_pro.modules.auto_apply.LOCATIONS", ["hyderabad"]):
                total = runner.run(limit=5)

        # Should handle exception and return 0 without raising
        assert total == 0


# ==============================================================================
# ADVERSARIAL SUITE 3: Excel Tracker 11-Column Schema Stress Testing
# ==============================================================================

class TestAdversarialExcelTrackerPreservation:
    """Stress tests on Excel schema invariants, corrupt input resistance, and multi-sheet integrity."""

    CANONICAL_COLUMNS = [
        "Application Date",
        "Company Name",
        "Job Title",
        "Match Score (%)",
        "Role Category",
        "Experience Required",
        "Location",
        "Application Type",
        "Status",
        "Job URL",
        "Recruiter Response"
    ]

    def test_100_diverse_writes_preserve_exact_11_columns_and_cell_invariants(self, tmp_path):
        """Stress-test with 100 heterogeneous rows containing emojis, special chars, quotes, commas, and edge scores."""
        test_tracker = str(tmp_path / "Adversarial_Tracker.xlsx")
        shutil.copyfile(TRACKER_PATH, test_tracker)

        orig_wb = openpyxl.load_workbook(test_tracker, data_only=True)
        base_rows = orig_wb["Applications_Log"].max_row

        test_cases = [
            # (title, company, status, match_score, app_type, salary, url)
            ("Senior Data Engineer, Cloud & BigData", "FinTech, Inc.", "Submitted (Pending)", 85, "1-Click Direct", "20 LPA", "https://naukri.com/job/1,2,3"),
            ("Data Engineer (Azure / Databricks)", "Global Tech Ltd & Co", "External Redirect", "90%", "External Portal", "25 LPA", "https://naukri.com/job/special?a=1&b=2"),
            ("Cloud Data Platform Engineer", "Startup 🚀", "Submitted (Pending)", 70.0, "1-Click Direct", "18 LPA", ""),
            ("PySpark Specialist", "Consulting Group", "Targeted Vacancy", None, "1-Click Apply", "15 LPA", "N/A"),
            ("Data Lakehouse Architect", "Retail Giant", "Submitted (Pending)", " 78 % ", "1-Click Direct", "30 LPA", "https://naukri.com/lakehouse"),
        ]

        # Append 100 entries cycling through edge cases
        for i in range(100):
            tc = test_cases[i % len(test_cases)]
            ok = log_to_excel(
                title=f"{tc[0]} #{i}",
                company=f"{tc[1]} #{i}",
                status=tc[2],
                match_score=tc[3],
                app_type=tc[4],
                salary=tc[5],
                url=tc[6],
                tracker_path=test_tracker
            )
            assert ok is True

        # Audit workbook
        audit_wb = openpyxl.load_workbook(test_tracker, data_only=True)
        ws = audit_wb["Applications_Log"]

        assert ws.max_row == base_rows + 100
        assert ws.max_column == 11, f"Schema drift detected! max_column is {ws.max_column}, expected 11"

        # Verify header is strictly preserved
        current_header = [ws.cell(1, c).value for c in range(1, 12)]
        assert current_header == self.CANONICAL_COLUMNS

        # Verify all newly written rows
        for r in range(base_rows + 1, ws.max_row + 1):
            row_vals = [ws.cell(r, c).value for c in range(1, 12)]
            assert len(row_vals) == 11
            assert all(v is not None for v in row_vals), f"Row {r} contains None cell: {row_vals}"
            assert str(row_vals[3]).endswith("%"), f"Score column not formatted with % at row {r}: {row_vals[3]}"
            assert row_vals[4] == "Azure Data Engineer"
            assert row_vals[5] == "2-5 Yrs"
            assert row_vals[6] == "Hyderabad / Bengaluru"
            assert row_vals[10] == "Awaiting Recruiter Review"

    def test_log_to_excel_creates_clean_canonical_file_from_scratch(self, tmp_path):
        """Verify log_to_excel can bootstrap an empty nonexistent file into an exact 11-column tracker."""
        new_tracker = str(tmp_path / "SubDir" / "Fresh_Tracker.xlsx")
        os.makedirs(os.path.dirname(new_tracker), exist_ok=True)

        ok = log_to_excel(
            title="Azure Data Engineer",
            company="Accenture",
            status="Submitted (Pending)",
            match_score=80,
            app_type="1-Click Direct",
            url="https://www.naukri.com/fresh-job",
            tracker_path=new_tracker
        )
        assert ok is True

        wb = openpyxl.load_workbook(new_tracker, data_only=True)
        assert "Applications_Log" in wb.sheetnames
        ws = wb["Applications_Log"]
        assert ws.max_row == 2
        assert ws.max_column == 11
        assert [ws.cell(1, c).value for c in range(1, 12)] == self.CANONICAL_COLUMNS
        assert ws.cell(2, 2).value == "Accenture"
        assert ws.cell(2, 3).value == "Azure Data Engineer"
        assert ws.cell(2, 4).value == "80%"
