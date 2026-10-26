"""
Adversarial Stress Test Suite: Milestone 4 Feature 16
Target: LinkedInApplicationTracker & LinkedInApplyRunner Facade
Governed by: PROJECT.md, ORIGINAL_REQUEST.md, and AGENTS.md

Tests:
1. Exact Canonical 10-column schema & visual styling (LinkedIn Blue #0A66C2, white bold font, freeze panes A2, widths)
2. Automatic legacy 11-column sheet migration with zero data loss
3. Tier 1 -> Tier 2 fallback: Missing LinkedIn sheet redirects to Applications_Log
4. Tier 1/2 -> Tier 3 emergency fallback: Corrupted non-zip Excel file falls back to CSV + JSON
5. OS PermissionError / File Lock resilience: Zero data loss under simulated locked Excel file
6. Dual failure resilience: Excel + CSV failure gracefully preserves JSON audit
7. Multi-source deduplication: get_applied_job_ids() 5-tier aggregation across Skip List, JSON Log, CSV, and both Excel sheets
8. Dry-run isolation: Guarantee dry-run records do NOT pollute skip list or applied IDs
9. Facade compliance: LinkedInApplyRunner.run(dry_run=True) and log_application_to_excel()
10. Idempotent initialization: Multiple concurrent tracker inits do not duplicate headers or corrupt schema
"""

import os
import sys
import csv
import json
import socket
import pytest
from unittest.mock import MagicMock, patch, AsyncMock
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment

from naukari_pro.modules.linkedin_apply import (
    LinkedInApplicationTracker,
    LinkedInApplyRunner,
    LinkedInApplyEngine,
    LINKEDIN_BLUE_FILL,
    WHITE_BOLD_FONT,
    HEADER_ALIGNMENT
)

# ==============================================================================
# OFFLINE SOCKET ISOLATION GUARD
# ==============================================================================

class BlockedNetworkError(RuntimeError):
    pass

def guard_socket_connect(*args, **kwargs):
    raise BlockedNetworkError("Adversarial Test Violation: Unauthorized network socket call!")

@pytest.fixture(autouse=True)
def enforce_offline_isolation(monkeypatch):
    """Intercept socket primitives to guarantee tests execute strictly offline."""
    monkeypatch.setattr(socket.socket, "connect", guard_socket_connect)
    monkeypatch.setattr(socket, "create_connection", guard_socket_connect)
    monkeypatch.setattr(socket, "getaddrinfo", guard_socket_connect)


# ==============================================================================
# TEST 1: CANONICAL 10 COLUMNS & EXACT VISUAL STYLING SPECIFICATIONS
# ==============================================================================

def test_adversarial_canonical_schema_and_visual_specifications(tmp_path):
    """
    Asserts canonical 10 columns in row 1 of LinkedIn_Applications_Log:
    - Exact header names and ordering
    - LinkedIn Blue #0A66C2 fill
    - Bold white Calibri 11 font
    - Freeze panes 'A2'
    - Column widths A through J matching COLUMN_WIDTHS specification
    """
    excel_path = str(tmp_path / "Tracker_Spec.xlsx")
    json_path = str(tmp_path / "linkedin_applications_log.json")
    skip_path = str(tmp_path / "skip_list_linkedin.json")

    tracker = LinkedInApplicationTracker(
        excel_path=excel_path,
        json_log_path=json_path,
        skip_file_path=skip_path
    )

    wb = openpyxl.load_workbook(excel_path)
    assert "LinkedIn_Applications_Log" in wb.sheetnames, "Sheet LinkedIn_Applications_Log must exist"
    ws = wb["LinkedIn_Applications_Log"]

    # 1. Exact 10 Canonical Headers
    expected_headers = [
        "Job ID",
        "Platform",
        "Company",
        "Job Title",
        "Location",
        "Match Score",
        "Status",
        "Steps Completed",
        "Applied Timestamp",
        "Job URL"
    ]
    actual_headers = [cell.value for cell in ws[1]]
    assert len(actual_headers) == 10, f"Expected exactly 10 columns, got {len(actual_headers)}"
    assert actual_headers == expected_headers, f"Header mismatch: {actual_headers} vs {expected_headers}"

    # 2. Visual Styling Verification
    assert ws.row_dimensions[1].height == 26, "Header row height must be 26pt"
    assert ws.freeze_panes == "A2", "Freeze panes must be set to 'A2'"

    for col_idx in range(1, 11):
        cell = ws.cell(1, col_idx)
        # Background fill must be #0A66C2
        assert cell.fill is not None and cell.fill.fill_type == "solid", f"Col {col_idx} missing solid fill"
        assert cell.fill.start_color.rgb in ("000A66C2", "0A66C2"), f"Col {col_idx} fill color was {cell.fill.start_color.rgb}"
        
        # Font must be bold white
        assert cell.font.bold is True, f"Col {col_idx} font is not bold"
        assert cell.font.color.rgb in ("00FFFFFF", "FFFFFF"), f"Col {col_idx} font color was {cell.font.color.rgb}"
        assert cell.font.name == "Calibri", f"Col {col_idx} font name was {cell.font.name}"

        # Alignment
        assert cell.alignment.horizontal == "center"
        assert cell.alignment.vertical == "center"
        assert cell.alignment.wrap_text is True

    # 3. Column Widths
    expected_widths = {
        "A": 16, "B": 20, "C": 25, "D": 28, "E": 25,
        "F": 14, "G": 18, "H": 16, "I": 20, "J": 45
    }
    for col_letter, width in expected_widths.items():
        assert ws.column_dimensions[col_letter].width == width, (
            f"Column {col_letter} width mismatch: {ws.column_dimensions[col_letter].width} != {width}"
        )


# ==============================================================================
# TEST 2: AUTOMATIC LEGACY 11-COLUMN SCHEMA MIGRATION WITH DATA PRESERVATION
# ==============================================================================

def test_adversarial_legacy_schema_migration_and_data_preservation(tmp_path):
    """
    Stress-tests auto-migration when existing workbook has 11 legacy headers:
    Verifies that existing data rows are transformed correctly into the canonical
    10 columns with Job ID extracted and fields mapped without loss.
    """
    excel_path = str(tmp_path / "Legacy_Tracker.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LinkedIn_Applications_Log"

    legacy_headers = [
        "Application Date", "Company Name", "Job Title", "Match Score (%)",
        "Role Category", "Location", "Application Type", "Status",
        "Job URL", "Easy Apply Steps", "Recruiter Response"
    ]
    ws.append(legacy_headers)
    
    # Add 2 legacy rows
    ws.append([
        "2026-10-25 14:00", "Microsoft IDC", "Senior Azure Data Engineer", "92%",
        "Data Engineering", "Bengaluru", "LinkedIn Easy Apply", "APPLIED",
        "https://www.linkedin.com/jobs/view/4211009988/", "4", "Under Review"
    ])
    ws.append([
        "2026-10-25 15:30", "Amazon AWS", "Data Architect", "86%",
        "Big Data", "Hyderabad", "LinkedIn Easy Apply", "APPLIED",
        "https://www.linkedin.com/jobs/view/4211009989/", "3", "Pending"
    ])
    wb.save(excel_path)

    # Initialize tracker — should detect legacy headers and migrate automatically
    tracker = LinkedInApplicationTracker(excel_path=excel_path)

    wb_migrated = openpyxl.load_workbook(excel_path)
    ws_migrated = wb_migrated["LinkedIn_Applications_Log"]

    # Verify canonical 10 headers
    assert [c.value for c in ws_migrated[1]] == tracker.CANONICAL_HEADERS
    assert ws_migrated.max_column == 10
    assert ws_migrated.max_row == 3

    # Row 1 migrated: Job ID 4211009988 extracted from URL
    row1 = [c.value for c in ws_migrated[2]]
    assert row1[0] == "4211009988"
    assert row1[1] == "LinkedIn Easy Apply"
    assert row1[2] == "Microsoft IDC"
    assert row1[3] == "Senior Azure Data Engineer"
    assert row1[4] == "Bengaluru"
    assert row1[5] == "92%"
    assert row1[6] == "APPLIED"
    assert row1[7] == "4" or row1[7] == 4
    assert row1[8] == "2026-10-25 14:00"
    assert row1[9] == "https://www.linkedin.com/jobs/view/4211009988/"

    # Row 2 migrated: Job ID 4211009989
    row2 = [c.value for c in ws_migrated[3]]
    assert row2[0] == "4211009989"
    assert row2[2] == "Amazon AWS"
    assert row2[6] == "APPLIED"

    # Styling must be applied after migration
    assert ws_migrated.cell(1, 1).fill.start_color.rgb in ("000A66C2", "0A66C2")
    assert ws_migrated.freeze_panes == "A2"


# ==============================================================================
# TEST 3: TIER 1 -> TIER 2 FALLBACK (MISSING LINKEDIN SHEET -> APPLICATIONS_LOG)
# ==============================================================================

def test_adversarial_fallback_to_applications_log_when_linkedin_sheet_missing(tmp_path):
    """
    Stress-tests scenario where 'LinkedIn_Applications_Log' worksheet is absent or
    fails to load, verifying graceful fallback to 'Applications_Log'.
    """
    excel_path = str(tmp_path / "Missing_LinkedIn_Sheet.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Applications_Log"
    ws.append([
        "Application Date", "Company Name", "Job Title", "Match Score (%)",
        "Role Category", "Experience Required", "Location", "Application Type",
        "Status", "Job URL", "Recruiter Response"
    ])
    wb.save(excel_path)

    # Patch _ensure_excel_sheet to avoid auto-creating the LinkedIn sheet
    with patch.object(LinkedInApplicationTracker, "_ensure_excel_sheet"):
        tracker = LinkedInApplicationTracker(excel_path=excel_path)
        job_data = {
            "job_id": "4399001122",
            "company": "Accenture AI",
            "title": "Lead PySpark Engineer",
            "location": "Pune",
            "url": "https://www.linkedin.com/jobs/view/4399001122/"
        }
        res = tracker.log_application(job_data, "APPLIED", match_score=89.0, steps=2)
        assert res is True, "log_application should return True upon fallback"

    # Verify row was written to Applications_Log
    wb_after = openpyxl.load_workbook(excel_path)
    assert "Applications_Log" in wb_after.sheetnames
    ws_after = wb_after["Applications_Log"]
    assert ws_after.max_row == 2
    row_vals = [c.value for c in ws_after[2]]
    assert row_vals[1] == "Accenture AI"
    assert row_vals[2] == "Lead PySpark Engineer"
    assert row_vals[3] == "89%"
    assert row_vals[6] == "Pune"
    assert row_vals[7] == "LinkedIn Easy Apply"
    assert row_vals[8] == "APPLIED"
    assert row_vals[9] == "https://www.linkedin.com/jobs/view/4399001122/"


# ==============================================================================
# TEST 4: TIER 1/2 -> TIER 3 EMERGENCY FALLBACK (CORRUPTED EXCEL FILE)
# ==============================================================================

def test_adversarial_emergency_csv_fallback_on_corrupt_non_zip_excel(tmp_path):
    """
    Stress-tests scenario where the Excel workbook is completely corrupted
    (e.g. truncated file or plain text instead of zip container).
    Verifies:
    1. Tracker initialization and logging does NOT crash or raise unhandled exception.
    2. Emergency fallback writes to linkedin_applications_fallback.csv with canonical headers.
    3. JSON audit log is also written.
    4. Subsequent writes append cleanly without duplicating CSV headers.
    """
    corrupt_excel_path = str(tmp_path / "Corrupt_Workbook.xlsx")
    fallback_csv = str(tmp_path / "linkedin_applications_fallback.csv")
    json_log = str(tmp_path / "linkedin_applications_log.json")

    # Create invalid non-zip file
    with open(corrupt_excel_path, "wb") as f:
        f.write(b"PLAIN_TEXT_GARBAGE_NOT_A_VALID_ZIP_OR_EXCEL_STREAM")

    tracker = LinkedInApplicationTracker(
        excel_path=corrupt_excel_path,
        fallback_csv_path=fallback_csv,
        json_log_path=json_log
    )

    job1 = {
        "job_id": "4500000001",
        "company": "Swiggy Databricks",
        "title": "Data Engineer II",
        "location": "Bengaluru",
        "url": "https://www.linkedin.com/jobs/view/4500000001/"
    }
    result1 = tracker.log_application(job1, "APPLIED", match_score=91.0, steps=3)
    assert result1 is True, "Must return True on emergency CSV fallback"

    # Verify CSV file created
    assert os.path.exists(fallback_csv), "Fallback CSV must exist"
    with open(fallback_csv, "r", encoding="utf-8") as f:
        rows = list(csv.reader(f))
        assert len(rows) == 2, f"Expected header + 1 row, got {len(rows)}"
        assert rows[0] == tracker.CANONICAL_HEADERS
        assert rows[1][0] == "4500000001"
        assert rows[1][2] == "Swiggy Databricks"
        assert rows[1][5] == "91%"
        assert rows[1][6] == "APPLIED"

    # Append second application — headers should NOT be repeated
    job2 = {
        "job_id": "4500000002",
        "company": "Flipkart Analytics",
        "title": "Delta Lake Specialist",
        "location": "Bengaluru",
        "url": "https://www.linkedin.com/jobs/view/4500000002/"
    }
    result2 = tracker.log_application(job2, "APPLIED", match_score=85.0, steps=2)
    assert result2 is True

    with open(fallback_csv, "r", encoding="utf-8") as f:
        rows2 = list(csv.reader(f))
        assert len(rows2) == 3, f"Expected header + 2 rows, got {len(rows2)}"
        assert rows2[2][0] == "4500000002"

    # Verify JSON log also captured both
    assert os.path.exists(json_log)
    with open(json_log, "r", encoding="utf-8") as f:
        json_entries = json.load(f)
        assert len(json_entries) == 2
        assert json_entries[0]["job_id"] == "4500000001"
        assert json_entries[1]["job_id"] == "4500000002"


# ==============================================================================
# TEST 5: SIMULATED PERMISSIONERROR & FILE LOCK (ZERO DATA LOSS)
# ==============================================================================

def test_adversarial_permission_error_file_lock_zero_data_loss(tmp_path):
    """
    Stress-tests scenario where another process has locked the Excel workbook
    or an OS PermissionError occurs during workbook load/save.
    Verifies that the tracker diverts to CSV and JSON without raising an exception.
    """
    excel_path = str(tmp_path / "Locked_Tracker.xlsx")
    fallback_csv = str(tmp_path / "linkedin_applications_fallback.csv")
    json_log = str(tmp_path / "linkedin_applications_log.json")

    # Pre-create valid workbook
    wb = openpyxl.Workbook()
    wb.active.title = "LinkedIn_Applications_Log"
    wb.save(excel_path)

    tracker = LinkedInApplicationTracker(
        excel_path=excel_path,
        fallback_csv_path=fallback_csv,
        json_log_path=json_log
    )

    job_data = {
        "job_id": "4611223344",
        "company": "Razorpay Cloud",
        "title": "Azure Data Platform Engineer",
        "location": "Bengaluru",
        "url": "https://www.linkedin.com/jobs/view/4611223344/"
    }

    # Simulate PermissionError on load_workbook
    with patch("openpyxl.load_workbook", side_effect=PermissionError("[Errno 13] Permission denied: file is locked by another process")):
        res = tracker.log_application(job_data, "APPLIED", match_score=87.0, steps=3)
        assert res is True, "Must return True and not bubble PermissionError"

    # Verify zero data loss: record captured in fallback CSV
    assert os.path.exists(fallback_csv)
    with open(fallback_csv, "r", encoding="utf-8") as f:
        content = f.read()
        assert "4611223344" in content
        assert "Razorpay Cloud" in content

    # And captured in JSON log
    assert os.path.exists(json_log)
    with open(json_log, "r", encoding="utf-8") as f:
        data = json.load(f)
        assert any(d["job_id"] == "4611223344" for d in data)


# ==============================================================================
# TEST 6: DUAL FAILURE RESILIENCE (EXCEL + CSV FAILURE STILL LOGS JSON)
# ==============================================================================

def test_adversarial_dual_failure_resilience_csv_and_excel_failure(tmp_path):
    """
    Extreme adversarial scenario: Both Excel AND CSV fail to write
    (e.g., read-only disk or permission lock on CSV).
    Verifies that JSON audit logging still succeeds and daemon does not crash.
    """
    excel_path = str(tmp_path / "Failing.xlsx")
    fallback_csv = str(tmp_path / "non_writable_dir" / "fallback.csv")  # Directory doesn't exist
    json_log = str(tmp_path / "linkedin_applications_log.json")

    tracker = LinkedInApplicationTracker(
        excel_path=excel_path,
        fallback_csv_path=fallback_csv,
        json_log_path=json_log
    )

    job = {
        "job_id": "4777889900",
        "company": "Databricks Partner",
        "title": "Spark Core Engineer",
        "location": "Remote",
        "url": "https://www.linkedin.com/jobs/view/4777889900/"
    }

    with patch("openpyxl.load_workbook", side_effect=OSError("Disk I/O error on Excel")):
        # Notice: CSV write to non_writable_dir will raise FileNotFoundError, which is caught
        res = tracker.log_application(job, "APPLIED", match_score=94.0, steps=4)
        assert res is True

    # JSON audit log must still have received the record
    assert os.path.exists(json_log)
    with open(json_log, "r", encoding="utf-8") as f:
        entries = json.load(f)
        assert len(entries) == 1
        assert entries[0]["job_id"] == "4777889900"
        assert entries[0]["company"] == "Databricks Partner"


# ==============================================================================
# TEST 7: 5-TIER MULTI-SOURCE DEDUPLICATION AGGREGATION
# ==============================================================================

def test_adversarial_get_applied_job_ids_comprehensive_5_tier_aggregation(tmp_path):
    """
    Verifies get_applied_job_ids() aggregates applied job IDs across all 5 tiers:
    Tier 1: LinkedIn_Applications_Log (Excel)
    Tier 2: Applications_Log (Excel, extracting ID from /jobs/view/<id>)
    Tier 3: linkedin_applications_fallback.csv (CSV)
    Tier 4: linkedin_applications_log.json (JSON audit log, filtering by status)
    Tier 5: skip_list_linkedin.json (Skip list)
    And verifies is_applied() returns True for IDs in ANY tier, and False otherwise.
    """
    excel_path = str(tmp_path / "Multi_Tier_Tracker.xlsx")
    fallback_csv = str(tmp_path / "linkedin_applications_fallback.csv")
    json_log = str(tmp_path / "linkedin_applications_log.json")
    skip_list = str(tmp_path / "skip_list_linkedin.json")

    # 1. Tier 5: Skip list
    with open(skip_list, "w", encoding="utf-8") as f:
        json.dump(["5001", "5002"], f)

    # 2. Tier 4: JSON log (include a FAILED job which must NOT be treated as applied)
    with open(json_log, "w", encoding="utf-8") as f:
        json.dump([
            {"job_id": "5003", "status": "APPLIED"},
            {"job_id": "5004", "status": "SUBMITTED"},
            {"job_id": "5005", "status": "FAILED_MODAL_ERROR"}
        ], f)

    # 3. Tier 3: Emergency CSV fallback
    with open(fallback_csv, "w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["Job ID", "Platform", "Company", "Job Title", "Location", "Match Score", "Status", "Steps Completed", "Applied Timestamp", "Job URL"])
        writer.writerow(["5006", "LinkedIn Easy Apply", "Company X", "Dev", "India", "80%", "APPLIED", "2", "2026-10-26", "https://linkedin.com/jobs/view/5006/"])

    # 4. Tier 1 & Tier 2: Excel workbook with two sheets
    wb = openpyxl.Workbook()
    # Sheet 1: LinkedIn_Applications_Log
    ws_li = wb.active
    ws_li.title = "LinkedIn_Applications_Log"
    ws_li.append(["Job ID", "Platform", "Company", "Job Title", "Location", "Match Score", "Status", "Steps Completed", "Applied Timestamp", "Job URL"])
    ws_li.append(["5007", "LinkedIn Easy Apply", "Company Y", "Dev", "India", "85%", "APPLIED", 2, "2026-10-26", "https://linkedin.com/jobs/view/5007/"])
    # Row with URL regex match
    ws_li.append(["", "LinkedIn Easy Apply", "Company Z", "Dev", "India", "85%", "APPLIED", 2, "2026-10-26", "https://www.linkedin.com/jobs/view/5008/"])

    # Sheet 2: Applications_Log (legacy structure with URL in col 10)
    ws_app = wb.create_sheet(title="Applications_Log")
    ws_app.append(["Date", "Company", "Title", "Score", "Role", "Exp", "Loc", "Type", "Status", "URL", "Response"])
    ws_app.append(["2026-10-26", "Company W", "Dev", "88%", "Data", "2 Yrs", "India", "LinkedIn Easy Apply", "APPLIED", "https://www.linkedin.com/jobs/view/5009/", "Pending"])

    wb.save(excel_path)

    tracker = LinkedInApplicationTracker(
        excel_path=excel_path,
        fallback_csv_path=fallback_csv,
        json_log_path=json_log,
        skip_file_path=skip_list
    )

    applied_ids = tracker.get_applied_job_ids()

    # All expected applied IDs must be present
    expected_ids = {"5001", "5002", "5003", "5004", "5006", "5007", "5008", "5009"}
    for expected_id in expected_ids:
        assert expected_id in applied_ids, f"Expected ID {expected_id} missing from applied_ids set"
        assert tracker.is_applied(expected_id) is True, f"is_applied({expected_id}) must return True"

    # Failed job 5005 must NOT be in applied IDs
    assert "5005" not in applied_ids, "Failed job 5005 should not be counted as applied"
    assert tracker.is_applied("5005") is False

    # Unseen ID
    assert tracker.is_applied("9999999") is False

    # Boundary conditions
    assert tracker.is_applied("") is False
    assert tracker.is_applied(None) is False
    assert tracker.is_applied("   ") is False


# ==============================================================================
# TEST 8: DRY-RUN PURITY & STATUS FILTERING (SKIP LIST IS NEVER MUTATED)
# ==============================================================================

def test_adversarial_dry_run_purity_and_status_filtering(tmp_path):
    """
    Validates the strict dry-run purity invariant:
    1. During dry-run (dry_run=True, status='SUBMITTED_DRY_RUN'), skip_list_linkedin.json is NOT modified.
    2. Dry-run rows in Excel are excluded from get_applied_job_ids() (status containing 'DRY_RUN').
    3. Authentic live applications (dry_run=False, status='APPLIED') ARE recorded in skip list.
    """
    excel_path = str(tmp_path / "DryRun_Tracker.xlsx")
    skip_path = str(tmp_path / "skip_list_linkedin.json")
    json_path = str(tmp_path / "linkedin_applications_log.json")

    tracker = LinkedInApplicationTracker(
        excel_path=excel_path,
        skip_file_path=skip_path,
        json_log_path=json_path
    )

    dry_job = {
        "job_id": "6001",
        "company": "Simulated Corp",
        "title": "Data Engineer Dry Run",
        "url": "https://www.linkedin.com/jobs/view/6001/"
    }

    # 1. Log Dry Run
    tracker.log_application(dry_job, status="SUBMITTED_DRY_RUN", match_score=85.0, steps=3, dry_run=True)

    # Assert skip list does NOT exist or is empty
    if os.path.exists(skip_path):
        with open(skip_path, "r", encoding="utf-8") as f:
            data = json.load(f)
            assert "6001" not in data, "Dry run job must NEVER enter skip_list_linkedin.json!"

    # Dry run should NOT be returned by get_applied_job_ids() or is_applied()
    assert tracker.is_applied("6001") is False, "Dry-run job ID 6001 must not be marked applied"
    assert "6001" not in tracker.get_applied_job_ids()

    # 2. Log Authentic Live Application
    live_job = {
        "job_id": "6002",
        "company": "Real Corp",
        "title": "Azure Data Engineer",
        "url": "https://www.linkedin.com/jobs/view/6002/"
    }
    tracker.log_application(live_job, status="APPLIED", match_score=92.0, steps=4, dry_run=False)

    # Live job MUST be present in skip list
    assert os.path.exists(skip_path)
    with open(skip_path, "r", encoding="utf-8") as f:
        live_skip_data = json.load(f)
        assert "6002" in live_skip_data, "Live job 6002 must be written to skip list"
        assert "6001" not in live_skip_data

    assert tracker.is_applied("6002") is True
    assert "6002" in tracker.get_applied_job_ids()


# ==============================================================================
# TEST 9: LINKEDIN APPLY RUNNER FACADE INTERFACE CONTRACT & OFFLINE DISPATCH
# ==============================================================================

def test_adversarial_linkedin_apply_runner_facade(tmp_path, monkeypatch):
    """
    Stress-tests the LinkedInApplyRunner facade:
    1. Instantiation with headless, dry_run, limit, min_score
    2. log_application_to_excel() delegation to LinkedInApplicationTracker
    3. run() dispatch under dry_run=True executes offline with zero sockets
    """
    import asyncio
    monkeypatch.setenv("WORKSPACE_DIR", str(tmp_path))
    excel_path = str(tmp_path / "Facade_Tracker.xlsx")

    runner = LinkedInApplyRunner(headless=True, dry_run=True, limit=5, min_score=70.0)

    # 1. Test log_application_to_excel
    entry = {
        "job_id": "7001",
        "company": "Facade Test Corp",
        "title": "Azure PySpark Architect",
        "location": "Bengaluru",
        "url": "https://www.linkedin.com/jobs/view/7001/",
        "status": "APPLIED",
        "match_score": 95.0,
        "steps": 4
    }
    log_success = runner.log_application_to_excel(entry, tracker_path=excel_path)
    assert log_success is True

    # Verify sheet created with canonical headers and styled
    wb = openpyxl.load_workbook(excel_path)
    assert "LinkedIn_Applications_Log" in wb.sheetnames
    ws = wb["LinkedIn_Applications_Log"]
    assert ws.max_row == 2
    assert ws.cell(2, 1).value == "7001"
    assert ws.cell(2, 3).value == "Facade Test Corp"

    # 2. Test runner.run() in dry_run mode (100% offline simulation)
    # Under enforce_offline_isolation, any network socket call raises BlockedNetworkError
    result = asyncio.run(runner.run(
        keywords=["Azure Data Engineer", "Databricks"],
        locations=["Bengaluru", "Hyderabad"]
    ))

    assert isinstance(result, dict)
    assert "discovered" in result
    assert "scored_above_threshold" in result
    assert "dry_run_verified" in result
    assert "jobs" in result
    assert result["dry_run_verified"] >= 3
    assert result["discovered"] >= 5


# ==============================================================================
# TEST 10: IDEMPOTENT INITIALIZATION & CONCURRENT MULTI-TRACKER SAFETY
# ==============================================================================

def test_adversarial_concurrent_or_repeated_initialization_safety(tmp_path):
    """
    Validates that repeated or concurrent creation of LinkedInApplicationTracker instances
    against the same Excel file does NOT:
    1. Duplicate header rows
    2. Overwrite existing rows
    3. Corrupt formatting or freeze panes
    """
    excel_path = str(tmp_path / "Idempotent_Tracker.xlsx")

    # Instance 1: Creates workbook and headers
    t1 = LinkedInApplicationTracker(excel_path=excel_path)
    t1.log_application({
        "job_id": "8001",
        "company": "Company Alpha",
        "title": "Data Engineer",
        "url": "https://www.linkedin.com/jobs/view/8001/"
    }, "APPLIED")

    # Instance 2: Initializes on same file
    t2 = LinkedInApplicationTracker(excel_path=excel_path)
    t2.log_application({
        "job_id": "8002",
        "company": "Company Beta",
        "title": "PySpark Lead",
        "url": "https://www.linkedin.com/jobs/view/8002/"
    }, "APPLIED")

    # Instance 3: Initializes again on same file
    t3 = LinkedInApplicationTracker(excel_path=excel_path)

    wb = openpyxl.load_workbook(excel_path)
    ws = wb["LinkedIn_Applications_Log"]

    # Verify exactly 1 header row + 2 data rows = 3 rows total
    assert ws.max_row == 3, f"Expected 3 rows (1 header + 2 data), got {ws.max_row}"
    assert [c.value for c in ws[1]] == t1.CANONICAL_HEADERS
    assert ws.cell(2, 1).value == "8001"
    assert ws.cell(3, 1).value == "8002"
    assert ws.freeze_panes == "A2"
    assert ws.cell(1, 1).fill.start_color.rgb in ("000A66C2", "0A66C2")
