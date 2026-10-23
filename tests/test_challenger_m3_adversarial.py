"""
Adversarial Stress Test Suite for Milestone 3 (Autonomous Multi-Channel Application Engine & Dry-Run)
Authored by Challenger 2 (teamwork_preview_challenger_m3_2)

Empirically verifies:
1. Complete offline dry-run isolation (network independent via socket blocking).
2. Absolute side-effect freedom and Excel tracker immutability over 50 consecutive cycles.
3. Zero browser process creation or subprocess execution during dry-run.
4. Complete structural and candidate fidelity across all 8 ATS portal types in CareerPortalFiller.
5. Boundary conditions and robustness under adversarial/corrupted inputs.
"""

import os
import sys
import json
import time
import socket
import pytest
import subprocess
from unittest.mock import patch, MagicMock
from pathlib import Path
import openpyxl

from naukari_pro.config import TRACKER_PATH, load_config
from naukari_pro.modules.auto_apply import (
    ApplicationPayload,
    AutoApplyEngine,
    build_dry_run_payload,
    generate_dry_run_payload,
    evaluate_job_match,
    is_salary_acceptable
)
from naukari_pro.modules.career_portals import (
    CareerPortalFiller,
    load_applicant_profile
)
from naukari_pro.modules.multi_search import (
    UniversalJobSearchEngine
)
from naukari_pro.cli import (
    handle_auto_apply,
    handle_portal_apply,
    handle_multi_search,
    handle_daily_cycle
)

ALL_8_ATS_SAMPLES = [
    ("Greenhouse", "https://boards.greenhouse.io/databricks/jobs/12345"),
    ("Greenhouse", "https://careers.company.com/apply?gh_jid=67890"),
    ("Ashby", "https://jobs.ashbyhq.com/openai/23456"),
    ("Lever", "https://jobs.lever.co/stripe/34567"),
    ("Workday", "https://walmart.myworkdayjobs.com/careers/job/45678"),
    ("Workday", "https://wd5.myworkday.com/recruiting/corp/job/56789"),
    ("Rippling", "https://ats.rippling.com/scaleai/jobs/67890"),
    ("LinkedIn", "https://www.linkedin.com/jobs/view/789012345/"),
    ("SmartRecruiters", "https://jobs.smartrecruiters.com/corp/89012"),
    ("Generic Portal", "https://careers.enterprise.com/job/90123")
]


# ==============================================================================
# SECTION 1: 100% OFFLINE ISOLATION & SOCKET BLOCKING VERIFICATION
# ==============================================================================

class BlockedNetworkError(RuntimeError):
    """Raised when an unauthorized network connection attempt is made in offline mode."""
    pass


def guard_socket_connect(*args, **kwargs):
    raise BlockedNetworkError("Adversarial Trap: Network socket connection attempted during dry-run!")


@pytest.fixture
def enforce_offline_isolation(monkeypatch):
    """
    Monkeypatches socket-level primitives to simulate an air-gapped / offline environment.
    Any attempt to open a socket or establish a TCP connection triggers BlockedNetworkError.
    """
    monkeypatch.setattr(socket.socket, "connect", guard_socket_connect)
    monkeypatch.setattr(socket, "create_connection", guard_socket_connect)
    monkeypatch.setattr(socket, "getaddrinfo", guard_socket_connect)


def test_offline_isolation_auto_apply_engine(enforce_offline_isolation):
    """
    Empirically proves AutoApplyEngine.run(dry_run=True) is 100% network-independent.
    Must never attempt socket connections or DNS resolution.
    """
    engine = AutoApplyEngine(dry_run=True)
    count = engine.run(limit=10)
    assert count >= 1
    assert len(engine.payloads) >= 1
    assert all(p.dry_run is True for p in engine.payloads)
    assert any(p.status == "VALIDATED_PAYLOAD" for p in engine.payloads)


def test_offline_isolation_career_portal_filler(enforce_offline_isolation):
    """
    Empirically proves CareerPortalFiller dry-run methods operate 100% offline.
    """
    filler = CareerPortalFiller()
    # Direct fill_dry_run
    payload = filler.fill_dry_run(
        portal_url="https://boards.greenhouse.io/databricks/jobs/999",
        company="Databricks",
        role="Lead Azure Data Engineer",
        raw_jd="Azure Databricks PySpark Delta Lake Kafka",
        salary="30 LPA"
    )
    assert isinstance(payload, ApplicationPayload)
    assert payload.dry_run is True
    assert payload.status == "VALIDATED_PAYLOAD"

    # Unified fill(dry_run=True)
    success, unified_payload = filler.fill(
        portal_url="https://jobs.lever.co/palantir/888",
        company="Palantir",
        role="Azure Data Platform Engineer",
        dry_run=True
    )
    assert success is True
    assert isinstance(unified_payload, ApplicationPayload)
    assert unified_payload.dry_run is True


def test_offline_isolation_multi_search_engine(enforce_offline_isolation):
    """
    Empirically proves UniversalJobSearchEngine handles offline conditions without unhandled crashes.
    When network is severed, search_hacker_news catches the socket error and falls back to offline curated fixtures.
    """
    engine = UniversalJobSearchEngine()

    # 1. HN payload generator with mock jobs
    mock_jobs = [
        {
            "source": "Hacker News (Who's Hiring)",
            "title": "Senior Azure Data Engineer",
            "company": "TechGCC",
            "location": "Remote",
            "url": "https://jobs.ashbyhq.com/techgcc/1",
            "snippet": "Azure Databricks PySpark Delta Lake Kafka Medallion"
        }
    ]
    hn_payloads = engine.generate_application_payloads_from_hn(limit=2, jobs=mock_jobs)
    assert len(hn_payloads) == 1
    assert hn_payloads[0].match_score >= 70

    # 2. LinkedIn payload generator (algorithmic URL generation, no network)
    li_payloads = engine.generate_application_payloads_from_linkedin(limit=3)
    assert len(li_payloads) >= 1
    assert all(p.dry_run is True for p in li_payloads)

    # 3. search_and_dry_run under socket block: falls back to offline curated fixtures gracefully
    search_payloads = engine.search_and_dry_run(limit=2)
    assert len(search_payloads) >= 2
    assert all(p.dry_run is True for p in search_payloads)


def test_offline_isolation_cli_dry_run_routes(enforce_offline_isolation):
    """
    Empirically proves all CLI dry-run handlers execute without making network requests.
    """
    # CLI auto-apply dry-run
    handle_auto_apply(dry_run=True, limit=5)

    # CLI portal-apply dry-run
    handle_portal_apply(
        url="https://jobs.ashbyhq.com/scale/123",
        company="Scale",
        role="Azure Data Engineer",
        dry_run=True
    )

    # CLI multi-search dry-run
    payloads = handle_multi_search(dry_run=True)
    assert len(payloads) >= 1

    # CLI daily cycle dry-run (Steps 1, 2, 3 must avoid network and browser)
    handle_daily_cycle(dry_run=True)


# ==============================================================================
# SECTION 2: TRACKER IMMUTABILITY ACROSS 50 CONSECUTIVE DRY-RUN CYCLES
# ==============================================================================

def test_tracker_immutability_50_consecutive_cycles():
    """
    Adversarial Stress Test:
    Executes 50 consecutive dry-run application cycles across all engines and verifies:
    - Naukri_Live_Application_Tracker.xlsx file modification timestamp is 100% UNCHANGED.
    - Naukri_Live_Application_Tracker.xlsx file byte size is 100% UNCHANGED.
    - openpyxl sheet max_row and max_column across all sheets are 100% UNCHANGED.
    - Complete byte/content fidelity of every cell in Applications_Log.
    - Zero writes to skip_list.json or applications logs.
    """
    assert os.path.exists(TRACKER_PATH), f"Tracker workbook must exist at {TRACKER_PATH}"

    # Initial snapshot
    initial_mtime = os.path.getmtime(TRACKER_PATH)
    initial_size = os.path.getsize(TRACKER_PATH)

    wb_initial = openpyxl.load_workbook(TRACKER_PATH, data_only=True)
    initial_sheets = list(wb_initial.sheetnames)
    initial_sheet_shapes = {s: (wb_initial[s].max_row, wb_initial[s].max_column) for s in initial_sheets}
    
    # Snapshot all cell values in Applications_Log
    ws_app_initial = wb_initial["Applications_Log"]
    initial_cells = [
        [ws_app_initial.cell(row=r, column=c).value for c in range(1, ws_app_initial.max_column + 1)]
        for r in range(1, ws_app_initial.max_row + 1)
    ]
    wb_initial.close()

    # Snapshot skip_list.json if it exists
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    skip_file = os.path.join(base_dir, "skip_list.json")
    initial_skip_mtime = os.path.getmtime(skip_file) if os.path.exists(skip_file) else None
    initial_skip_content = Path(skip_file).read_text(encoding="utf-8") if os.path.exists(skip_file) else None

    # Execute 50 consecutive adversarial dry-run cycles
    for cycle in range(1, 51):
        # 1. AutoApplyEngine dry-run
        engine = AutoApplyEngine(dry_run=True)
        count = engine.run(limit=10)
        assert count == 3, f"Cycle {cycle}: AutoApplyEngine expected 3 validated payloads"

        # 2. CareerPortalFiller dry-run
        filler = CareerPortalFiller()
        p1 = filler.fill_dry_run(
            portal_url=f"https://boards.greenhouse.io/databricks/jobs/{cycle}",
            company="Databricks",
            role="Azure Data Engineer",
            raw_jd="Azure Databricks PySpark Delta Lake Kafka",
            salary="28 LPA"
        )
        assert p1.status == "VALIDATED_PAYLOAD"

        success, p2 = filler.fill(
            portal_url=f"https://jobs.lever.co/stripe/{cycle}",
            company="Stripe",
            role="Azure Data Platform Engineer",
            dry_run=True
        )
        assert success is True and p2.dry_run is True

        # 3. UniversalJobSearchEngine dry-run
        search_engine = UniversalJobSearchEngine()
        li_payloads = search_engine.generate_application_payloads_from_linkedin(limit=2)
        assert len(li_payloads) >= 1

    # Forensic Verification after 50 cycles
    final_mtime = os.path.getmtime(TRACKER_PATH)
    final_size = os.path.getsize(TRACKER_PATH)

    assert final_mtime == initial_mtime, (
        f"VIOLATION: Tracker file mtime modified during 50 dry-run cycles! "
        f"Initial: {initial_mtime}, Final: {final_mtime}"
    )
    assert final_size == initial_size, (
        f"VIOLATION: Tracker file size changed during 50 dry-run cycles! "
        f"Initial: {initial_size} bytes, Final: {final_size} bytes"
    )

    wb_final = openpyxl.load_workbook(TRACKER_PATH, data_only=True)
    final_sheets = list(wb_final.sheetnames)
    assert final_sheets == initial_sheets, "Sheet list mutated!"

    for s in initial_sheets:
        init_r, init_c = initial_sheet_shapes[s]
        fin_r, fin_c = wb_final[s].max_row, wb_final[s].max_column
        assert (fin_r, fin_c) == (init_r, init_c), (
            f"VIOLATION: Sheet '{s}' dimensions mutated from ({init_r}, {init_c}) to ({fin_r}, {fin_c})!"
        )

    ws_app_final = wb_final["Applications_Log"]
    final_cells = [
        [ws_app_final.cell(row=r, column=c).value for c in range(1, ws_app_final.max_column + 1)]
        for r in range(1, ws_app_final.max_row + 1)
    ]
    wb_final.close()

    assert final_cells == initial_cells, "VIOLATION: Cell contents in Applications_Log altered!"

    # Verify skip_list was not altered
    if initial_skip_mtime is not None:
        assert os.path.getmtime(skip_file) == initial_skip_mtime, "skip_list.json was modified!"
        assert Path(skip_file).read_text(encoding="utf-8") == initial_skip_content, "skip_list.json content changed!"


# ==============================================================================
# SECTION 3: BROWSER PROCESS & SUBPROCESS SUPPRESSION
# ==============================================================================

def test_zero_browser_processes_spawned_during_dry_run(monkeypatch):
    """
    Adversarial Process Verification:
    Spies on subprocess and Playwright context launchers to verify that zero browser
    processes (Google Chrome, Chromium, osascript GUI automation) are initiated in dry-run mode.
    """
    forbidden_calls = []

    original_run = subprocess.run
    def spied_subprocess_run(cmd, *args, **kwargs):
        cmd_str = " ".join(cmd) if isinstance(cmd, list) else str(cmd)
        if any(needle in cmd_str.lower() for needle in ["chrome", "osascript", "chromium", "playwright"]):
            forbidden_calls.append(f"subprocess.run: {cmd_str}")
        return original_run(cmd, *args, **kwargs)

    original_popen = subprocess.Popen
    def spied_subprocess_popen(cmd, *args, **kwargs):
        cmd_str = " ".join(cmd) if isinstance(cmd, list) else str(cmd)
        if any(needle in cmd_str.lower() for needle in ["chrome", "osascript", "chromium", "playwright"]):
            forbidden_calls.append(f"subprocess.Popen: {cmd_str}")
        return original_popen(cmd, *args, **kwargs)

    monkeypatch.setattr(subprocess, "run", spied_subprocess_run)
    monkeypatch.setattr(subprocess, "Popen", spied_subprocess_popen)

    # 1. AutoApplyEngine dry-run
    auto_engine = AutoApplyEngine(dry_run=True)
    auto_engine.run(limit=10)

    # 2. CareerPortalFiller dry-run
    filler = CareerPortalFiller()
    filler.fill_dry_run("https://boards.greenhouse.io/databricks/jobs/1", company="Databricks")
    filler.fill("https://jobs.ashbyhq.com/scale/2", company="Scale", dry_run=True)

    # 3. UniversalJobSearchEngine dry-run
    search_engine = UniversalJobSearchEngine()
    search_engine.generate_application_payloads_from_linkedin(limit=3)

    # 4. CLI daily cycle dry-run
    handle_daily_cycle(dry_run=True)

    assert len(forbidden_calls) == 0, (
        f"VIOLATION: Browser or GUI subprocesses initiated during dry-run!\n"
        f"Calls detected: {forbidden_calls}"
    )


# ==============================================================================
# SECTION 4: ALL 8 ATS PORTAL TYPES CANDIDATE AND STRUCTURE FIDELITY
# ==============================================================================

@pytest.mark.parametrize("expected_ats, sample_url", ALL_8_ATS_SAMPLES)
def test_all_8_ats_portal_types_structural_and_candidate_fidelity(expected_ats, sample_url):
    """
    Validates that CareerPortalFiller correctly recognizes all 8 supported ATS types,
    generates legitimate autofill JavaScript containing candidate facts, and emits
    ApplicationPayload conforming to the standard schema with candidate ground truth.
    """
    filler = CareerPortalFiller()

    # 1. ATS Type Detection
    detected = filler.detect_ats(sample_url)
    assert detected == expected_ats, f"URL {sample_url} should detect as {expected_ats}, got {detected}"

    # 2. Generate Dry-Run ApplicationPayload
    payload = filler.fill_dry_run(
        portal_url=sample_url,
        company="Enterprise GCC",
        role="Azure Data Engineer",
        raw_jd="Azure Databricks PySpark Delta Lake Medallion Architecture Kafka SQL",
        salary="25 - 32 LPA",
        calculate_score=True
    )

    # 3. Contract Schema Verification
    assert isinstance(payload, ApplicationPayload)
    assert payload.dry_run is True
    assert payload.status == "VALIDATED_PAYLOAD"
    assert payload.company == "Enterprise GCC"
    assert payload.role == "Azure Data Engineer"
    assert payload.title == "Azure Data Engineer"
    assert payload.job_title == "Azure Data Engineer"
    assert payload.url == sample_url
    assert payload.salary == "25 - 32 LPA"
    assert payload.match_score >= 70
    assert payload.portal_type == f"Career Portal ({expected_ats})"

    # 4. Candidate Ground Truth Verification (AGENTS.md Section 9)
    answers = payload.answers
    assert answers["first_name"] == "Lavkush"
    assert answers["last_name"] == "Kumar"
    assert answers["full_name"] == "Lavkush Kumar"
    assert answers["email"] == "lavkushry@gmail.com"
    assert answers["phone"] == "+91-9122036484"
    assert answers["location"] == "Hyderabad"
    assert "linkedin.com/in/lavkushry" in answers["linkedin"]
    assert "github.com/lavkushry" in answers["github"]
    assert answers["experience"] == "2.5"
    assert answers["notice_period"] == "30"
    assert float(answers["current_ctc_lpa"]) == 12.0
    assert float(answers["expected_ctc_lpa"]) in [20.0, 25.0]
    assert answers["work_auth"] == "Yes"
    assert answers["visa_sponsorship"] == "No"
    assert answers["ats_type"] == expected_ats

    # 5. Payload Data & Autofill Script Verification
    payload_data = payload.payload_data
    assert payload_data["ats_type"] == expected_ats
    assert payload_data["has_script"] is True
    assert payload_data["script_length"] > 100

    # 6. JavaScript Generator Verification
    js_script = filler.generate_fill_script(expected_ats)
    assert isinstance(js_script, str)
    assert f"ats: '{expected_ats}'" in js_script
    assert "Lavkush" in js_script
    assert "lavkushry@gmail.com" in js_script
    assert "+91-9122036484" in js_script

    # 7. Serialization Integrity
    as_dict = payload.to_dict()
    assert as_dict["role"] == "Azure Data Engineer"
    assert as_dict["dry_run"] is True

    as_json = payload.to_json()
    assert isinstance(as_json, str)
    parsed = json.loads(as_json)
    assert parsed["company"] == "Enterprise GCC"
    assert parsed["match_score"] >= 70


# ==============================================================================
# SECTION 5: ADVERSARIAL INPUTS & BOUNDARY CONDITIONS
# ==============================================================================

def test_adversarial_corrupted_job_feed_dry_run():
    """
    Adversarial Stress Test:
    Passes malformed, incomplete, and extreme job tuples into AutoApplyEngine.run_offline_dry_run.
    Must never raise unhandled exceptions and must flag items with correct rejection statuses.
    """
    corrupted_feed = [
        # Missing title and company
        {"salary": "25 LPA", "jd": "Azure Databricks PySpark Delta Lake"},
        # None salary
        {"title": "Azure Data Engineer", "company": "Co", "salary": None, "jd": "Databricks"},
        # Extremely low salary below 12 LPA
        {"title": "Azure Data Engineer", "company": "Co", "salary": "4 - 6 Lacs PA", "jd": "Azure Databricks PySpark"},
        # Negative / absurd salary
        {"title": "Azure Data Engineer", "company": "Co", "salary": "Free Internship (0 LPA)", "jd": "Azure Databricks"},
        # Completely irrelevant skills (<70% match)
        {"title": "Junior Chef", "company": "Restaurant", "salary": "20 LPA", "jd": "Italian cuisine pizza pasta"},
        # Missing all fields except empty dict
        {}
    ]

    engine = AutoApplyEngine(dry_run=True)
    payloads = engine.run_offline_dry_run(job_feed=corrupted_feed, limit=10)

    assert len(payloads) == 6
    # Low salary items must be flagged REJECTED_LOW_SALARY
    assert payloads[2].status == "REJECTED_LOW_SALARY"
    assert payloads[3].status == "REJECTED_LOW_SALARY"
    # Chef and empty dict must be flagged REJECTED_BELOW_THRESHOLD
    assert payloads[4].status == "REJECTED_BELOW_THRESHOLD"
    assert payloads[5].status == "REJECTED_BELOW_THRESHOLD"
    assert all(p.dry_run is True for p in payloads)


def test_adversarial_malformed_portal_urls():
    """
    Adversarial Stress Test:
    Tests CareerPortalFiller with malformed, non-standard, or extreme URL inputs.
    Must fall back safely without crashing.
    """
    filler = CareerPortalFiller()

    bizarre_urls = [
        "",
        "not_a_valid_url",
        "https://",
        "ftp://invalid.scheme.org/job",
        "https://example.com/" + "a" * 2000,
        "https://unknown-portal-provider.io/openings/123?candidate=xyz"
    ]

    for u in bizarre_urls:
        detected = filler.detect_ats(u)
        assert detected == "Generic Portal"

        payload = filler.fill_dry_run(portal_url=u, company="BizarreCo", role="Azure Data Engineer")
        assert isinstance(payload, ApplicationPayload)
        assert payload.dry_run is True
        assert payload.portal_type == "Career Portal (Generic Portal)"


def test_adversarial_match_score_exact_boundary():
    """
    Adversarial Boundary Test:
    Ensures exact boundary gating in build_dry_run_payload and evaluate_job_match:
    - score < 70 -> REJECTED_BELOW_THRESHOLD
    - score >= 70 -> VALIDATED_PAYLOAD
    """
    # 55 points: Azure (20) + Databricks (20) + PySpark (15) = 55 (<70)
    p_55 = build_dry_run_payload(
        company="Co55",
        title="Data Engineer",
        url="https://test.com/55",
        salary="25 LPA",
        raw_jd="Azure Databricks PySpark"
    )
    assert p_55.match_score == 55
    assert p_55.status == "REJECTED_BELOW_THRESHOLD"

    # 65 points: Azure (20) + Databricks (20) + PySpark (15) + SQL (10) = 65 (<70)
    p_65 = build_dry_run_payload(
        company="Co65",
        title="Data Engineer",
        url="https://test.com/65",
        salary="25 LPA",
        raw_jd="Azure Databricks PySpark SQL"
    )
    assert p_65.match_score == 65
    assert p_65.status == "REJECTED_BELOW_THRESHOLD"

    # 70 points: Azure (20) + Databricks (20) + PySpark (15) + Delta Lake (15) = 70 (>=70)
    p_70 = build_dry_run_payload(
        company="Co70",
        title="Data Engineer",
        url="https://test.com/70",
        salary="25 LPA",
        raw_jd="Azure Databricks PySpark Delta Lake"
    )
    assert p_70.match_score == 70
    assert p_70.status == "VALIDATED_PAYLOAD"


# ==============================================================================
# SECTION 6: EMPIRICAL CHALLENGES (SURFACING IMPLEMENTATION DEFECTS)
# ==============================================================================

def test_adversarial_portal_dry_run_rejects_sub_threshold_role_without_jd():
    """
    CRITICAL CHALLENGE:
    Exposes Bug in CareerPortalFiller.fill_dry_run (career_portals.py line 394):
    When raw_jd is empty or not passed, 'computed < 70' is overridden to 85,
    allowing completely irrelevant roles (Frontend React Developer, Cashier)
    to receive 'VALIDATED_PAYLOAD' and 85% score, bypassing the gate.
    """
    filler = CareerPortalFiller()
    payload = filler.fill_dry_run(
        portal_url="https://boards.greenhouse.io/acme/jobs/123",
        company="Startup",
        role="Frontend React Developer"
    )
    assert payload.match_score == 0, (
        f"VIOLATION: Sub-threshold role received fabricated score {payload.match_score}% (expected 0%)!"
    )
    assert payload.status == "REJECTED_BELOW_THRESHOLD", (
        f"VIOLATION: Sub-threshold role received status {payload.status} (expected REJECTED_BELOW_THRESHOLD)!"
    )


def test_adversarial_detect_ats_none_safety():
    """
    CRITICAL CHALLENGE:
    Exposes Bug in CareerPortalFiller.detect_ats (career_portals.py line 180):
    Passing None as url raises unhandled AttributeError: 'NoneType' object has no attribute 'lower'.
    Must return 'Generic Portal' safely.
    """
    filler = CareerPortalFiller()
    ats = filler.detect_ats(None)
    assert ats == "Generic Portal"

