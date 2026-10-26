"""
Comprehensive Offline Test Suite: Autonomous LinkedIn Career & Easy Apply Engine
Target: Milestone 1 (LinkedIn Stealth Session & Cookie Architecture)
Governed by TEST_INFRA.md, ORIGINAL_REQUEST.md, and PROJECT.md

100% Offline, Zero Real Browser Processes, Zero Network Calls.
"""

import os
import sys
import json
import socket
import argparse
import urllib.parse
import pytest
from unittest.mock import AsyncMock, MagicMock, patch, call

from naukari_pro.modules.linkedin_session import (
    LinkedInSessionManager,
    LinkedInError,
    LinkedInSessionError,
    LinkedInCookieMissingError,
    LinkedInCookieExpiredError,
    LinkedInCheckpointError,
    DEFAULT_LINKEDIN_SESSION_DIR,
    DEFAULT_LINKEDIN_LOCK_FILE,
    DEFAULT_LINKEDIN_COOKIE_FILE,
    STEALTH_JS,
    mask_cookie_token,
    clean_avatar_name
)
from naukari_pro.modules.linkedin_apply import (
    LinkedInScreeningSolver,
    LinkedInEasyApplyNavigator,
    LinkedInApplicationTracker,
    LinkedInApplyEngine,
    LinkedInApplyError,
    LinkedInModalNotFoundError,
    LinkedInSubmissionFailedError,
    get_candidate_profile
)
from naukari_pro.cli import handle_linkedin_verify, handle_linkedin_apply
import asyncio
import functools

def run_async(coro):
    @functools.wraps(coro)
    def wrapper(*args, **kwargs):
        return asyncio.run(coro(*args, **kwargs))
    return wrapper


# ==============================================================================
# SECTION 0: OFFLINE ISOLATION & SYNTHETIC MOCK FACTORIES
# ==============================================================================

class BlockedNetworkError(RuntimeError):
    pass

def guard_socket_connect(*args, **kwargs):
    raise BlockedNetworkError("Offline Test Violation: Unauthorized network connection attempt!")

@pytest.fixture(autouse=True)
def enforce_offline_isolation(monkeypatch):
    """Intercept socket primitives to guarantee tests execute strictly offline."""
    monkeypatch.setattr(socket.socket, "connect", guard_socket_connect)
    monkeypatch.setattr(socket, "create_connection", guard_socket_connect)
    monkeypatch.setattr(socket, "getaddrinfo", guard_socket_connect)

def make_mock_locator(count: int = 1, visible: bool = True, text: str = "", attr_map: dict = None):
    """Creates a deterministic Playwright locator mock supporting async operations."""
    loc = MagicMock()
    loc.count = AsyncMock(return_value=count)
    loc.is_visible = AsyncMock(return_value=visible)
    loc.inner_text = AsyncMock(return_value=text)
    attr_map = attr_map or {}
    loc.get_attribute = AsyncMock(side_effect=lambda a: attr_map.get(a, None))
    loc.first = loc
    loc.last = loc
    loc.nth = MagicMock(return_value=loc)
    return loc

def make_mock_page(current_url: str = "https://www.linkedin.com/feed/", locators: dict = None, body_text: str = ""):
    """Creates an async MockPage."""
    page = MagicMock()
    page.url = current_url
    page.goto = AsyncMock(return_value=None)
    page.wait_for_selector = AsyncMock(return_value=None)
    page.wait_for_load_state = AsyncMock(return_value=None)
    page.content = AsyncMock(return_value=f"<html><body>{body_text}</body></html>")
    page.close = AsyncMock(return_value=None)
    page.evaluate = AsyncMock(return_value=None)

    locators = locators or {}
    body_loc = make_mock_locator(count=1, visible=True, text=body_text)

    def get_locator(selector: str):
        if selector == "body":
            return body_loc
        if selector in locators:
            return locators[selector]
        return make_mock_locator(count=0, visible=False, text="")

    page.locator = MagicMock(side_effect=get_locator)
    return page

def make_mock_context(page=None):
    """Creates an async MockBrowserContext."""
    ctx = MagicMock()
    p = page or make_mock_page()
    ctx.pages = [p]
    ctx.new_page = AsyncMock(return_value=p)
    ctx.add_init_script = AsyncMock(return_value=None)
    ctx.add_cookies = AsyncMock(return_value=None)
    ctx.cookies = AsyncMock(return_value=[])
    
    close_listeners = []
    def on_func(event, handler):
        if event == "close":
            close_listeners.append(handler)
    ctx.on = MagicMock(side_effect=on_func)
    
    async def close_func():
        for handler in close_listeners:
            try:
                handler()
            except TypeError:
                try:
                    handler(ctx)
                except Exception:
                    pass
            except Exception:
                pass
    ctx.close = AsyncMock(side_effect=close_func)
    return ctx

def make_mock_playwright(context=None):
    """Creates a mock Playwright instance."""
    pw = MagicMock()
    ctx = context or make_mock_context()
    pw.chromium = MagicMock()
    pw.chromium.launch_persistent_context = AsyncMock(return_value=ctx)
    pw.chromium.launch = AsyncMock(return_value=MagicMock(new_context=AsyncMock(return_value=ctx)))
    return pw

def make_playwright_cm_mock(context=None, page=None):
    """Creates a mock Playwright context manager compatible with both sync and async Playwright."""
    ctx = context or make_mock_context(page=page)
    pw = make_mock_playwright(context=ctx)
    cm = MagicMock()
    cm.__enter__ = MagicMock(return_value=pw)
    cm.__exit__ = MagicMock(return_value=None)
    cm.__aenter__ = AsyncMock(return_value=pw)
    cm.__aexit__ = AsyncMock(return_value=None)
    factory = MagicMock(return_value=cm)
    return factory


# ==============================================================================
# SECTION 1: FEATURE 1 TESTS — `li_at` COOKIE AUTH & RESOLUTION
# ==============================================================================

def test_cookie_loader_from_env(monkeypatch):
    """Tier 1: When LINKEDIN_LI_AT is set in env, loads correctly with valid Playwright schema."""
    monkeypatch.setenv("LINKEDIN_LI_AT", "AQEDAQ_test_env_token_123")
    manager = LinkedInSessionManager(cookie_path="nonexistent_file.json")
    
    token = manager.get_li_at_cookie()
    assert token == "AQEDAQ_test_env_token_123"
    
    cookie_dict = manager.get_cookie_dict()
    assert cookie_dict["name"] == "li_at"
    assert cookie_dict["value"] == "AQEDAQ_test_env_token_123"
    assert cookie_dict["domain"] == ".linkedin.com"
    assert cookie_dict["path"] == "/"
    assert cookie_dict["httpOnly"] is True
    assert cookie_dict["secure"] is True
    assert cookie_dict["sameSite"] == "Lax"

def test_cookie_loader_from_file_list(tmp_path, monkeypatch):
    """Tier 1: When env is absent, loads token from session_cookies.json (Playwright list format)."""
    monkeypatch.delenv("LINKEDIN_LI_AT", raising=False)
    monkeypatch.delenv("LI_AT", raising=False)
    
    cookie_file = tmp_path / "session_cookies.json"
    cookie_file.write_text(json.dumps([
        {"name": "JSESSIONID", "value": "ajax:123"},
        {"name": "li_at", "value": "AQEDAQ_file_token_list_456", "domain": ".linkedin.com"}
    ]))
    
    with patch("naukari_pro.modules.linkedin_session.WORKSPACE_DIR", str(tmp_path)):
        manager = LinkedInSessionManager(cookie_path=str(cookie_file))
        assert manager.get_li_at_cookie() == "AQEDAQ_file_token_list_456"

def test_cookie_loader_from_file_dict(tmp_path, monkeypatch):
    """Tier 1: Loads token from session_cookies.json (key-value dictionary format)."""
    monkeypatch.delenv("LINKEDIN_LI_AT", raising=False)
    monkeypatch.delenv("LI_AT", raising=False)
    
    cookie_file = tmp_path / "session_cookies.json"
    cookie_file.write_text(json.dumps({
        "li_at": "AQEDAQ_file_token_dict_789",
        "JSESSIONID": "ajax:456"
    }))
    
    with patch("naukari_pro.modules.linkedin_session.WORKSPACE_DIR", str(tmp_path)):
        manager = LinkedInSessionManager(cookie_path=str(cookie_file))
        assert manager.get_li_at_cookie() == "AQEDAQ_file_token_dict_789"

def test_cookie_loader_priority_order(tmp_path, monkeypatch):
    """Tier 1: Environment variable has absolute precedence over disk files."""
    monkeypatch.setenv("LINKEDIN_LI_AT", "ENV_TOKEN_WINNER")
    
    cookie_file = tmp_path / "session_cookies.json"
    cookie_file.write_text(json.dumps({"li_at": "FILE_TOKEN_LOSER"}))
    
    manager = LinkedInSessionManager(cookie_path=str(cookie_file))
    assert manager.get_li_at_cookie() == "ENV_TOKEN_WINNER"

def test_cookie_missing_raises(tmp_path, monkeypatch):
    """Tier 2 Boundary: When no cookie source is available, raises LinkedInCookieMissingError."""
    monkeypatch.delenv("LINKEDIN_LI_AT", raising=False)
    monkeypatch.delenv("LI_AT", raising=False)
    
    with patch("naukari_pro.modules.linkedin_session.WORKSPACE_DIR", str(tmp_path)):
        manager = LinkedInSessionManager(cookie_path=str(tmp_path / "missing.json"))
        with pytest.raises(LinkedInCookieMissingError) as exc_info:
            manager.get_li_at_cookie()
        
        assert "LINKEDIN_LI_AT" in str(exc_info.value)
        assert "session_cookies.json" in str(exc_info.value)

def test_cookie_empty_string_raises(tmp_path, monkeypatch):
    """Tier 2 Boundary: Empty or whitespace-only env token is rejected as missing."""
    monkeypatch.setenv("LINKEDIN_LI_AT", "    ")
    with patch("naukari_pro.modules.linkedin_session.WORKSPACE_DIR", str(tmp_path)):
        manager = LinkedInSessionManager(cookie_path="nonexistent.json")
        with pytest.raises(LinkedInCookieMissingError):
            manager.get_li_at_cookie()

def test_cookie_token_sanitization(monkeypatch):
    """Tier 2 Boundary: Strips surrounding quotes and whitespace from tokens."""
    monkeypatch.setenv("LINKEDIN_LI_AT", '  "AQEDAQ_quoted_token_abc"  \n')
    manager = LinkedInSessionManager(cookie_path="nonexistent.json")
    assert manager.get_li_at_cookie() == "AQEDAQ_quoted_token_abc"

def test_cookie_corrupted_json_handling(tmp_path, monkeypatch):
    """Tier 2 Edge: When session_cookies.json contains malformed JSON, falls back gracefully."""
    monkeypatch.delenv("LINKEDIN_LI_AT", raising=False)
    monkeypatch.delenv("LI_AT", raising=False)
    
    corrupted_file = tmp_path / "session_cookies.json"
    corrupted_file.write_text("{corrupted json, not valid")
    
    with patch("naukari_pro.modules.linkedin_session.WORKSPACE_DIR", str(tmp_path)):
        manager = LinkedInSessionManager(cookie_path=str(corrupted_file))
        with pytest.raises(LinkedInCookieMissingError):
            manager.get_li_at_cookie()

def test_cookie_masking_security():
    """Tier 2 Security: Validates cookie token masking for terminal and logs."""
    assert mask_cookie_token("AQEDARuXYZ123456789") == "AQED...6789 (19 chars, masked)"
    assert mask_cookie_token(None) == "None"
    assert mask_cookie_token("") == "None"
    assert mask_cookie_token("   ") == "None"
    assert mask_cookie_token("1234") == "****"

def test_clean_avatar_name():
    """Tier 2 Parsing: Strips photo prefixes from nav avatar alt text."""
    assert clean_avatar_name("Photo of Lavkush Kumar") == "Lavkush Kumar"
    assert clean_avatar_name("Profile photo of Lavkush Kumar") == "Lavkush Kumar"
    assert clean_avatar_name("Lavkush Kumar") == "Lavkush Kumar"
    assert clean_avatar_name("Photo of ") is None
    assert clean_avatar_name("Profile photo of ") is None
    assert clean_avatar_name("Photo of") is None
    assert clean_avatar_name("Profile photo of") is None
    assert clean_avatar_name(None) is None
    assert clean_avatar_name("") is None
    assert clean_avatar_name("   ") is None

def test_cookie_token_quoted_internal_whitespace(monkeypatch):
    """Validates tokens with internal quoted spaces are properly trimmed."""
    monkeypatch.setenv("LINKEDIN_LI_AT", '\'  "  AQEDAQ_quoted_token_abc  "  \'')
    manager = LinkedInSessionManager(cookie_path="nonexistent.json")
    assert manager.get_li_at_cookie() == "AQEDAQ_quoted_token_abc"

def test_cookie_env_inline_comment_stripping(tmp_path, monkeypatch):
    """Validates inline comments in .env files do not poison the token."""
    monkeypatch.delenv("LINKEDIN_LI_AT", raising=False)
    monkeypatch.delenv("LI_AT", raising=False)
    
    env_file = tmp_path / ".env"
    env_file.write_text('LINKEDIN_LI_AT="AQEDAQ_comment_token_123" # Primary LinkedIn session cookie\n')
    
    with patch("naukari_pro.modules.linkedin_session.WORKSPACE_DIR", str(tmp_path)):
        manager = LinkedInSessionManager(cookie_path="missing.json")
        assert manager.get_li_at_cookie() == "AQEDAQ_comment_token_123"


# ==============================================================================
# SECTION 2: FEATURE 2 TESTS — STEALTH CONTEXT & ANTI-DETECTION
# ==============================================================================

@run_async
async def test_stealth_context_initialization(tmp_path, monkeypatch):
    """Tier 1: Context is created with desktop UA, viewport, STEALTH_JS, and li_at cookie."""
    monkeypatch.setenv("LINKEDIN_LI_AT", "AQEDAQ_stealth_token")
    session_dir = str(tmp_path / "session")
    manager = LinkedInSessionManager(session_dir=session_dir)
    
    mock_ctx = make_mock_context()
    mock_pw = make_mock_playwright(mock_ctx)
    
    context = await manager.create_stealth_context(mock_pw, headless=True)
    assert context == mock_ctx
    
    # Assert launch_persistent_context parameters
    launch_kwargs = mock_pw.chromium.launch_persistent_context.call_args[1]
    assert launch_kwargs["user_data_dir"] == session_dir
    assert launch_kwargs["headless"] is True
    assert launch_kwargs["viewport"] == {"width": 1920, "height": 1080}
    assert launch_kwargs["locale"] == "en-US"
    assert launch_kwargs["timezone_id"] == "Asia/Kolkata"
    assert any(os_id in launch_kwargs["user_agent"] for os_id in ["Windows NT 10.0", "Macintosh", "Linux"])
    
    # Assert stealth init script and cookies added
    mock_ctx.add_init_script.assert_awaited_once_with(STEALTH_JS)
    mock_ctx.add_cookies.assert_awaited_once()
    added_cookies = mock_ctx.add_cookies.call_args[0][0]
    assert any(c.get("name") == "li_at" and c.get("value") == "AQEDAQ_stealth_token" for c in added_cookies)
    await context.close()

def test_stealth_js_contents():
    """Tier 1: STEALTH_JS contains necessary anti-fingerprinting patches."""
    assert "navigator, 'webdriver'" in STEALTH_JS
    assert "window.chrome" in STEALTH_JS
    assert "Intel Inc." in STEALTH_JS
    assert "Intel Iris OpenGL Engine" in STEALTH_JS

@run_async
async def test_stealth_context_headless_flag(tmp_path, monkeypatch):
    """Tier 1: Verifies headless=False flag is propagated to Chromium."""
    monkeypatch.setenv("LINKEDIN_LI_AT", "TOKEN")
    session_dir = str(tmp_path / "session")
    manager = LinkedInSessionManager(session_dir=session_dir)
    
    mock_ctx = make_mock_context()
    mock_pw = make_mock_playwright(mock_ctx)
    
    context = await manager.create_stealth_context(mock_pw, headless=False)
    launch_kwargs = mock_pw.chromium.launch_persistent_context.call_args[1]
    # On macOS or Linux with display, headless is False
    if not (sys.platform.startswith("linux") and not os.environ.get("DISPLAY")):
        assert launch_kwargs["headless"] is False
    await context.close()

@run_async
async def test_stealth_context_locking(tmp_path, monkeypatch):
    """Tier 1: Verifies process lock acquisition on .linkedin_session.lock."""
    monkeypatch.setenv("LINKEDIN_LI_AT", "TOKEN")
    session_dir = str(tmp_path / "session")
    manager = LinkedInSessionManager(session_dir=session_dir)
    
    mock_ctx = make_mock_context()
    mock_pw = make_mock_playwright(mock_ctx)
    
    with patch("naukari_pro.modules.linkedin_session.SessionLock") as MockLock:
        lock_instance = MagicMock()
        MockLock.return_value = lock_instance
        
        context = await manager.create_stealth_context(mock_pw, headless=True)
        
        MockLock.assert_called_once()
        lock_instance.acquire.assert_called_once()
        mock_ctx.on.assert_called_once()
        call_args = mock_ctx.on.call_args[0]
        assert call_args[0] == "close"
        # Invoke closure with mock argument to verify argument-tolerant release
        call_args[1](mock_ctx)
        lock_instance.release.assert_called()
        await context.close()

@run_async
async def test_create_stealth_context_lock_released_on_cookie_error(tmp_path, monkeypatch):
    """Tier 2 Concurrency: SessionLock must be released if cookie discovery fails."""
    monkeypatch.delenv("LINKEDIN_LI_AT", raising=False)
    monkeypatch.delenv("LI_AT", raising=False)
    session_dir = str(tmp_path / "session")
    manager = LinkedInSessionManager(session_dir=session_dir, cookie_path=str(tmp_path / "nonexistent.json"))
    
    mock_pw = make_mock_playwright()
    with patch("naukari_pro.modules.linkedin_session.WORKSPACE_DIR", str(tmp_path)):
        with pytest.raises(LinkedInCookieMissingError):
            await manager.create_stealth_context(mock_pw)
    
    lock = manager.get_lock()
    assert not lock._acquired, "SessionLock must not remain acquired when cookie is missing"

@run_async
async def test_create_stealth_context_lock_released_on_launch_failure(tmp_path, monkeypatch):
    """Tier 2 Concurrency: SessionLock must be released if browser launch raises exception."""
    monkeypatch.setenv("LINKEDIN_LI_AT", "TOKEN")
    session_dir = str(tmp_path / "session")
    manager = LinkedInSessionManager(session_dir=session_dir)
    
    mock_pw = MagicMock()
    mock_pw.chromium.launch_persistent_context = AsyncMock(side_effect=RuntimeError("Browser launch failure"))
    
    with pytest.raises(RuntimeError):
        await manager.create_stealth_context(mock_pw)
    
    lock = manager.get_lock()
    assert not lock._acquired, "SessionLock must not remain acquired when browser launch fails"

@run_async
async def test_stealth_context_stale_lock_cleanup(tmp_path, monkeypatch):
    """Tier 2: Verifies cleanup_stale_chromium_locks is executed."""
    monkeypatch.setenv("LINKEDIN_LI_AT", "TOKEN")
    session_dir = str(tmp_path / "session")
    manager = LinkedInSessionManager(session_dir=session_dir)
    
    mock_ctx = make_mock_context()
    mock_pw = make_mock_playwright(mock_ctx)
    
    with patch("naukari_pro.modules.linkedin_session.cleanup_stale_chromium_locks") as mock_cleanup:
        context = await manager.create_stealth_context(mock_pw, headless=True)
        mock_cleanup.assert_called_once_with(manager.session_dir)
        await context.close()

@run_async
async def test_stealth_context_linux_no_display(tmp_path, monkeypatch):
    """Tier 2 Edge: Linux with DISPLAY unset forces headless=True."""
    monkeypatch.setenv("LINKEDIN_LI_AT", "TOKEN")
    monkeypatch.setattr(sys, "platform", "linux")
    monkeypatch.delenv("DISPLAY", raising=False)
    
    session_dir = str(tmp_path / "session")
    manager = LinkedInSessionManager(session_dir=session_dir)
    
    mock_ctx = make_mock_context()
    mock_pw = make_mock_playwright(mock_ctx)
    
    context = await manager.create_stealth_context(mock_pw, headless=False)
    launch_kwargs = mock_pw.chromium.launch_persistent_context.call_args[1]
    assert launch_kwargs["headless"] is True
    await context.close()


# ==============================================================================
# SECTION 3: FEATURE 3 TESTS — CHECKPOINT & EXPIRY VERIFICATION
# ==============================================================================

@run_async
async def test_verify_session_valid(monkeypatch):
    """Tier 1: Authenticated feed DOM correctly extracts candidate identity."""
    monkeypatch.setenv("LINKEDIN_LI_AT", "TOKEN")
    manager = LinkedInSessionManager()
    
    mock_name_loc = make_mock_locator(count=1, visible=True, text="Lavkush Kumar")
    mock_headline_loc = make_mock_locator(count=1, visible=True, text="Azure Data Engineer | 2.5 YOE")
    mock_nav_loc = make_mock_locator(count=1, visible=True, text="")
    
    locators = {
        "div.feed-identity-module__actor-meta a span": mock_name_loc,
        "div.feed-identity-module .feed-identity-module__headline": mock_headline_loc,
        "nav.global-nav": mock_nav_loc
    }
    
    mock_page = make_mock_page(
        current_url="https://www.linkedin.com/feed/",
        locators=locators
    )
    
    result = await manager.verify_session(mock_page)
    assert result["status"] == "VALID"
    assert result["authenticated"] is True
    assert result["candidate_name"] == "Lavkush Kumar"
    assert result["headline"] == "Azure Data Engineer | 2.5 YOE"

@run_async
async def test_verify_session_valid_avatar_fallback(monkeypatch):
    """Tier 1: Feed avatar alt fallback successfully extracts candidate name."""
    monkeypatch.setenv("LINKEDIN_LI_AT", "TOKEN")
    manager = LinkedInSessionManager()
    
    avatar_loc = make_mock_locator(count=1, visible=True, attr_map={"alt": "Photo of Lavkush Kumar"})
    nav_loc = make_mock_locator(count=1, visible=True, text="")
    
    locators = {
        "img.global-nav__me-photo": avatar_loc,
        "nav.global-nav": nav_loc
    }
    
    mock_page = make_mock_page(
        current_url="https://www.linkedin.com/feed/",
        locators=locators
    )
    
    result = await manager.verify_session(mock_page)
    assert result["status"] == "VALID"
    assert result["candidate_name"] == "Lavkush Kumar"

@run_async
async def test_verify_session_expired_authwall(monkeypatch):
    """Tier 1: Redirect to /authwall is recognized as EXPIRED."""
    monkeypatch.setenv("LINKEDIN_LI_AT", "TOKEN")
    manager = LinkedInSessionManager()
    
    mock_page = make_mock_page(current_url="https://www.linkedin.com/authwall?trk=rip")
    result = await manager.verify_session(mock_page)
    
    assert result["status"] == "EXPIRED"
    assert result["authenticated"] is False
    assert "authwall" in result["detail"].lower()

@run_async
async def test_verify_session_expired_login_form(monkeypatch):
    """Tier 1: Login form visible in DOM is recognized as EXPIRED."""
    monkeypatch.setenv("LINKEDIN_LI_AT", "TOKEN")
    manager = LinkedInSessionManager()
    
    login_loc = make_mock_locator(count=1, visible=True, text="")
    mock_page = make_mock_page(
        current_url="https://www.linkedin.com/feed/",
        locators={"input#session_key": login_loc}
    )
    
    result = await manager.verify_session(mock_page)
    assert result["status"] == "EXPIRED"
    assert result["authenticated"] is False

@run_async
async def test_verify_session_checkpoint_url(monkeypatch):
    """Tier 1: URL matching /checkpoint/challenge is recognized as CHECKPOINT."""
    monkeypatch.setenv("LINKEDIN_LI_AT", "TOKEN")
    manager = LinkedInSessionManager()
    
    mock_page = make_mock_page(current_url="https://www.linkedin.com/checkpoint/challenge/12345")
    result = await manager.verify_session(mock_page)
    
    assert result["status"] == "CHECKPOINT"
    assert result["authenticated"] is False
    assert result["checkpoint_detected"] is True

@run_async
async def test_verify_session_checkpoint_captcha_dom(monkeypatch):
    """Tier 2: DOM selector input#captcha detected on feed URL is recognized as CHECKPOINT."""
    monkeypatch.setenv("LINKEDIN_LI_AT", "TOKEN")
    manager = LinkedInSessionManager()
    
    captcha_loc = make_mock_locator(count=1, visible=True, text="")
    mock_page = make_mock_page(
        current_url="https://www.linkedin.com/feed/",
        locators={"input#captcha": captcha_loc}
    )
    
    result = await manager.verify_session(mock_page)
    assert result["status"] == "CHECKPOINT"
    assert result["checkpoint_detected"] is True

@run_async
async def test_verify_session_rate_limited(monkeypatch):
    """Tier 2: HTTP 999 or request-denied recognized as RATE_LIMITED."""
    monkeypatch.setenv("LINKEDIN_LI_AT", "TOKEN")
    manager = LinkedInSessionManager()
    
    mock_page = make_mock_page(
        current_url="https://www.linkedin.com/uas/request-denied",
        body_text="HTTP 999 Request Denied"
    )
    result = await manager.verify_session(mock_page)
    assert result["status"] == "RATE_LIMITED"

@run_async
async def test_verify_session_blank_page_rejected(monkeypatch):
    """Negative Verification: Blank page (about:blank) must return UNKNOWN, never VALID."""
    monkeypatch.setenv("LINKEDIN_LI_AT", "TOKEN")
    manager = LinkedInSessionManager()
    
    mock_page = make_mock_page(current_url="about:blank")
    result = await manager.verify_session(mock_page)
    
    assert result["status"] == "UNKNOWN"
    assert result["authenticated"] is False
    assert result["candidate_name"] is None

@run_async
async def test_verify_session_arbitrary_domain_rejected(monkeypatch):
    """Negative Verification: Non-LinkedIn domain (google.com) must return UNKNOWN, never VALID."""
    monkeypatch.setenv("LINKEDIN_LI_AT", "TOKEN")
    manager = LinkedInSessionManager()
    
    mock_page = make_mock_page(current_url="https://www.google.com", body_text="Google Search")
    result = await manager.verify_session(mock_page)
    
    assert result["status"] == "UNKNOWN"
    assert result["authenticated"] is False
    assert result["candidate_name"] is None

@pytest.mark.parametrize("checkpoint_url", [
    "https://www.linkedin.com/checkpoint/rp/request-password-reset",
    "https://www.linkedin.com/checkpoint/security-check",
    "https://www.linkedin.com/checkpoint/identity-verification",
    "https://www.linkedin.com/checkpoint/id-verification",
    "https://www.linkedin.com/checkpoint/phone-pin",
    "https://www.linkedin.com/uas/captcha",
    "https://www.linkedin.com/uas/consumer-captcha"
])
@run_async
async def test_verify_session_unlisted_checkpoint_rejected(checkpoint_url, monkeypatch):
    """Negative Verification: All /checkpoint/ and /uas/captcha URLs must be recognized as CHECKPOINT."""
    monkeypatch.setenv("LINKEDIN_LI_AT", "TOKEN")
    manager = LinkedInSessionManager()
    
    mock_page = make_mock_page(current_url=checkpoint_url)
    result = await manager.verify_session(mock_page)
    
    assert result["status"] == "CHECKPOINT"
    assert result["checkpoint_detected"] is True
    assert result["authenticated"] is False

@run_async
async def test_verify_session_navigation_failure_rejected(monkeypatch):
    """Negative Verification: Navigation timeout on page.goto must return unauthenticated status."""
    monkeypatch.setenv("LINKEDIN_LI_AT", "TOKEN")
    manager = LinkedInSessionManager()
    
    mock_page = make_mock_page(current_url="about:blank")
    mock_page.goto = AsyncMock(side_effect=TimeoutError("Navigation timeout of 30000ms exceeded"))
    
    result = await manager.verify_session(mock_page)
    assert result["status"] != "VALID"
    assert result["authenticated"] is False
    assert "Navigation failed" in result["detail"]

@run_async
async def test_verify_session_navigation_failure_raises(monkeypatch):
    """Negative Verification: Navigation failure raises LinkedInSessionError when raise_for_status=True."""
    monkeypatch.setenv("LINKEDIN_LI_AT", "TOKEN")
    manager = LinkedInSessionManager()
    
    mock_page = make_mock_page(current_url="about:blank")
    mock_page.goto = AsyncMock(side_effect=TimeoutError("Navigation timeout"))
    
    with pytest.raises(LinkedInSessionError):
        await manager.verify_session(mock_page, raise_for_status=True)

@run_async
async def test_verify_session_unauthenticated_generic_page_rejected(monkeypatch):
    """Negative Verification: Generic unauthenticated pages without nav elements return EXPIRED."""
    monkeypatch.setenv("LINKEDIN_LI_AT", "TOKEN")
    manager = LinkedInSessionManager()
    
    mock_page = make_mock_page(current_url="https://www.linkedin.com/legal/privacy-policy")
    result = await manager.verify_session(mock_page)
    
    assert result["status"] == "EXPIRED"
    assert result["authenticated"] is False
    assert result["candidate_name"] is None

@run_async
async def test_verify_session_unauthenticated_scaffold_rejected(monkeypatch):
    """Negative Verification: Unauthenticated layout with only scaffold-layout__main is rejected as EXPIRED."""
    monkeypatch.setenv("LINKEDIN_LI_AT", "TOKEN")
    manager = LinkedInSessionManager()
    
    loc = make_mock_locator(count=1, visible=True, text="Logged out page")
    mock_page = make_mock_page(
        "https://www.linkedin.com/feed/",
        locators={"main.scaffold-layout__main": loc}
    )
    result = await manager.verify_session(mock_page)
    
    assert result["status"] == "EXPIRED"
    assert result["authenticated"] is False
    assert result["candidate_name"] is None

@run_async
async def test_verify_session_port_443_url_accepted(monkeypatch):
    """Verification: URL with explicit :443 port passes is_linkedin_domain check."""
    monkeypatch.setenv("LINKEDIN_LI_AT", "TOKEN")
    manager = LinkedInSessionManager()
    
    mock_name_loc = make_mock_locator(count=1, visible=True, text="Lavkush Kumar")
    mock_nav_loc = make_mock_locator(count=1, visible=True, text="")
    mock_page = make_mock_page(
        "https://www.linkedin.com:443/feed/",
        locators={
            "div.feed-identity-module__actor-meta a span": mock_name_loc,
            "nav.global-nav": mock_nav_loc,
        }
    )
    result = await manager.verify_session(mock_page)
    assert result["status"] == "VALID"
    assert result["authenticated"] is True
    assert result["candidate_name"] == "Lavkush Kumar"


# ==============================================================================
# SECTION 4: FEATURE 4 TESTS — CLI DIAGNOSTIC `linkedin-verify`
# ==============================================================================

def test_cli_linkedin_verify_valid(capsys, monkeypatch):
    """Tier 1: CLI invocation with valid session returns exit code 0 and reports VALID."""
    monkeypatch.setenv("LINKEDIN_LI_AT", "AQEDAQ_valid_cli_token_123")
    
    name_loc = make_mock_locator(count=1, visible=True, text="Lavkush Kumar")
    nav_loc = make_mock_locator(count=1, visible=True, text="")
    mock_page = make_mock_page(
        current_url="https://www.linkedin.com/feed/",
        locators={
            "div.feed-identity-module__actor-meta a span": name_loc,
            "nav.global-nav": nav_loc,
            "div.feed-identity-module": name_loc
        }
    )
    pw_mock = make_playwright_cm_mock(page=mock_page)
    
    with patch("naukari_pro.cli.async_playwright", pw_mock), \
         patch("naukari_pro.cli.sync_playwright", pw_mock):
        
        exit_code = handle_linkedin_verify(headless=True)
        assert exit_code == 0
        
        captured = capsys.readouterr().out
        assert "Session Status:         VALID" in captured
        assert "Authenticated Identity: Lavkush Kumar" in captured
        assert "Cookie Present:         YES" in captured

def test_cli_linkedin_verify_expired(capsys, monkeypatch):
    """Tier 1: CLI invocation with expired session returns exit code 1 and remediation."""
    monkeypatch.setenv("LINKEDIN_LI_AT", "AQEDAQ_expired_token")
    
    mock_page = make_mock_page(current_url="https://www.linkedin.com/authwall")
    pw_mock = make_playwright_cm_mock(page=mock_page)
    
    with patch("naukari_pro.cli.async_playwright", pw_mock), \
         patch("naukari_pro.cli.sync_playwright", pw_mock):
        
        exit_code = handle_linkedin_verify(headless=True)
        assert exit_code == 1
        
        captured = capsys.readouterr().out
        assert "Session Status:         EXPIRED" in captured
        assert "Remediation Instructions" in captured

def test_cli_linkedin_verify_checkpoint(capsys, monkeypatch):
    """Tier 1: CLI invocation with checkpoint challenge returns exit code 2."""
    monkeypatch.setenv("LINKEDIN_LI_AT", "AQEDAQ_checkpoint_token")
    
    mock_page = make_mock_page(current_url="https://www.linkedin.com/checkpoint/challenge/12345")
    pw_mock = make_playwright_cm_mock(page=mock_page)
    
    with patch("naukari_pro.cli.async_playwright", pw_mock), \
         patch("naukari_pro.cli.sync_playwright", pw_mock):
        
        exit_code = handle_linkedin_verify(headless=True)
        assert exit_code == 2
        
        captured = capsys.readouterr().out
        assert "Session Status:         CHECKPOINT" in captured
        assert "--headed" in captured

def test_cli_linkedin_verify_real_async_timeout(capsys, monkeypatch):
    """Verifies CLI execution on navigation timeout yields exit code 1 (NOT 0)."""
    monkeypatch.setenv("LINKEDIN_LI_AT", "AQEDAQ_timeout_token")
    
    mock_page = make_mock_page(current_url="about:blank")
    mock_page.goto = AsyncMock(side_effect=TimeoutError("Navigation timeout of 30000ms exceeded"))
    pw_mock = make_playwright_cm_mock(page=mock_page)
    
    with patch("naukari_pro.cli.async_playwright", pw_mock), \
         patch("naukari_pro.cli.sync_playwright", pw_mock):
        exit_code = handle_linkedin_verify(headless=True)
        assert exit_code == 1
        
        captured = capsys.readouterr().out
        assert "Session Status:         EXPIRED" in captured

def test_cli_linkedin_verify_missing_cookie(capsys, monkeypatch):
    """Tier 1: CLI invocation without cookie returns exit code 1 and instructions."""
    monkeypatch.delenv("LINKEDIN_LI_AT", raising=False)
    monkeypatch.delenv("LI_AT", raising=False)
    
    with patch("naukari_pro.modules.linkedin_session.LinkedInSessionManager.get_li_at_cookie", side_effect=LinkedInCookieMissingError):
        exit_code = handle_linkedin_verify(headless=True)
        assert exit_code == 1
        
        captured = capsys.readouterr().out
        assert "Cookie Present:         NO" in captured
        assert "COOKIE_MISSING" in captured

def test_cli_subparser_arguments():
    """Tier 2: Verifies argument parsing for linkedin-verify subcommand."""
    from naukari_pro.cli import main
    # Create test parser matching cli.py structure
    common_parser = argparse.ArgumentParser(add_help=False)
    common_parser.add_argument("--headless", action="store_true")
    parser = argparse.ArgumentParser(parents=[common_parser])
    subparsers = parser.add_subparsers(dest="command")
    verify_parser = subparsers.add_parser("linkedin-verify", parents=[common_parser])
    verify_parser.add_argument("--headed", action="store_true", default=False)
    
    args = parser.parse_args(["linkedin-verify"])
    assert args.command == "linkedin-verify"
    assert args.headed is False
    
    args_headed = parser.parse_args(["linkedin-verify", "--headed"])
    assert args_headed.headed is True


# ==============================================================================
# SECTION 5: TIER 3 & 4 INTEGRATION CONTRACTS & SCENARIOS
# ==============================================================================

def test_tier3_cookie_to_stealth_contract(monkeypatch):
    """Tier 3: Pairwise contract between cookie resolution and Playwright cookie formatting."""
    monkeypatch.setenv("LINKEDIN_LI_AT", "AQEDAQ_pair_token")
    manager = LinkedInSessionManager()
    
    cookie_dict = manager.get_cookie_dict()
    assert isinstance(cookie_dict, dict)
    assert cookie_dict["name"] == "li_at"
    assert cookie_dict["value"] == "AQEDAQ_pair_token"
    assert cookie_dict["domain"] == ".linkedin.com"
    assert cookie_dict["path"] == "/"

@run_async
async def test_tier4_full_session_lifecycle_scenario(tmp_path, monkeypatch):
    """Tier 4: Realistic end-to-end offline lifecycle simulation."""
    monkeypatch.setenv("LINKEDIN_LI_AT", "AQEDAQ_e2e_scenario_token")
    session_dir = str(tmp_path / "e2e_session")
    
    manager = LinkedInSessionManager(session_dir=session_dir)
    assert manager.get_li_at_cookie() == "AQEDAQ_e2e_scenario_token"
    
    mock_ctx = make_mock_context()
    mock_pw = make_mock_playwright(mock_ctx)
    
    ctx = await manager.create_stealth_context(mock_pw, headless=True)
    assert ctx == mock_ctx
    
    # Simulate page verification
    mock_page = mock_ctx.pages[0]
    mock_page.url = "https://www.linkedin.com/feed/"
    name_loc = make_mock_locator(count=1, visible=True, text="Lavkush Kumar")
    def locator_dispatch(sel):
        if "actor-meta" in sel or "global-nav" in sel or "identity" in sel:
            return name_loc
        return make_mock_locator(count=0, visible=False)
    mock_page.locator = MagicMock(side_effect=locator_dispatch)
    
    report = await manager.verify_session(mock_page)
    assert report["status"] == "VALID"
    assert report["candidate_name"] == "Lavkush Kumar"
    
    # Teardown
    await ctx.close()
    mock_ctx.close.assert_awaited_once()


# ==============================================================================
# SECTION 6: MILESTONE 2 — URL BUILDER & KEYWORD MATRIX TESTS (FEATURES 5 & 6)
# ==============================================================================

import urllib.parse

from naukari_pro.modules.linkedin_search import (
    TECH_HUBS,
    CANONICAL_LOCATIONS,
    LOCATION_ALIASES,
    DEFAULT_KEYWORDS,
    TARGET_KEYWORDS_TIER1,
    TARGET_KEYWORDS_TIER2,
    ALL_TARGET_KEYWORDS,
    LinkedInSearchUrlBuilder,
    JobDeduplicator,
    LinkedInJobDeduplicator,
    JobMatchScorer,
    PrecisionMatchScorer,
    MatchScoreResult,
    LinkedInJobCardParser,
    LinkedInSearchEngine,
    LinkedInSearchError,
    LinkedInInvalidLocationError,
)


def test_url_builder_all_7_tech_hubs():
    """Validates URL generation for all 7 Tier-1 Indian tech hubs and Remote India."""
    builder = LinkedInSearchUrlBuilder()
    for hub, expected_geoid in TECH_HUBS.items():
        url = builder.build("Azure Data Engineer", location=hub)
        assert f"geoId={expected_geoid}" in url
        assert "keywords=Azure%20Data%20Engineer" in url
        assert "f_AL=true" in url
        assert "sortBy=DD" in url
        if hub == "Remote":
            assert "f_WT=2" in url
            assert "f_CF=f_WRA" in url
            assert "location=India" in url
        else:
            canonical_loc = urllib.parse.quote(CANONICAL_LOCATIONS[hub], safe=",")
            assert canonical_loc in url


def test_url_builder_easy_apply_filter():
    """Validates f_AL=true Easy Apply filter toggling."""
    url_easy = LinkedInSearchUrlBuilder.build("Databricks", easy_apply_only=True)
    assert "f_AL=true" in url_easy

    url_all = LinkedInSearchUrlBuilder.build("Databricks", easy_apply_only=False)
    assert "f_AL" not in url_all


def test_url_builder_freshness_filters():
    """Validates freshness filter mappings for 24h, 7d, integer seconds, and None."""
    # 24 hours (default)
    url_24h = LinkedInSearchUrlBuilder.build("PySpark", freshness_sec=86400)
    assert "f_TPR=r86400" in url_24h

    # String aliases
    url_alias_24h = LinkedInSearchUrlBuilder.build("PySpark", freshness_sec="24h")
    assert "f_TPR=r86400" in url_alias_24h

    url_7d = LinkedInSearchUrlBuilder.build("PySpark", freshness_sec="7d")
    assert "f_TPR=r604800" in url_7d

    url_past_week = LinkedInSearchUrlBuilder.build("PySpark", freshness_sec="past_week")
    assert "f_TPR=r604800" in url_past_week

    # None / 0 omits f_TPR
    url_any = LinkedInSearchUrlBuilder.build("PySpark", freshness_sec=None)
    assert "f_TPR" not in url_any

    url_zero = LinkedInSearchUrlBuilder.build("PySpark", freshness_sec=0)
    assert "f_TPR" not in url_zero


def test_url_builder_experience_level_filter():
    """Validates f_E=2,3 experience filter for Entry-level and Associate (2.5 YOE)."""
    url_default = LinkedInSearchUrlBuilder.build("Delta Lake")
    assert "f_E=2,3" in url_default

    url_custom = LinkedInSearchUrlBuilder.build("Delta Lake", experience_levels="2")
    assert "f_E=2" in url_custom

    url_none = LinkedInSearchUrlBuilder.build("Delta Lake", experience_levels=None)
    assert "f_E" not in url_none


def test_url_builder_keyword_queries_matrix():
    """Validates that all default keywords can be converted into valid search URLs."""
    for kw in DEFAULT_KEYWORDS:
        url = LinkedInSearchUrlBuilder.build(kw, location="Bengaluru")
        parsed = LinkedInSearchUrlBuilder.parse(url)
        assert parsed["keywords"] == kw
        assert parsed["geoId"] == TECH_HUBS["Bengaluru"]
        assert parsed["f_AL"] == "true"


def test_url_builder_unknown_hub_handling():
    """Validates strict vs non-strict behavior for unsupported locations."""
    # Non-strict mode: falls back to custom location
    url_freeform = LinkedInSearchUrlBuilder.build("Azure", location="Chandigarh", strict_location=False)
    assert "location=Chandigarh" in url_freeform
    assert "geoId" not in url_freeform

    # Strict mode: raises LinkedInInvalidLocationError
    with pytest.raises(LinkedInInvalidLocationError):
        LinkedInSearchUrlBuilder.build("Azure", location="UnknownCityXYZ", strict_location=True)


def test_url_builder_special_characters_escaping():
    """Validates encoding for complex Boolean strings, symbols, and quotes."""
    boolean_kw = '"Azure Data Engineer" AND PySpark'
    url = LinkedInSearchUrlBuilder.build(boolean_kw, location="Hyderabad")
    parsed = LinkedInSearchUrlBuilder.parse(url)
    assert parsed["keywords"] == boolean_kw

    symbol_kw = "Python & SQL / CI/CD"
    url_sym = LinkedInSearchUrlBuilder.build(symbol_kw)
    parsed_sym = LinkedInSearchUrlBuilder.parse(url_sym)
    assert parsed_sym["keywords"] == symbol_kw


def test_url_builder_pagination_offset():
    """Validates start offset and 1-indexed page conversion."""
    # start offset
    url_start_0 = LinkedInSearchUrlBuilder.build("ADF", start=0)
    assert "start=" not in url_start_0

    url_start_25 = LinkedInSearchUrlBuilder.build("ADF", start=25)
    assert "start=25" in url_start_25

    url_start_50 = LinkedInSearchUrlBuilder.build("ADF", start=50)
    assert "start=50" in url_start_50

    # page conversion
    url_page_1 = LinkedInSearchUrlBuilder.build("ADF", page=1)
    assert "start=" not in url_page_1

    url_page_2 = LinkedInSearchUrlBuilder.build("ADF", page=2)
    assert "start=25" in url_page_2

    url_page_3 = LinkedInSearchUrlBuilder.build("ADF", page=3)
    assert "start=50" in url_page_3


def test_url_builder_remote_workplace_type():
    """Validates remote workplace filters for Pan-India remote searches."""
    for loc in ["Remote", "remote india", "India Remote", "pan-india"]:
        url = LinkedInSearchUrlBuilder.build("Databricks", location=loc)
        assert "f_WT=2" in url
        assert "f_CF=f_WRA" in url
        assert f"geoId={TECH_HUBS['Remote']}" in url


def test_url_builder_empty_keyword_rejected():
    """Validates ValueError is raised for blank keywords or invalid pagination."""
    with pytest.raises(ValueError, match="Keyword parameter must be a non-empty string"):
        LinkedInSearchUrlBuilder.build("")

    with pytest.raises(ValueError, match="Keyword parameter must be a non-empty string"):
        LinkedInSearchUrlBuilder.build("   ")

    with pytest.raises(ValueError, match="Page number must be >= 1"):
        LinkedInSearchUrlBuilder.build("Azure", page=0)

    with pytest.raises(ValueError, match="Pagination start offset must be >= 0"):
        LinkedInSearchUrlBuilder.build("Azure", start=-5)


def test_url_builder_parse_roundtrip():
    """Validates reverse-parsing of constructed LinkedIn search URL."""
    original_url = LinkedInSearchUrlBuilder.build(
        keyword="Medallion Architecture",
        location="Pune",
        freshness_sec=86400,
        easy_apply_only=True,
        start=25
    )
    parsed = LinkedInSearchUrlBuilder.parse(original_url)
    assert parsed["keywords"] == "Medallion Architecture"
    assert parsed["location"] == CANONICAL_LOCATIONS["Pune"]
    assert parsed["geoId"] == TECH_HUBS["Pune"]
    assert parsed["f_AL"] == "true"
    assert parsed["f_TPR"] == "r86400"
    assert parsed["f_E"] == "2,3"
    assert parsed["sortBy"] == "DD"
    assert parsed["start"] == "25"


# ==============================================================================
# SECTION 7: MILESTONE 2 — DOM JOB CARD PARSING & DEDUPLICATION (FEATURE 7)
# ==============================================================================

def make_mock_job_card(
    job_id: str = "4159823456",
    title: str = "Azure Data Engineer",
    company: str = "Microsoft",
    location: str = "Bengaluru, Karnataka, India",
    is_easy_apply: bool = True,
    is_applied: bool = False,
    time_posted: str = "1 day ago",
    raw_snippet: str = "Looking for Azure Data Engineer with Databricks, PySpark and ADF experience."
):
    """Creates a deterministic Playwright job card locator mock."""
    card = MagicMock()
    attr_map = {
        "data-job-id": job_id,
        "data-occludable-job-id": job_id,
        "data-entity-urn": f"urn:li:jobPosting:{job_id}"
    }
    card.get_attribute = AsyncMock(side_effect=lambda a: attr_map.get(a, None))
    card.inner_text = AsyncMock(return_value=f"{title}\n{company}\n{location}\n{raw_snippet}")
    card.scroll_into_view_if_needed = AsyncMock(return_value=None)

    title_loc = make_mock_locator(count=1, visible=True, text=title, attr_map={"href": f"/jobs/view/{job_id}/"})
    company_loc = make_mock_locator(count=1, visible=True, text=company)
    loc_loc = make_mock_locator(count=1, visible=True, text=location)
    easy_apply_loc = make_mock_locator(count=1 if is_easy_apply else 0, visible=is_easy_apply, text="Easy Apply" if is_easy_apply else "")
    applied_loc = make_mock_locator(count=1 if is_applied else 0, visible=is_applied, text="Applied" if is_applied else "")
    time_loc = make_mock_locator(count=1, visible=True, text=time_posted)

    def sub_locator(selector):
        sel_lower = selector.lower()
        if "applied" in sel_lower or "footer-job-state" in sel_lower:
            return applied_loc
        if "apply" in sel_lower or "linkedin-bug" in sel_lower:
            return easy_apply_loc
        if "title" in sel_lower or "container__link" in sel_lower:
            return title_loc
        if "company" in sel_lower or "subtitle" in sel_lower or "primary-description" in sel_lower:
            return company_loc
        if "metadata" in sel_lower or "caption" in sel_lower:
            return loc_loc
        if "time" in sel_lower or "date" in sel_lower:
            return time_loc
        return make_mock_locator(count=0, visible=False, text="")

    card.locator = MagicMock(side_effect=sub_locator)
    return card


@run_async
async def test_parse_job_cards_standard_page():
    """Validates card parsing on standard search results page."""
    card1 = make_mock_job_card(job_id="4159823451", title="Azure Data Engineer", company="Wipro")
    card2 = make_mock_job_card(job_id="4159823452", title="Databricks Engineer", company="Infosys")

    page = MagicMock()
    cards_list = [card1, card2]
    list_loc = MagicMock()
    list_loc.count = AsyncMock(return_value=len(cards_list))
    list_loc.nth = MagicMock(side_effect=lambda idx: cards_list[idx])

    def page_locator(sel):
        if "no-results" in sel or "empty-state" in sel:
            return make_mock_locator(count=0, visible=False)
        return list_loc

    page.locator = MagicMock(side_effect=page_locator)
    page.evaluate = AsyncMock(return_value=True)

    parser = LinkedInJobCardParser()
    results = await parser.parse_job_cards(page, scroll_to_load=False)
    assert len(results) == 2
    assert results[0]["job_id"] == "4159823451"
    assert results[0]["title"] == "Azure Data Engineer"
    assert results[0]["company"] == "Wipro"
    assert results[0]["is_easy_apply"] is True
    assert results[1]["job_id"] == "4159823452"


@run_async
async def test_parse_job_cards_extract_job_id_variants():
    """Validates 5-tier job ID extraction across data-job-id, occludable-id, URN, and href."""
    parser = LinkedInJobCardParser()

    # Variant 1: data-job-id attribute
    card_attr = MagicMock()
    card_attr.get_attribute = AsyncMock(side_effect=lambda a: "4159823451" if a == "data-job-id" else None)
    card_attr.locator = MagicMock(return_value=make_mock_locator(count=0))
    jid1 = await parser.extract_job_id(card_attr)
    assert jid1 == "4159823451"

    # Variant 2: data-occludable-job-id attribute
    card_occ = MagicMock()
    card_occ.get_attribute = AsyncMock(side_effect=lambda a: "4159823452" if a == "data-occludable-job-id" else None)
    card_occ.locator = MagicMock(return_value=make_mock_locator(count=0))
    jid2 = await parser.extract_job_id(card_occ)
    assert jid2 == "4159823452"

    # Variant 3: data-entity-urn attribute
    card_urn = MagicMock()
    card_urn.get_attribute = AsyncMock(side_effect=lambda a: "urn:li:jobPosting:4159823453" if a == "data-entity-urn" else None)
    card_urn.locator = MagicMock(return_value=make_mock_locator(count=0))
    jid3 = await parser.extract_job_id(card_urn)
    assert jid3 == "4159823453"

    # Variant 4: title href regex
    card_href = MagicMock()
    card_href.get_attribute = AsyncMock(return_value=None)
    link_loc = make_mock_locator(count=1, visible=True, attr_map={"href": "https://www.linkedin.com/jobs/view/4159823454/?trk=jobs"})
    card_href.locator = MagicMock(return_value=link_loc)
    jid4 = await parser.extract_job_id(card_href)
    assert jid4 == "4159823454"


@run_async
async def test_parse_job_cards_easy_apply_detection():
    """Validates discrimination between Easy Apply badges vs external apply."""
    parser = LinkedInJobCardParser()

    card_easy = make_mock_job_card(is_easy_apply=True)
    is_easy = await parser.check_easy_apply(card_easy)
    assert is_easy is True

    card_ext = make_mock_job_card(is_easy_apply=False, raw_snippet="Apply on company website")
    is_ext = await parser.check_easy_apply(card_ext)
    assert is_ext is False


@run_async
async def test_parse_job_cards_already_applied_badge():
    """Validates detection of already applied status in card footer."""
    parser = LinkedInJobCardParser()

    card_applied = make_mock_job_card(is_applied=True)
    assert await parser.check_already_applied(card_applied) is True

    card_unapplied = make_mock_job_card(is_applied=False)
    assert await parser.check_already_applied(card_unapplied) is False


def test_deduplication_cache_hit_and_skip(tmp_path):
    """Validates in-memory and disk deduplication check."""
    applied_f = str(tmp_path / "applied.json")
    seen_f = str(tmp_path / "seen.json")
    dedup = JobDeduplicator(applied_file=applied_f, seen_file=seen_f)

    assert dedup.is_applied("4159823456") is False
    assert dedup.is_seen("4159823456") is False

    dedup.mark_applied("4159823456", job_info={"title": "Azure Data Engineer"})
    assert dedup.is_applied("4159823456") is True
    assert dedup.is_seen("4159823456") is True

    dedup.mark_seen("4159823457", status="SKIPPED_LOW_MATCH")
    assert dedup.is_applied("4159823457") is False
    assert dedup.is_seen("4159823457") is True


def test_deduplication_persistence_roundtrip(tmp_path):
    """Validates atomic save and re-load from disk."""
    applied_f = str(tmp_path / "applied.json")
    seen_f = str(tmp_path / "seen.json")

    dedup1 = JobDeduplicator(applied_file=applied_f, seen_file=seen_f)
    dedup1.mark_applied("4159823451", job_info={"title": "Job 1"})
    dedup1.mark_seen("4159823452", job_info={"title": "Job 2"})

    # Create new instance pointing to same files
    dedup2 = JobDeduplicator(applied_file=applied_f, seen_file=seen_f)
    assert dedup2.is_applied("4159823451") is True
    assert dedup2.is_seen("4159823452") is True
    assert dedup2.is_applied("4159823452") is False


@run_async
async def test_parse_job_cards_missing_fields_graceful():
    """Validates fallback defaults when card has missing company, location, or malformed title."""
    card = MagicMock()
    attr_map = {"data-job-id": "4159823499"}
    card.get_attribute = AsyncMock(side_effect=lambda a: attr_map.get(a, None))
    card.inner_text = AsyncMock(return_value="")
    card.locator = MagicMock(return_value=make_mock_locator(count=0))

    parser = LinkedInJobCardParser()
    parsed = await parser.parse_single_card(card)
    assert parsed is not None
    assert parsed["job_id"] == "4159823499"
    assert parsed["title"] == "Azure Data Engineer"  # default fallback
    assert parsed["company"] == "Unknown Employer"  # default fallback
    assert parsed["location"] == "India"


@run_async
async def test_parse_job_cards_zero_results():
    """Validates that empty-results banner returns an empty list without crashing."""
    page = MagicMock()
    def page_locator(sel):
        if "no-results" in sel or "empty-state" in sel:
            return make_mock_locator(count=1, visible=True, text="No matching jobs found")
        return make_mock_locator(count=0, visible=False)

    page.locator = MagicMock(side_effect=page_locator)

    parser = LinkedInJobCardParser()
    results = await parser.parse_job_cards(page)
    assert results == []


def test_deduplication_corrupted_cache_recovery(tmp_path):
    """Validates that corrupted JSON cache files reinitialize safely without exceptions."""
    applied_f = str(tmp_path / "corrupt_applied.json")
    seen_f = str(tmp_path / "corrupt_seen.json")

    with open(applied_f, "w", encoding="utf-8") as f:
        f.write("{invalid json content")
    with open(seen_f, "w", encoding="utf-8") as f:
        f.write("[not a dict")

    dedup = JobDeduplicator(applied_file=applied_f, seen_file=seen_f)
    assert dedup.is_applied("123") is False
    assert dedup.is_seen("123") is False

    # Should be able to save new values normally
    dedup.mark_applied("4159823456")
    assert dedup.is_applied("4159823456") is True


def test_deduplication_filter_unapplied_jobs(tmp_path):
    """Validates filtering out already-applied roles from parsed job list."""
    dedup = JobDeduplicator(applied_file=str(tmp_path / "app.json"), seen_file=str(tmp_path / "seen.json"))
    dedup.mark_applied("4159823451")

    jobs = [
        {"job_id": "4159823451", "title": "Job 1"},
        {"job_id": "4159823452", "title": "Job 2"},
        {"job_id": "4159823453", "title": "Job 3"},
    ]
    unapplied = dedup.filter_unapplied_jobs(jobs)
    assert len(unapplied) == 2
    assert [j["job_id"] for j in unapplied] == ["4159823452", "4159823453"]


def test_deduplication_sync_from_excel_tracker(tmp_path, monkeypatch):
    """Validates synchronization of applied job IDs from Excel tracker."""
    tracker_file = str(tmp_path / "Tracker.xlsx")

    # Mock openpyxl workbook
    mock_wb = MagicMock()
    mock_wb.sheetnames = ["LinkedIn_Applications_Log", "Applications_Log"]
    ws_mock = MagicMock()
    ws_mock.iter_rows = MagicMock(return_value=[
        ("Header1", "Header2", "Job URL"),
        ("TCS", "Azure Data Engineer", "https://www.linkedin.com/jobs/view/4159823455/"),
        ("Infosys", "PySpark Developer", "https://www.linkedin.com/jobs/view/4159823456/"),
    ])
    mock_wb.__getitem__ = MagicMock(return_value=ws_mock)

    with patch("openpyxl.load_workbook", return_value=mock_wb), \
         patch("os.path.exists", return_value=True):
        dedup = JobDeduplicator(
            applied_file=str(tmp_path / "app.json"),
            seen_file=str(tmp_path / "seen.json"),
            tracker_path=tracker_file
        )
        count = dedup.sync_from_excel_tracker(tracker_file)
        assert count == 2
        assert dedup.is_applied("4159823455") is True
        assert dedup.is_applied("4159823456") is True


# ==============================================================================
# SECTION 8: MILESTONE 2 — STRICT PRECISION MATCH SCORING >=70% (FEATURE 8)
# ==============================================================================

def test_match_scorer_high_match_azure_data_engineer():
    """Validates near-perfect score (>= 90%) for ideal Azure Data Engineer role."""
    scorer = JobMatchScorer()
    title = "Azure Data Engineer"
    jd = """
    We are looking for an Azure Data Engineer with 2-4 years of experience.
    Must have hands-on experience in Azure Data Factory (ADF), Azure Databricks, PySpark,
    Delta Lake, Medallion Architecture, and strong SQL and Python skills.
    Knowledge of CI/CD and Docker is a plus.
    """
    res = scorer.evaluate(title, jd)
    assert res.score >= 90.0
    assert res.passed is True
    assert res.breakdown["title"] == 35.0
    assert res.breakdown["tech"] >= 40.0
    assert res.breakdown["experience"] == 20.0
    assert "Azure" in res.matched_skills
    assert "Databricks" in res.matched_skills
    assert "PySpark/Spark" in res.matched_skills


def test_match_scorer_borderline_pass():
    """Validates borderline pass (70.0% - 85.0%) for adjacent data engineer role."""
    scorer = JobMatchScorer()
    title = "Data Engineer"
    jd = """
    Looking for Data Engineer with 2-4 yrs experience.
    Required skills: Python, SQL, Apache Spark, Azure cloud, Docker and Star Schema modeling.
    """
    res = scorer.evaluate(title, jd)
    assert 70.0 <= res.score <= 85.0
    assert res.passed is True
    assert res.breakdown["title"] == 26.25  # Tier 2


def test_match_scorer_borderline_fail():
    """Validates borderline fail (< 70.0%) for AWS/Snowflake role without Azure ecosystem."""
    scorer = JobMatchScorer()
    title = "Data Engineer"
    jd = """
    Data Engineer needed with 3-5 years experience.
    Strong in Python, SQL, PostgreSQL, AWS Glue, AWS Redshift, and Snowflake.
    """
    res = scorer.evaluate(title, jd)
    assert res.score < 70.0
    assert res.passed is False
    assert any("GATE_REJECTED" in r for r in res.match_reasons)


def test_match_scorer_disqualified_negative_title():
    """Validates 0.0% score and rejection for completely irrelevant roles."""
    scorer = JobMatchScorer()
    irrelevant_roles = [
        ("Frontend React Developer", "Looking for React, Next.js, Redux, CSS, HTML5 developer."),
        ("Lead UI/UX Designer", "Figma, user research, wireframing, design systems."),
        ("Enterprise Sales Executive", "B2B SaaS quota carrying sales professional."),
        ("HR Talent Acquisition Partner", "Technical recruitment, sourcing candidates, Workday."),
    ]
    for title, jd in irrelevant_roles:
        res = scorer.evaluate(title, jd)
        assert res.score == 0.0
        assert res.passed is False
        assert any("DISQUALIFIED_TITLE" in r for r in res.match_reasons)


def test_match_scorer_seniority_mismatch():
    """Validates rejection for senior overqualified roles (10+ YOE)."""
    scorer = JobMatchScorer()
    title = "Data Engineer"
    jd = """
    Requires 10+ years of enterprise experience building petabyte scale platforms.
    Must have managed teams of 10+ engineers. Azure and PySpark required.
    """
    res = scorer.evaluate(title, jd)
    assert res.score <= 35.0
    assert res.passed is False
    assert any("EXP_DISQUALIFIED" in r for r in res.match_reasons)


def test_match_scorer_exact_boundary_69_9_vs_70_0():
    """Validates strict mathematical boundary gating at 70.0%."""
    # Test strict gating condition in MatchScoreResult creation
    res_fail = MatchScoreResult(
        score=69.9,
        passed=(69.9 >= 70.0),
        match_reasons=[],
        breakdown={},
        job_title="Test"
    )
    assert res_fail.passed is False

    res_pass = MatchScoreResult(
        score=70.0,
        passed=(70.0 >= 70.0),
        match_reasons=[],
        breakdown={},
        job_title="Test"
    )
    assert res_pass.passed is True


def test_match_scorer_empty_jd_fallback():
    """Validates handling of empty JD snippet without crashing."""
    scorer = JobMatchScorer()
    res = scorer.evaluate("Azure Data Engineer", "")
    assert res.score < 70.0
    assert res.passed is False
    assert res.breakdown["title"] == 35.0


def test_match_scorer_whitespace_and_case_insensitivity():
    """Validates that casing and surrounding whitespace do not impact scoring."""
    scorer = JobMatchScorer()
    res1 = scorer.evaluate(
        "   AZURE DATA ENGINEER   ",
        "  PYSPARK, DATABRICKS, ADF, DELTA LAKE, SQL, PYTHON. 2-4 YEARS.  "
    )
    res2 = scorer.evaluate(
        "azure data engineer",
        "pyspark, databricks, adf, delta lake, sql, python. 2-4 years."
    )
    assert res1.score == res2.score
    assert res1.passed is True


def test_match_scorer_custom_candidate_skills_override():
    """Validates that custom candidate skills supplement matched skills."""
    scorer = JobMatchScorer()
    res = scorer.evaluate(
        "Azure Data Engineer",
        "Experience in Azure, Databricks, PySpark, 2-3 yrs, and Kafka streaming.",
        candidate_skills=["Kafka", "Delta Lake", "Terraform"]
    )
    assert res.passed is True
    assert "Kafka" in res.matched_skills or any("kafka" in s.lower() for s in res.matched_skills)


def test_match_scorer_structured_breakdown_contract():
    """Validates that breakdown dictionary contains all components summing to total."""
    scorer = JobMatchScorer()
    res = scorer.evaluate("Azure Data Engineer", "Databricks, PySpark, ADF, SQL, Python, 2-4 yrs")
    assert "title" in res.breakdown
    assert "tech" in res.breakdown
    assert "experience" in res.breakdown
    assert "penalty" in res.breakdown
    assert "total" in res.breakdown
    assert isinstance(res.match_reasons, list)
    assert len(res.match_reasons) > 0


# ==============================================================================
# SECTION 9: TIER 3 — INTEGRATION CONTRACT TESTS
# ==============================================================================

@run_async
async def test_tier3_url_to_search_to_scoring_pipeline(tmp_path):
    """Tier 3: Pairwise contract between URL builder, search page DOM, card parser, and match scorer."""
    applied_f = str(tmp_path / "app.json")
    seen_f = str(tmp_path / "seen.json")
    dedup = JobDeduplicator(applied_file=applied_f, seen_file=seen_f)
    scorer = JobMatchScorer()

    card_high = make_mock_job_card(
        job_id="4159823401",
        title="Azure Data Engineer",
        company="Microsoft",
        raw_snippet="Databricks, PySpark, ADF, Delta Lake, SQL, 2-4 years"
    )
    card_low = make_mock_job_card(
        job_id="4159823402",
        title="Frontend React Developer",
        company="DesignCo",
        raw_snippet="React, CSS, HTML"
    )

    page = MagicMock()
    cards_list = [card_high, card_low]
    list_loc = MagicMock()
    list_loc.count = AsyncMock(return_value=len(cards_list))
    list_loc.nth = MagicMock(side_effect=lambda idx: cards_list[idx])

    def page_locator(sel):
        if "no-results" in sel or "empty-state" in sel:
            return make_mock_locator(count=0, visible=False)
        return list_loc

    page.locator = MagicMock(side_effect=page_locator)
    page.evaluate = AsyncMock(return_value=True)
    page.goto = AsyncMock(return_value=None)

    engine = LinkedInSearchEngine(deduplicator=dedup, scorer=scorer)
    qualified = await engine.search_jobs(page, keyword="Azure Data Engineer", location="Bengaluru")

    # High match must qualify, low match must be skipped
    assert len(qualified) == 1
    assert qualified[0]["job_id"] == "4159823401"
    assert qualified[0]["match_score"] >= 70.0
    assert dedup.is_seen("4159823401") is True
    assert dedup.is_seen("4159823402") is True


@run_async
async def test_tier3_dedup_filtering_before_scoring(tmp_path):
    """Tier 3: Verifies that previously applied jobs are filtered out before evaluation."""
    applied_f = str(tmp_path / "app.json")
    seen_f = str(tmp_path / "seen.json")
    dedup = JobDeduplicator(applied_file=applied_f, seen_file=seen_f)
    dedup.mark_applied("4159823401")

    card_applied = make_mock_job_card(job_id="4159823401", title="Azure Data Engineer")
    card_new = make_mock_job_card(
        job_id="4159823402",
        title="Azure Data Engineer",
        raw_snippet="Databricks, PySpark, ADF, SQL, 2-4 yrs"
    )

    page = MagicMock()
    cards_list = [card_applied, card_new]
    list_loc = MagicMock()
    list_loc.count = AsyncMock(return_value=len(cards_list))
    list_loc.nth = MagicMock(side_effect=lambda idx: cards_list[idx])

    def page_locator2(sel):
        if "no-results" in sel or "empty-state" in sel:
            return make_mock_locator(count=0, visible=False)
        return list_loc

    page.locator = MagicMock(side_effect=page_locator2)
    page.evaluate = AsyncMock(return_value=True)
    page.goto = AsyncMock(return_value=None)

    engine = LinkedInSearchEngine(deduplicator=dedup)
    qualified = await engine.search_jobs(page, keyword="Azure Data Engineer")

    assert len(qualified) == 1
    assert qualified[0]["job_id"] == "4159823402"


# ==============================================================================
# SECTION 10: TIER 4 — REAL-WORLD WORKLOAD SCENARIOS
# ==============================================================================

def test_tier4_multihub_search_matrix_and_deduplication():
    """Tier 4: Realistic multi-hub search matrix generation across 3 hubs and 2 keywords."""
    engine = LinkedInSearchEngine()
    matrix = engine.generate_search_matrix(
        keywords=["Azure Data Engineer", "Databricks"],
        locations=["Bengaluru", "Hyderabad", "Remote"]
    )
    assert len(matrix) == 6
    urls = [item["url"] for item in matrix]
    assert all("https://www.linkedin.com/jobs/search/?" in u for u in urls)
    assert all("f_AL=true" in u for u in urls)
    assert all("f_E=2,3" in u for u in urls)


# ==============================================================================
# SECTION 11: MILESTONE 3 — SCREENING QUESTIONNAIRE SOLVER TESTS
# ==============================================================================

def test_screening_solver_calibrated_skills_experience():
    """Milestone 3: Evaluates calibrated numeric experience across core tech skills."""
    solver = LinkedInScreeningSolver()
    
    assert solver.get_numeric_experience_for_skill("How many years of experience do you have with PySpark?") == "2"
    assert solver.get_numeric_experience_for_skill("Years of experience with Azure Databricks:") == "2"
    assert solver.get_numeric_experience_for_skill("Experience in Azure Data Factory (ADF)") == "2"
    assert solver.get_numeric_experience_for_skill("How many years of SQL experience?") == "2.5"
    assert solver.get_numeric_experience_for_skill("Kafka event streaming experience") == "1.5"
    assert solver.get_numeric_experience_for_skill("Total years of data engineering experience") == "2.5"


def test_screening_solver_text_numeric_questions():
    """Milestone 3: Solves common text and numeric screening questions."""
    solver = LinkedInScreeningSolver()
    
    # Contact info
    assert solver.determine_text_answer("What is your mobile phone number?") == "9122036484"
    assert solver.determine_text_answer("Country phone code") == "+91"
    assert solver.determine_text_answer("Email address") == "lavkushry@gmail.com"
    
    # Notice period
    assert solver.determine_text_answer("What is your notice period (in days)?") == "15"
    assert solver.determine_text_answer("Notice period") == "15 Days or less"
    
    # CTC
    assert solver.determine_text_answer("What is your current CTC (in LPA)?") == "12"
    assert solver.determine_text_answer("What is your expected salary (LPA)?") == "20"
    
    # Location & Education
    assert solver.determine_text_answer("What is your current city?") == "Hyderabad, India"
    assert solver.determine_text_answer("Highest degree or qualification:") == "Bachelor's Degree"
    assert solver.determine_text_answer("University or College attended:") == "Maharishi Dayanand University (MDU)"


def test_screening_solver_radio_questions():
    """Milestone 3: Evaluates binary affirmative/negative radio button questions."""
    solver = LinkedInScreeningSolver()
    
    # Affirmative standard responses
    assert solver.determine_radio_answer("Are you legally authorized to work in India?") is True
    assert solver.determine_radio_answer("Are you open to relocating to Bengaluru?") is True
    assert solver.determine_radio_answer("Are you comfortable working in a hybrid model (3 days in office)?") is True
    assert solver.determine_radio_answer("Do you consent to a standard background check?") is True
    assert solver.determine_radio_answer("Have you completed a Bachelor's degree?") is True
    
    # Negative standard responses (Visa sponsorship)
    assert solver.determine_radio_answer("Will you now or in the future require visa sponsorship?") is False
    assert solver.determine_radio_answer("Do you need sponsorship for an employment authorization?") is False


def test_screening_solver_dropdown_matching():
    """Milestone 3: Selects the most accurate option from select dropdowns."""
    solver = LinkedInScreeningSolver()
    
    # Notice period
    notice_opts = ["-- Select --", "Immediate", "15 days or less", "1 month", "2 months", "3 months"]
    assert solver.find_best_dropdown_option(notice_opts, "What is your official notice period?") == "15 days or less"
    
    # Education
    edu_opts = ["Select degree", "High School", "Associate", "Bachelor's Degree", "Master's Degree", "Doctorate"]
    assert solver.find_best_dropdown_option(edu_opts, "Highest level of education completed:") == "Bachelor's Degree"
    
    # Visa sponsorship
    sponsorship_opts = ["-- Choose --", "Yes, I require sponsorship", "No, I do not require sponsorship"]
    assert solver.find_best_dropdown_option(sponsorship_opts, "Do you require visa sponsorship?") == "No, I do not require sponsorship"
    
    # Work authorization
    auth_opts = ["Select", "Yes, I am authorized", "No, I am not authorized"]
    assert solver.find_best_dropdown_option(auth_opts, "Are you legally authorized to work in India?") == "Yes, I am authorized"
    
    # Experience bracket
    bracket_opts = ["0-1 years", "2-3 years", "4-6 years", "7+ years"]
    assert solver.find_best_dropdown_option(bracket_opts, "Years of experience in PySpark:") == "2-3 years"


@run_async
async def test_screening_solver_solve_step_inputs():
    """Milestone 3: Verifies that solver interacts with and fills visible form elements."""
    solver = LinkedInScreeningSolver()
    
    mock_input = MagicMock()
    mock_input.is_visible = AsyncMock(return_value=True)
    mock_input.input_value = AsyncMock(return_value="")
    mock_input.get_attribute = AsyncMock(side_effect=lambda a: "years-exp-input" if a == "id" else None)
    mock_input.fill = AsyncMock(return_value=None)
    
    mock_label = MagicMock()
    mock_label.inner_text = AsyncMock(return_value="How many years of PySpark experience do you have?")
    
    input_loc = MagicMock()
    input_loc.count = AsyncMock(return_value=1)
    input_loc.nth = MagicMock(return_value=mock_input)
    
    label_loc = MagicMock()
    label_loc.count = AsyncMock(return_value=1)
    label_loc.first = mock_label
    
    modal = MagicMock()
    def modal_locator(sel):
        if "label" in sel:
            return label_loc
        if sel == "input[type='text']":
            return input_loc
        return make_mock_locator(count=0, visible=False)
        
    modal.locator = MagicMock(side_effect=modal_locator)
    
    records = await solver.solve_all_questions_in_step(modal)
    assert len(records) == 1
    assert records[0]["type"] == "text"
    assert records[0]["answer"] == "2"
    mock_input.fill.assert_called_with("2")


# ==============================================================================
# SECTION 12: MILESTONE 3 — EASY APPLY MODAL NAVIGATOR TESTS
# ==============================================================================

@run_async
async def test_easy_apply_button_detection():
    """Milestone 3: Accurately distinguishes Easy Apply buttons from external Apply."""
    nav = LinkedInEasyApplyNavigator()
    
    # Page with native Easy Apply button
    easy_btn = make_mock_locator(count=1, visible=True, text="Easy Apply")
    page_easy = make_mock_page(locators={"button.jobs-apply-button": easy_btn})
    assert await nav.is_easy_apply_available(page_easy) is True
    
    # Page with external Apply button
    ext_btn = make_mock_locator(count=1, visible=True, text="Apply on company website")
    page_ext = make_mock_page(locators={"button.jobs-apply-button": ext_btn})
    assert await nav.is_easy_apply_available(page_ext) is False


@run_async
async def test_easy_apply_uncheck_follow_company():
    """Milestone 3: Unchecks 'Follow company' checkbox before application submission."""
    nav = LinkedInEasyApplyNavigator()
    
    mock_checkbox = MagicMock()
    mock_checkbox.is_checked = AsyncMock(return_value=True)
    mock_checkbox.uncheck = AsyncMock(return_value=None)
    
    chk_loc = MagicMock()
    chk_loc.count = AsyncMock(return_value=1)
    chk_loc.nth = MagicMock(return_value=mock_checkbox)
    
    modal = MagicMock()
    modal.locator = MagicMock(return_value=chk_loc)
    
    res = await nav.uncheck_follow_company(modal)
    assert res is True
    mock_checkbox.uncheck.assert_called_once()


@run_async
async def test_easy_apply_dry_run_workflow(tmp_path):
    """Milestone 3: Executes multi-step dry-run without submitting, returning DRY_RUN_VERIFIED."""
    nav = LinkedInEasyApplyNavigator()
    
    # Create dummy resume file
    resume_file = tmp_path / "Lavkush_Kumar_Resume.pdf"
    resume_file.write_text("Dummy Resume Content", encoding="utf-8")
    
    # Mock modal buttons
    apply_btn = make_mock_locator(count=1, visible=True, text="Easy Apply")
    apply_btn.first.click = AsyncMock(return_value=None)
    
    submit_btn = make_mock_locator(count=1, visible=True, text="Submit application")
    dismiss_btn = make_mock_locator(count=1, visible=True, text="Dismiss")
    dismiss_btn.first.click = AsyncMock(return_value=None)
    
    discard_btn = make_mock_locator(count=1, visible=True, text="Discard")
    discard_btn.first.click = AsyncMock(return_value=None)
    
    modal_elem = MagicMock()
    def modal_loc(sel):
        if "submit" in sel.lower():
            return submit_btn
        if "input" in sel:
            return make_mock_locator(count=0, visible=False)
        return make_mock_locator(count=0, visible=False)
    modal_elem.locator = MagicMock(side_effect=modal_loc)
    
    modal_loc_page = make_mock_locator(count=1, visible=True)
    modal_loc_page.first = modal_elem
    
    page = MagicMock()
    def page_loc(sel):
        if "jobs-apply-button" in sel:
            return apply_btn
        if "jobs-easy-apply-modal" in sel or "[data-test-modal]" in sel:
            return modal_loc_page
        if "dismiss" in sel.lower():
            return dismiss_btn
        if "discard" in sel.lower():
            return discard_btn
        return make_mock_locator(count=0, visible=False)
    page.locator = MagicMock(side_effect=page_loc)
    
    job_info = {
        "job_id": "4199990001",
        "title": "Azure Data Engineer",
        "company": "Tech Corp",
        "url": "https://www.linkedin.com/jobs/view/4199990001/"
    }
    
    with patch("naukari_pro.modules.linkedin_apply.humanized_delay", AsyncMock()):
        result = await nav.execute_easy_apply(
            page=page,
            job_info=job_info,
            resume_path=str(resume_file),
            dry_run=True
        )
    
    assert result["success"] is True
    assert result["status"] == "DRY_RUN_VERIFIED"
    assert result["job_id"] == "4199990001"
    # Modal dismiss must be invoked
    assert dismiss_btn.first.click.called


@run_async
async def test_easy_apply_live_submission_workflow(tmp_path):
    """Milestone 3: Executes live application submission and confirms DOM banner."""
    nav = LinkedInEasyApplyNavigator()
    
    resume_file = tmp_path / "Lavkush_Kumar_Resume.pdf"
    resume_file.write_text("Dummy Resume Content", encoding="utf-8")
    
    apply_btn = make_mock_locator(count=1, visible=True, text="Easy Apply")
    apply_btn.first.click = AsyncMock(return_value=None)
    
    submit_btn = make_mock_locator(count=1, visible=True, text="Submit application")
    submit_btn.first.click = AsyncMock(return_value=None)
    
    confirm_banner = make_mock_locator(count=1, visible=True, text="Application submitted")
    dismiss_btn = make_mock_locator(count=1, visible=True, text="Dismiss")
    dismiss_btn.first.click = AsyncMock(return_value=None)
    
    modal_elem = MagicMock()
    def modal_loc(sel):
        if "submit" in sel.lower():
            return submit_btn
        return make_mock_locator(count=0, visible=False)
    modal_elem.locator = MagicMock(side_effect=modal_loc)
    
    modal_loc_page = make_mock_locator(count=1, visible=True)
    modal_loc_page.first = modal_elem
    
    page = MagicMock()
    def page_loc(sel):
        if "jobs-apply-button" in sel:
            return apply_btn
        if "jobs-easy-apply-modal" in sel or "[data-test-modal]" in sel:
            return modal_loc_page
        if "Application submitted" in sel:
            return confirm_banner
        if "dismiss" in sel.lower():
            return dismiss_btn
        return make_mock_locator(count=0, visible=False)
    page.locator = MagicMock(side_effect=page_loc)
    
    job_info = {
        "job_id": "4199990002",
        "title": "Senior Azure Data Engineer",
        "company": "Enterprise Tech",
        "url": "https://www.linkedin.com/jobs/view/4199990002/"
    }
    
    with patch("naukari_pro.modules.linkedin_apply.humanized_delay", AsyncMock()):
        result = await nav.execute_easy_apply(
            page=page,
            job_info=job_info,
            resume_path=str(resume_file),
            dry_run=False
        )
    
    assert result["success"] is True
    assert result["status"] == "APPLIED"
    assert submit_btn.first.click.called


# ==============================================================================
# SECTION 13: MILESTONE 4 — TRACKER LOGGING & DEDUPLICATION TESTS
# ==============================================================================

def test_tracker_excel_sheet_creation_and_logging(tmp_path):
    """Milestone 4: Confirms Excel logging to LinkedIn_Applications_Log and JSON sync."""
    excel_path = str(tmp_path / "Test_Tracker.xlsx")
    json_path = str(tmp_path / "linkedin_applications_log.json")
    skip_path = str(tmp_path / "skip_list_linkedin.json")
    
    tracker = LinkedInApplicationTracker(
        excel_path=excel_path,
        json_log_path=json_path,
        skip_file_path=skip_path
    )
    
    job_sample = {
        "job_id": "4188880001",
        "title": "Azure Data Engineer",
        "company": "Cloud GCC",
        "location": "Bengaluru, Karnataka, India",
        "url": "https://www.linkedin.com/jobs/view/4188880001/"
    }
    
    tracker.log_application(job_sample, "APPLIED", match_score=88.5, steps=3)
    
    # 1. Verify Excel worksheet and headers
    import openpyxl
    wb = openpyxl.load_workbook(excel_path)
    assert "LinkedIn_Applications_Log" in wb.sheetnames
    ws = wb["LinkedIn_Applications_Log"]
    assert ws.max_row == 2
    row_vals = [cell.value for cell in ws[2]]
    assert row_vals[1] == "Cloud GCC"
    assert row_vals[2] == "Azure Data Engineer"
    assert row_vals[3] == "88%"
    assert row_vals[6] == "LinkedIn Easy Apply"
    assert row_vals[7] == "APPLIED"
    
    # 2. Verify JSON log
    assert os.path.exists(json_path)
    with open(json_path, "r", encoding="utf-8") as f:
        json_data = json.load(f)
        assert len(json_data) == 1
        assert json_data[0]["job_id"] == "4188880001"
        assert json_data[0]["status"] == "APPLIED"
        
    # 3. Verify Skip List
    assert tracker.is_applied("4188880001") is True
    assert tracker.is_applied("nonexistent_id") is False


# ==============================================================================
# SECTION 14: MILESTONE 4 — MASTER APPLY ENGINE & CLI INTEGRATION
# ==============================================================================

@run_async
async def test_apply_engine_run_cycle_dry_run(tmp_path, monkeypatch):
    """Milestone 4: Orchestrates discovery, scoring gating (>= 70.0%), and dry-run apply cycle."""
    monkeypatch.setenv("WORKSPACE_DIR", str(tmp_path))
    monkeypatch.setenv("LINKEDIN_LI_AT", "AQEDAQ_DUMMY_TOKEN")
    
    lock_file = str(tmp_path / ".linkedin_session.lock")
    session_mgr = LinkedInSessionManager(
        session_dir=str(tmp_path / "session"),
        lock_path=lock_file
    )
    
    # Mock vacancies (one high match, one gated out < 70%)
    job_qualified = {
        "job_id": "4177770001",
        "title": "Azure Data Engineer",
        "company": "Fintech GCC",
        "location": "Hyderabad",
        "url": "https://www.linkedin.com/jobs/view/4177770001/",
        "is_easy_apply": True,
        "raw_snippet": "PySpark, Databricks, ADF, SQL, Delta Lake, 2-4 yrs exp"
    }
    job_gated = {
        "job_id": "4177770002",
        "title": "Frontend React Engineer",
        "company": "Web Agency",
        "location": "Mumbai",
        "url": "https://www.linkedin.com/jobs/view/4177770002/",
        "is_easy_apply": True,
        "raw_snippet": "React, JavaScript, CSS, HTML"
    }
    
    mock_search = MagicMock()
    mock_search.search_jobs = AsyncMock(return_value=[job_qualified, job_gated])
    
    mock_nav = MagicMock()
    mock_nav.execute_easy_apply = AsyncMock(return_value={
        "success": True,
        "status": "DRY_RUN_VERIFIED",
        "job_id": "4177770001",
        "steps": 3
    })
    
    tracker = LinkedInApplicationTracker(
        excel_path=str(tmp_path / "Tracker.xlsx"),
        json_log_path=str(tmp_path / "log.json"),
        skip_file_path=str(tmp_path / "skip.json")
    )
    
    engine = LinkedInApplyEngine(
        session_manager=session_mgr,
        search_engine=mock_search,
        navigator=mock_nav,
        tracker=tracker,
        min_score=70.0
    )
    
    # Mock session verify
    session_mgr.verify_session = AsyncMock(return_value={
        "status": "VALID",
        "authenticated": True,
        "candidate_name": "Lavkush Kumar"
    })
    
    mock_page = make_mock_page()
    mock_ctx = MagicMock(pages=[mock_page], close=AsyncMock())
    session_mgr.create_stealth_context = AsyncMock(return_value=mock_ctx)
    
    with patch("naukari_pro.modules.linkedin_apply.humanized_delay", AsyncMock()), \
         patch("playwright.async_api.async_playwright", make_playwright_cm_mock(mock_ctx)):
        
        stats = await engine.run_apply_cycle(
            keywords=["Azure Data Engineer"],
            locations=["Bengaluru"],
            limit=5,
            dry_run=True
        )
    
    assert stats["dry_run_verified"] == 1
    assert stats["skipped"] >= 1  # The frontend React job was gated out!
    assert len(stats["jobs"]) == 1
    assert stats["jobs"][0]["job_id"] == "4177770001"
    
    # Session lock must be released
    assert session_mgr.get_lock()._acquired is False


def test_cli_handle_linkedin_apply(monkeypatch, capsys):
    """Milestone 4: CLI command handler executes cleanly with structured output."""
    class MockArgs:
        limit = 5
        dry_run = True
        headed = False
        min_score = 70.0
        keyword = "Azure Data Engineer"
        location = "Bengaluru"
        
    mock_stats = {
        "discovered": 10,
        "scored_above_threshold": 4,
        "applied": 0,
        "dry_run_verified": 4,
        "skipped": 6,
        "failed": 0,
        "jobs": []
    }
    
    with patch("naukari_pro.modules.linkedin_apply.LinkedInApplyEngine.run_apply_cycle", new_callable=AsyncMock) as mock_cycle:
        mock_cycle.return_value = mock_stats
        
        exit_code = handle_linkedin_apply(MockArgs())
        assert exit_code == 0
        
        out = capsys.readouterr().out
        assert "LINKEDIN AUTONOMOUS EASY APPLY ENGINE" in out
        assert "DRY-RUN SIMULATION" in out
        assert "LINKEDIN APPLY CYCLE RESULTS SUMMARY" in out
        assert "Dry-Run Verified:      4" in out
        assert "LinkedIn_Applications_Log" in out
